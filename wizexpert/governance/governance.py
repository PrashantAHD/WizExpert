"""Create the ten-slide operational deck from a reviewed Wiz snapshot."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
import textwrap
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from wizexpert.governance.assignment_history import action_lines, write_actions_csv
from wizexpert.governance.dashboard import full_date, pct, replace_text, shape
from wizexpert.governance.dashboard_data import DashboardError, require
from wizexpert.governance.governance_data import (
    DISPLAY_SEVERITIES,
    GovernanceData,
    closed_counts,
    load_snapshot,
)
from wizexpert.governance.monthly_files import (
    check_recovery,
    check_unchanged,
    digest,
    monthly_lock,
    publish_pair,
)
from wizexpert.governance.profile import active as active_profile
from wizexpert.governance.profile import add_profile_argument
from wizexpert.governance.profile import ensure as ensure_profile
from wizexpert.governance.resolution_evidence import changed_text, risk_text

TEXT_IDS = {
    0: (4, 5, 6, 7, 11),
    1: (4, 5, 6, 8),
    2: (4, 5, 6, 9, 12, 15, 18, 20),
    3: (4, 5, 6, 9, 12, 15, 18, 21, 25, 26),
    4: (4, 5, 6),
    5: (4, 5, 6),
    6: (4, 5, 6),
    7: (4, 5, 6),
    8: (4, 5, 6),
    9: (4, 5, 6, 8),
}
TABLE_CAPACITY = {4: 4, 5: 6, 6: 4, 7: 6, 8: 4}
TABLE_BODY_FONT_SIZE = 12


def validate_template(deck) -> None:
    require(
        len(deck.slides) == 10 and (deck.slide_width, deck.slide_height) == (12192000, 6858000),
        "Governance template must be the approved ten-slide 16:9 deck.",
    )
    for index, ids in TEXT_IDS.items():
        for sid in ids:
            require(
                shape(deck.slides[index], sid).has_text_frame,
                f"Slide {index + 1}, shape {sid}: missing text placeholder.",
            )
    require(
        shape(deck.slides[0], 5).text == "Monthly Cloud Security Governance",
        "Operational governance cover not recognized; do not use the executive dashboard template.",
    )
    for index, capacity in TABLE_CAPACITY.items():
        target = shape(deck.slides[index], 7)
        require(
            target.has_table
            and len(target.table.rows) == capacity + 1
            and len(target.table.columns) == 4,
            f"Slide {index + 1}: expected a four-column table with {capacity} data rows.",
        )


def metric_text(value: dict) -> str:
    return "Unavailable" if value["value"] is None else pct(value["value"])


def delta(before: int, after: int) -> str:
    if after == before:
        return "unchanged"
    absolute = f"{abs(after - before):,} {'fewer' if after < before else 'more'}"
    return absolute if before == 0 else f"{absolute} ({pct(abs(after - before) / before * 100)})"


def write_text(slide, sid: int, paragraphs: list[str], size: int | None = None) -> None:
    replace_text(slide, sid, paragraphs)
    frame = shape(slide, sid).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    if size is not None:
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size)


def cell_text(cell, value: str, size: int) -> None:
    paragraph = cell.text_frame.paragraphs[0]
    first = next(iter(paragraph.runs), None)
    style = deepcopy(first._r.find(qn("a:rPr"))) if first is not None else None
    cell.text = value
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.text_frame.word_wrap = True
    for paragraph in cell.text_frame.paragraphs:
        paragraph.space_before = paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            if style is not None:
                run._r.insert(0, deepcopy(style))
            run.font.size = Pt(size)


def excerpt(value: str, width: int = 175) -> str:
    return textwrap.shorten(" ".join(value.split()), width=width, placeholder="...")


def first_sentence(value: str, width: int = 175) -> str:
    """First complete sentence of a recommendation, so slides never end mid-thought."""
    text = " ".join(value.split())
    match = re.match(r"(.+?[.!?])(\s|$)", text)
    sentence = match.group(1) if match else text
    return sentence if len(sentence) <= width else excerpt(sentence, width)


GENERIC_BULLETS = (
    "this finding",
    "use the \u201cleast privilege\u201d",
    'use the "least privilege"',
    "to find unused permissions",
    "find unused permissions",
    "minimize the attack surface",
    "avoid using access keys",
    "enable mfa",
    "sensitive information should not",
    "for detailed remediation steps",
)


def next_step(recommendation: str, rule_name: str = "") -> str:
    """First concrete action in a Wiz recommendation, as one complete sentence.

    Wiz recommendations are templates shared across rules, so the first bullet is often
    generic advice or a description of the finding. Skip those and sub-headings and keep
    Wiz's own order; fall back to the first bullet when nothing else remains. Markdown code
    formatting is removed.
    """
    bullets = [
        line.strip().lstrip("*- ").strip()
        for line in recommendation.splitlines()
        if line.strip().startswith(("*", "-"))
    ]
    candidates = [
        b
        for b in bullets
        if b and not b.endswith(":") and not b.lower().startswith(GENERIC_BULLETS)
    ]
    chosen = (candidates or bullets or [recommendation])[0]
    return first_sentence(chosen.replace("`", ""))


def table_row(
    group: dict,
    resolved: bool,
    rule_names: list[str] | None = None,
    credited: bool = False,
    owner_status: dict[str, str] | None = None,
    rule_status: dict[str, str] | None = None,
    wording: dict[str, dict[str, str]] | None = None,
    next_steps: dict[str, str] | None = None,
) -> tuple[str, str, str, str]:
    closed = closed_counts(group)
    count = sum(closed[s] for s in DISPLAY_SEVERITIES)
    noun = "finding" if count == 1 else "findings"
    severity = "/".join(s.title() for s in DISPLAY_SEVERITIES if closed[s])
    annotation = group.get("annotation")
    assignments = group.get("assignments")
    names = (
        list(
            dict.fromkeys(
                entry["name"] for entry in assignments["entries"] if entry["id"] is not None
            )
        )
        if assignments
        else []
    )
    if rule_names is not None:
        names = rule_names
    assignees = "/".join(names) if names else "Assignee not verified"
    if (
        rule_names is None
        and assignments
        and not names
        and sum(entry["count"] for entry in assignments["entries"]) == count
    ):
        assignees = "Unassigned"
    title = f"{group['name'].strip()}\n({assignees}) ({count:,}\u00a0{noun})"
    if resolved and credited:
        resolved_count = sum(group["counts"][s] for s in DISPLAY_SEVERITIES)
        accepted_count = count - resolved_count
        # Counts appear once, in the issue column; the other columns stay plain text.
        parts = ([f"{resolved_count:,}\u00a0resolved"] if resolved_count else []) + (
            [f"{accepted_count:,}\u00a0accepted\u00a0risk"] if accepted_count else []
        )
        title = f"{group['name'].strip()}\n({assignees}) ({', '.join(parts)})"
        resolution = group.get("resolution")
        team = (wording or {}).get(group["id"], {})
        if resolution:
            change = changed_text(resolution, resolved_count)
            risk = risk_text(resolution, resolved_count, accepted_count)
        else:
            change = "Change details: team to confirm."
            risk = "Closure status is confirmed; actual risk reduction needs evidence."
        change = team.get("changed", change)
        risk = team.get("risk", risk)
        return title, severity, change, excerpt(risk, 200)
    if resolved:
        # Current ownership cannot establish assignment or contribution before closure.
        title = f"{group['name'].strip()}\n({count:,}\u00a0{noun})"
        if annotation:
            change = f"Evidence covers {len(annotation['issue_ids'])} of {count:,}.\n" + excerpt(
                annotation["change"]
            )
            risk = excerpt(annotation["risk"])
        else:
            change = f"Wiz recorded {count:,} resolved.\nSpecific remediation not verified."
            risk = "Closure status is confirmed; actual risk reduction needs evidence."
        return title, severity, change, risk
    recommendation = group.get("recommendation")
    if recommendation:
        action = next_step(recommendation, group["name"])
        action = action if action.endswith((".", "!", "?")) else action + "."
    else:
        action = "Review Wiz recommendation; confirm remediation owner and target date."
    pending = "ETA and blockers" if names else "Assignee, ETA and blockers"
    # The issue column already carries the count, so the status is text only.
    status = f"{pending} need team confirmation."
    updates = [
        (name, owner_status[name]) for name in names if owner_status and name in owner_status
    ]
    if updates:
        if len(updates) == len(names) and len({s for _, s in updates}) == 1:
            status = f"Status: {updates[0][1]}."
        else:
            status = "; ".join(f"{name}: {state}" for name, state in updates) + "."
    if rule_status and group["id"] in rule_status:
        status = rule_status[group["id"]]
    if annotation:
        status = f"Update covers {len(annotation['issue_ids'])} of {count:,}.\n" + excerpt(
            annotation["change"]
        )
    if next_steps and group["id"] in next_steps:
        return title, severity, next_steps[group["id"]], status
    return title, severity, action, status


def fill_tables(deck, data: GovernanceData, field: str, indices: tuple[int, ...]) -> None:
    credited = field == "resolved_groups" and "team_resolved" in data.payload
    groups = data.ranked("team_resolved" if credited else field)
    if field == "in_progress_groups":
        # Rules the owner is not actively working this period stay in the counts but off the slides.
        deferred = {entry["id"] for entry in data.payload.get("deferred_in_progress", [])}
        groups = [g for g in groups if g["id"] not in deferred]
    # Rows beyond table capacity are omitted from the slides; the full list stays in the
    # notes/companion JSON.
    shown = groups[: sum(TABLE_CAPACITY[i] for i in indices)]
    resolved = field == "resolved_groups"
    rule_names = (
        data.rule_assignees()
        if data.payload.get("assignment_scope") == "RULE" and not credited
        else None
    )
    offset = 0
    for page, index in enumerate(indices, 1):
        slide = deck.slides[index]
        table = shape(slide, 7).table
        shape(slide, 7).top = Inches(0.9)
        table.rows[0].height = Inches(0.4)
        for row in list(table.rows)[1:]:
            row.height = Inches(5.05 / TABLE_CAPACITY[index])
        chunk = shown[offset : offset + TABLE_CAPACITY[index]]
        offset += len(chunk)
        title = (
            (
                ("Resolved / Risk Accepted by Team" if data.team_accepted() else "Resolved by Team")
                if credited
                else "Tenant Closures - Attribution Pending"
            )
            if resolved
            else "Work in Progress - Priorities"
        )
        write_text(slide, 4, [f"{title} ({page}/{len(indices)})"], 24)
        headers = (
            ("Resolved issue", "Severity", "Closure evidence", "Risk reduction")
            if resolved
            else ("In-progress issue", "Severity", "Next step", "Current status")
        )
        widths = (4.0, 0.9, 4.0, shape(slide, 7).width / 914400 - 8.9)
        for column, width in zip(table.columns, widths, strict=False):
            column.width = Inches(width)
        for col, header in enumerate(headers):
            cell_text(table.cell(0, col), header, 14)
        for row in range(1, len(table.rows)):
            group = chunk[row - 1] if row <= len(chunk) else None
            names = (
                rule_names.get(group["id"], [])
                if rule_names is not None and group is not None
                else None
            )
            values = (
                table_row(
                    group,
                    resolved,
                    names,
                    credited,
                    None if resolved else data.owner_status(),
                    None if resolved else data.rule_status(),
                    data.closure_wording() if resolved else None,
                    None if resolved else data.rule_next_step(),
                )
                if group is not None
                else ("", "", "", "")
            )
            if not chunk and row == 1:
                values = ("No findings in this section.", "", "", "")
            for col, value in enumerate(values):
                cell_text(table.cell(row, col), value, TABLE_BODY_FONT_SIZE)


def trend_statement(current: dict, before: dict, in_progress: int, prior_month: str) -> list[str]:
    """Narrative posture trend in the style of the approved deck's "Trend statement" box.

    Returns a heading plus 3 lines.
    """
    top, old_top = current["CRITICAL"] + current["HIGH"], before["CRITICAL"] + before["HIGH"]
    if top == 0:
        severity = (
            "Critical and High remain at zero, sustaining the downward trend."
            if old_top == 0
            else "Critical and High findings dropped to zero."
        )
    else:
        severity = f"Critical and High: {top:,} open ({delta(old_top, top)} vs {prior_month})."
    total, old_total = (
        sum(current[s] for s in DISPLAY_SEVERITIES),
        sum(before[s] for s in DISPLAY_SEVERITIES),
    )
    largest = max(("MEDIUM", "LOW"), key=lambda s: current[s])
    share = pct(current[largest] / total * 100) if total else "0%"
    focus = (
        f"Focus: {largest.title()} remains the largest volume "
        f"({current[largest]:,}; {share} of backlog), "
        f"{in_progress:,} in progress."
    )
    outlook = (
        "Medium and Low still call for steady remediation focus."
        if top <= old_top
        else "Prioritise top-severity findings ahead of Medium and Low."
    )
    overall = f"Backlog {delta(old_total, total)} since {prior_month}; {outlook}"
    return ["Trend statement", severity, focus, overall]


def metric_phrase(label: str, value: dict, prior: dict) -> str:
    """E.g. ``Wiz Security Score held at 100%`` or ``HIPAA compliance rose to 93% (up 1 point)``."""
    now, before = value["value"], prior["value"]
    if now is None:
        return f"{label} unavailable"
    if before is None:
        return f"{label} at {pct(now)}"
    if now == before:
        return f"{label} held at {pct(now)}"
    change = abs(now - before)
    unit = "point" if change == 1 else "points"
    return (
        f"{label} {'rose' if now > before else 'fell'} to {pct(now)} "
        f"({'up' if now > before else 'down'} {change:g} {unit})"
    )


def client_messages(data: GovernanceData) -> tuple[list[str], list[str]]:
    """Client-facing key messages (slide 3) and key takeaways (slide 10)."""
    payload, prior = data.payload, data.payload["prior"]
    current, before = payload["open"], prior["open"]
    total, old_total = data.total("open"), sum(before[s] for s in DISPLAY_SEVERITIES)
    prior_month = data.start.strftime("%B")
    top, old_top = current["CRITICAL"] + current["HIGH"], before["CRITICAL"] + before["HIGH"]
    if top == 0:
        severity = "Critical and High findings " + (
            "remain at zero" if old_top == 0 else "dropped to zero"
        )
    else:
        severity = (
            f"Critical and High findings stand at {top:,} "
            f"({delta(old_top, top)} than {prior_month})"
        )
    score = metric_phrase(
        "the Wiz Security Score", payload["security_score"], prior["security_score"]
    )
    hipaa = metric_phrase("HIPAA compliance", payload["hipaa"], prior["hipaa"])
    resolved, accepted = data.total("resolved"), data.total("accepted")
    closed = resolved + accepted
    split = f"{resolved:,} remediated" + (f", {accepted:,} risk-accepted" if accepted else "")
    backlog = (
        f"reducing the open backlog to {total:,}, {delta(old_total, total)} than {prior_month}"
        if total < old_total
        else f"with the open backlog at {total:,}, {delta(old_total, total)} than {prior_month}"
        if total > old_total
        else f"with the open backlog unchanged at {total:,}"
    )
    groups = data.ranked("in_progress_groups")[:2]
    drivers = " and ".join(
        f"{g['name'].strip()} ({sum(g['counts'][s] for s in DISPLAY_SEVERITIES):,})" for g in groups
    )
    in_progress = data.total("in_progress")
    largest = max(("MEDIUM", "LOW"), key=lambda s: current[s])
    focus = (
        "sustain zero Critical and High findings"
        if top == 0
        else "close the remaining Critical and High findings"
    )
    messages = [
        f"{severity}; {score} and {hipaa}.",
        (
            f"{closed:,} {'finding was' if closed == 1 else 'findings were'} "
            f"closed this period ({split}), {backlog}."
        ),
        (
            f"Remediation is under way on {in_progress:,} findings, led by {drivers}."
            if in_progress
            else "No findings are currently in remediation."
        ),
        f"Next-month focus: {focus} and keep reducing the {largest.title()} backlog "
        f"({current[largest]:,} open).",
    ]
    takeaways = [
        f"{score[0].upper()}{score[1:]}.",
        (
            "Critical and High issues remain at zero."
            if top == 0 and old_top == 0
            else f"Critical {current['CRITICAL']:,} "
            f"({delta(before['CRITICAL'], current['CRITICAL'])}); "
            f"High {current['HIGH']:,} ({delta(before['HIGH'], current['HIGH'])})."
        ),
        f"{hipaa}; HITRUST CSF at {metric_text(payload['hitrust'])}.",
        f"{closed:,} findings closed this period ({split}); open backlog "
        f"{'down' if total < old_total else 'up'} "
        f"{pct(abs(total - old_total) / old_total * 100)} to {total:,}."
        if total != old_total and old_total
        else f"{closed:,} findings closed this period ({split}); open backlog at {total:,}.",
        f"Next focus: {focus}; progress the {in_progress:,} in-progress findings"
        + (f", led by {groups[0]['name'].strip()}." if groups else "."),
    ]
    return messages, takeaways


def render_deck(data: GovernanceData, template: Path, destination: Path, source_hash: str) -> None:
    deck = Presentation(template)
    validate_template(deck)
    payload, prior = data.payload, data.payload["prior"]
    require(
        prior["source_sha256"] == digest(template),
        "Prior comparison source hash must match the approved template. "
        "Review prior figures before rerunning.",
    )
    # Remove stale narrative from non-content shapes, including hidden background text.
    for index, slide in enumerate(deck.slides):
        for target in list(slide.shapes):
            # Coverage captions from older generated decks are removed; captions are no
            # longer shown.
            if target.name == "Governance coverage" or (
                target.has_table and (index not in TABLE_CAPACITY or target.shape_id != 7)
            ):
                target._element.getparent().remove(target._element)
            elif target.has_text_frame and target.shape_id not in TEXT_IDS[index]:
                target.text_frame.clear()
        background = shape(slide, 2)
        background.top, background.height = 0, deck.slide_height

    as_of = data.backlog_as_of.strftime("%b %d, %Y")
    label = (
        f"PRELIMINARY | Data collected through {as_of}"
        if data.preliminary
        else f"Data as of {as_of} EOD"
    )
    footer = f"Confidential - Proprietary to {active_profile().provider_name} | {label}"
    for index, slide in enumerate(deck.slides):
        footer_id = 11 if index == 0 else 6
        write_text(slide, footer_id, [footer], 10)
        footer_frame = shape(slide, footer_id).text_frame
        footer_frame.margin_top = footer_frame.margin_bottom = 0
        if index:
            heading = shape(slide, 4)
            heading.width, heading.height = Inches(10), Inches(0.46)
            heading.text_frame.margin_top = heading.text_frame.margin_bottom = 0
    write_text(
        deck.slides[0], 6, [f"{full_date(data.start.date())} - {full_date(data.end.date())}"], 20
    )
    shape(deck.slides[0], 6).width = Inches(10)
    if data.preliminary:
        period_note = "Scheduled period end; live snapshot, not historical EOD."
    else:
        period_note = (
            f"Reporting period closed {full_date(data.end.date())} EOD; "
            "later changes roll into the next report."
        )
    write_text(deck.slides[0], 7, [label, period_note], 16)
    shape(deck.slides[0], 7).width = Inches(11.4)
    write_text(
        deck.slides[1],
        8,
        [
            "Executive summary",
            "Current cloud security posture",
            "Resolved findings and evidence",
            "Work in progress",
            "Key takeaways",
        ],
    )
    current, before = payload["open"], prior["open"]
    total = data.total("open")
    score, old_score = payload["security_score"]["value"], prior["security_score"]["value"]
    score_change = "Comparison unavailable"
    if score is not None and old_score is not None:
        score_change = (
            "No change vs prior" if score == old_score else f"{score - old_score:+g} pp vs prior"
        )
    summary = deck.slides[2]
    write_text(
        summary, 9, ["Wiz Security Score", metric_text(payload["security_score"]), score_change]
    )
    for sid, severity in ((12, "CRITICAL"), (15, "HIGH")):
        write_text(
            summary,
            sid,
            [
                f"{severity.title()} Issues",
                f"{current[severity]:,}",
                delta(before[severity], current[severity]),
            ],
        )
    write_text(
        summary,
        18,
        [
            "Compliance",
            f"HIPAA {metric_text(payload['hipaa'])}",
            f"HITRUST {metric_text(payload['hitrust'])}",
        ],
    )
    for sid in (9, 12, 15, 18):
        frame = shape(summary, sid).text_frame
        value_size = 18 if "Unavailable" in frame.paragraphs[1].text else 28
        for paragraph, size in zip(frame.paragraphs, (14, value_size, 11), strict=False):
            for run in paragraph.runs:
                run.font.size = Pt(size)
    team = "team_resolved" in payload
    messages, takeaways = client_messages(data)
    write_text(summary, 20, ["Key messages", *messages], 19)
    posture = deck.slides[3]
    for sid, heading, number in (
        (9, "Open backlog", total),
        (12, "Critical", current["CRITICAL"]),
        (15, "High", current["HIGH"]),
        (18, "Medium", current["MEDIUM"]),
        (21, "Low", current["LOW"]),
    ):
        write_text(posture, sid, [heading, f"{number:,}", "C/H/M/L only" if sid == 9 else as_of])
        target = shape(posture, sid)
        target.width = Inches(2.05 if sid == 9 else 1.8)
        target.height = Inches(1.0)
        target.text_frame.margin_top = target.text_frame.margin_bottom = 0
        for paragraph, size in zip(target.text_frame.paragraphs, (14, 28, 11), strict=False):
            paragraph.space_before = paragraph.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(size)
    write_text(
        posture,
        26,
        [
            "Posture highlights",
            f"Critical: {current['CRITICAL']:,}; High: {current['HIGH']:,}.",
            f"Medium: {delta(before['MEDIUM'], current['MEDIUM'])}.",
            f"Low: {delta(before['LOW'], current['LOW'])}.",
            f"{data.total('in_progress'):,} in progress; "
            f"{total - data.total('in_progress'):,} OPEN.",
        ],
        19,
    )
    write_text(
        posture,
        25,
        trend_statement(current, before, data.total("in_progress"), data.start.strftime("%B")),
        16,
    )
    bullets = shape(posture, 25).text_frame.paragraphs[1:]
    reference = bullets[0]._p.pPr if bullets else None
    for paragraph in bullets[1:]:
        if reference is not None and paragraph._p.pPr is not None:
            for attr in ("marL", "indent"):
                if reference.get(attr) is not None:
                    paragraph._p.pPr.set(attr, reference.get(attr))
    fill_tables(deck, data, "resolved_groups", (4, 5, 6))
    fill_tables(deck, data, "in_progress_groups", (7, 8))
    write_text(deck.slides[9], 8, takeaways, 20)
    provenance = {
        "snapshot_sha256": source_hash,
        "template_sha256": digest(template),
        "scope": payload["scope"],
        "period": payload["period"],
        "prior": prior,
        "limitations": payload["limitations"],
        "assignment_scope": payload.get("assignment_scope", "STATUS"),
        "resolved_attribution": (
            payload["team_resolved"]["basis"]
            if team
            else "UNVERIFIED: tenant closure statistics only; current assignees and "
            "shared rule owners do not establish contribution before resolution."
        ),
        "rule_assignment_evidence": payload.get("rule_assignment_evidence", []),
        "rule_assignees": data.rule_assignees()
        if payload.get("assignment_scope") == "RULE"
        else {},
        "assignment_gaps": payload.get("assignment_gaps", []),
        "owner_status": payload.get("owner_status", []),
        "deferred_in_progress": payload.get("deferred_in_progress", []),
        "final_approval": payload["period"].get("final_approval"),
        "rule_status": payload.get("rule_status", []),
        "semantics": "Counts are current statuses, not a complete event history. "
        "Resolved timestamps exclude rejected findings. "
        "Samples do not prove changes for an entire group.",
    }
    for index, slide in enumerate(deck.slides):
        notes = dict(provenance)
        if index in TABLE_CAPACITY:
            field = (
                ("team_resolved" if team else "resolved_groups")
                if index <= 6
                else "in_progress_groups"
            )
            notes["groups"] = data.ranked(field)
        slide.notes_slide.notes_text_frame.text = json.dumps(notes, ensure_ascii=True, indent=2)
    deck.save(destination)


def _cell(value: str) -> str:
    """Slide cell text with whitespace normalised, so layout-only differences are not edits."""
    return " ".join(value.replace("\u00a0", " ").split())


def import_slide_edits(snapshot: Path, deck_path: Path, editor: str) -> list[str]:
    """Record text the team edited on a generated deck as snapshot statements.

    Each resolved / in-progress table row is matched to its rule group by issue name and
    compared with what the snapshot would render. Edited Closure evidence / Risk reduction
    cells become ``closure_wording``; edited Next step / Current status cells become
    ``rule_status``. Issue names, owners, counts and severities always come from Wiz and are
    never imported.
    """
    snapshot, deck_path = snapshot.absolute(), deck_path.absolute()
    data = load_snapshot(snapshot)
    payload = data.payload
    credited = "team_resolved" in payload
    groups = {
        True: {g["name"].strip(): g for g in payload["team_resolved"]["groups"]}
        if credited
        else {},
        False: {g["name"].strip(): g for g in payload["in_progress_groups"]["groups"]},
    }
    rule_names = (
        data.rule_assignees()
        if payload.get("assignment_scope") == "RULE" and not credited
        else None
    )
    now = datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")
    source = f"{editor} slide edit imported from {deck_path.name}"
    wording = {e["id"]: e for e in payload.setdefault("closure_wording", [])}
    statuses = {e["id"]: e for e in payload.setdefault("rule_status", [])}
    changes: list[str] = []

    def record(store: dict, entries: list, rule_id: str, key: str, value: str, label: str) -> None:
        entry = store.get(rule_id)
        if entry is None:
            entry = store[rule_id] = {"id": rule_id}
            entries.append(entry)
        entry.update({key: value, "source": source, "as_of": now})
        changes.append(f"{label}: {value}")

    for slide in Presentation(deck_path).slides:
        for item in slide.shapes:
            if not item.has_table or _cell(item.table.cell(0, 0).text) not in (
                "Resolved issue",
                "In-progress issue",
            ):
                continue
            resolved = _cell(item.table.cell(0, 0).text) == "Resolved issue"
            for row in list(item.table.rows)[1:]:
                cells = [c.text for c in row.cells]
                group = groups[resolved].get(cells[0].split("\n")[0].strip())
                if group is None or (resolved and not credited):
                    continue
                names = rule_names.get(group["id"], []) if rule_names is not None else None
                expected = table_row(
                    group,
                    resolved,
                    names,
                    credited,
                    None if resolved else data.owner_status(),
                    None if resolved else data.rule_status(),
                    data.closure_wording() if resolved else None,
                    None if resolved else data.rule_next_step(),
                )
                label = f"{group['name'].strip()[:60]} ({group['id']})"
                for column, key in (
                    (2, "changed" if resolved else "next_step"),
                    (3, "risk" if resolved else "status"),
                ):
                    edited = _cell(cells[column])
                    if edited and edited != _cell(expected[column]):
                        record(
                            wording if resolved else statuses,
                            payload["closure_wording"] if resolved else payload["rule_status"],
                            group["id"],
                            key,
                            edited,
                            f"{label} {key}",
                        )
    if changes:
        if not payload["closure_wording"]:
            del payload["closure_wording"]
        if not payload["rule_status"]:
            del payload["rule_status"]
        staged = snapshot.with_suffix(".import.tmp")
        staged.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        load_snapshot(staged)  # validate before replacing the reviewed snapshot
        backup = snapshot.with_suffix(f".before-import-{datetime.now(UTC):%Y%m%dT%H%M%SZ}.json")
        shutil.copy2(snapshot, backup)
        staged.replace(snapshot)
    return changes


def build_governance(
    snapshot: Path, template: Path, output: Path | None = None, refresh: bool = False
) -> Path:
    snapshot, template = snapshot.absolute(), template.absolute()
    expected = {snapshot: digest(snapshot), template: digest(template)}
    require(all(expected.values()), "Snapshot and template must exist.")
    data = load_snapshot(snapshot)
    output = (
        output
        or template.parent
        / (
            f"{data.end.strftime('%B%Y')}_{active_profile().file_prefix}_Monthly_Cloud_Security_Governance.pptx"
        )
    ).absolute()
    require(output.suffix.lower() == ".pptx", "Output must be a .pptx file.")
    companion = output.with_suffix(".governance.json")
    require(
        len({p.resolve() for p in (snapshot, template, output, companion)}) == 4,
        "Source snapshot, approved template and output pair must be distinct files.",
    )
    require(output.parent.is_dir(), f"Output directory does not exist: {output.parent}")
    month = f"governance-{data.end:%Y-%m}"
    with monthly_lock(output.parent, month):
        check_recovery(output.parent, month)
        expected.update({output: digest(output), companion: digest(companion)})
        require(
            refresh or (expected[output] is None and expected[companion] is None),
            "Output exists. Use --refresh to back up and replace this month's deck and snapshot.",
        )
        with tempfile.TemporaryDirectory(
            prefix=".wiz-monthly-governance-", dir=output.parent
        ) as folder:
            staged_deck, staged_json = Path(folder) / output.name, Path(folder) / companion.name
            render_deck(data, template, staged_deck, expected[snapshot])
            validate_template(Presentation(staged_deck))
            shutil.copy2(snapshot, staged_json)
            check_unchanged(expected)
            publish_pair((staged_deck, staged_json), (output, companion), expected, month)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--snapshot", required=True, type=Path, help="Reviewed local Wiz governance JSON."
    )
    parser.add_argument("--template", type=Path, help="Approved prior ten-slide governance deck.")
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--refresh", action="store_true", help="Back up and replace the output pair."
    )
    parser.add_argument(
        "--actions-csv",
        type=Path,
        help="Write Wiz assignee actions (issues missing an assignee) to this CSV.",
    )
    parser.add_argument(
        "--import-edits",
        type=Path,
        metavar="DECK",
        help="Record team-edited table text from a generated deck into --snapshot, then exit.",
    )
    parser.add_argument(
        "--editor",
        default="Team",
        help="Who made the slide edits (recorded as the source with --import-edits).",
    )
    add_profile_argument(parser)
    args = parser.parse_args(argv)
    ensure_profile(args.profile)
    try:
        if args.import_edits:
            changes = import_slide_edits(args.snapshot, args.import_edits, args.editor)
            print(f"Imported {len(changes)} slide edit(s) into {args.snapshot}")
            for line in changes:
                print(f"  - {line}")
            return 0
        require(args.template is not None, "--template is required to build the deck.")
        result = build_governance(args.snapshot, args.template, args.output, args.refresh)
        print(f"Created {result}\nEvidence: {result.with_suffix('.governance.json')}")
        gaps = load_snapshot(args.snapshot.absolute()).payload.get("assignment_gaps", [])
        print(f"Wiz assignee actions: {len(gaps)}")
        for line in action_lines(gaps):
            print(f"  - {line}")
        if args.actions_csv:
            write_actions_csv(gaps, args.actions_csv)
            print(f"Assignee actions CSV: {args.actions_csv}")
        return 0
    except (DashboardError, OSError, ValueError) as exc:
        parser.exit(1, f"Governance generation failed: {exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
