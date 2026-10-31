"""Public, placeholder-only structural templates for Monthly Governance Reporting.

Structure only: generic tabs, slides and ``{{placeholder}}`` fields. No client
data, branding or production layout. Outputs are generated on demand into a
caller-chosen directory; private templates and pinned hashes stay private.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

TEMPLATE_VERSION = 1
CREATOR = "wizexpert-public-template"
FIXED_TIME = datetime(2000, 1, 1)
# 13.333 x 7.5 in, as EMU (identical to pptx.util.Inches(13.333) / Inches(7.5)).
DECK_WIDTH_EMU = int(13.333 * 914400)
DECK_HEIGHT_EMU = int(7.5 * 914400)

PLACEHOLDER = re.compile(r"\{\{\s*\.?[a-z][a-z0-9_]*\s*\}\}")
REPEAT_MARKER = "<!-- repeat:{name}_rows -->"

WORKBOOK_FILE = "monthly-governance-workbook.template.xlsx"
OPERATIONAL_FILE = "monthly-governance-slide-deck.template.pptx"
EXECUTIVE_FILE = "executive-dashboard.template.pptx"

WORKBOOK_SLOT = "monthly-governance-workbook"
OPERATIONAL_SLOT = "monthly-governance-slide-deck"
EXECUTIVE_SLOT = "executive-dashboard"

COLUMNS = 8
WORKBOOK_TABS = [(f"Tab {n} Placeholder", f"tab_{n}") for n in range(1, 7)]
SCALARS = [
    ("report_title", "Report title"),
    ("snapshot_period", "Snapshot period"),
    ("baseline_period", "Baseline period"),
    ("snapshot_cutoff", "Snapshot cutoff"),
    ("draft_status", "Draft status"),
    ("scope_statement", "Scope statement"),
]
ALLOWED_LABELS = {label for _, label in SCALARS}

OPERATIONAL_SLIDES = [
    {
        "title": f"{{{{slide_{n}_title}}}}",
        "fields": [f"{{{{slide_{n}_body_1}}}}", f"{{{{slide_{n}_body_2}}}}"],
    }
    for n in range(1, 11)
]
EXECUTIVE_SLIDES = [
    {
        "title": f"{{{{exec_{n}_title}}}}",
        "fields": [f"{{{{exec_{n}_metric_1}}}}", f"{{{{exec_{n}_metric_2}}}}"],
    }
    for n in range(1, 4)
]
FOOTER = "{{footer_text}}"
OPERATIONAL_SEVERITIES = ["Critical", "High", "Medium", "Low"]
EXECUTIVE_SEVERITIES = ["Critical", "High"]

_CLIENT = "care" + "first"
_PRIVACY = [
    ("account-like number", re.compile(r"\d{6,}")),
    ("CVE identifier", re.compile(r"CVE-\d", re.I)),
    ("workload/case id", re.compile(r"wc-id-\d+", re.I)),
    ("email address", re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")),
    ("private path", re.compile(r"[A-Za-z]:[\\/]|/(?:Users|home)/|\\\\[^\s\\]+\\")),
    ("URL", re.compile(r"https?://|www\.", re.I)),
    ("restricted name", re.compile(_CLIENT, re.I)),
]


def _slide_placeholders(slides: list[dict]) -> list[list[str]]:
    return [[s["title"], *s["fields"], FOOTER] for s in slides]


def _tab_placeholders(key: str) -> list[str]:
    cols = [f"{{{{col{c}_label}}}}" for c in range(1, COLUMNS + 1)]
    cols += [f"{{{{.col{c}}}}}" for c in range(1, COLUMNS + 1)]
    return [*cols, f"{{{{{key}_empty_message}}}}"]


def schema() -> dict:
    return {
        "template_version": TEMPLATE_VERSION,
        WORKBOOK_SLOT: {
            "file": WORKBOOK_FILE,
            "scalars": [name for name, _ in SCALARS],
            "tabs": [{"sheet": s, "key": k, "columns": COLUMNS} for s, k in WORKBOOK_TABS],
        },
        OPERATIONAL_SLOT: {
            "file": OPERATIONAL_FILE,
            "slides": [dict(s, fields=list(s["fields"])) for s in OPERATIONAL_SLIDES],
            "severities": list(OPERATIONAL_SEVERITIES),
            "footer": FOOTER,
        },
        EXECUTIVE_SLOT: {
            "file": EXECUTIVE_FILE,
            "slides": [dict(s, fields=list(s["fields"])) for s in EXECUTIVE_SLIDES],
            "severities": list(EXECUTIVE_SEVERITIES),
            "footer": FOOTER,
        },
    }


# ---------------------------------------------------------------- builders


def build_workbook(out_dir: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for index, (sheet, key) in enumerate(WORKBOOK_TABS):
        ws = wb.create_sheet(sheet)
        row = 1
        if index == 0:
            for name, label in SCALARS:
                ws.cell(row, 1, label)
                ws.cell(row, 2, f"{{{{{name}}}}}")
                row += 1
        for c in range(1, COLUMNS + 1):
            ws.cell(row, c, f"{{{{col{c}_label}}}}")
        row += 1
        for c in range(1, COLUMNS + 1):
            ws.cell(row, c, f"{{{{.col{c}}}}}")
        ws.cell(row, COLUMNS + 1, REPEAT_MARKER.format(name=key))
        row += 1
        ws.cell(row, 1, f"{{{{{key}_empty_message}}}}")
    wb.properties.creator = CREATOR
    wb.properties.lastModifiedBy = CREATOR
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / WORKBOOK_FILE
    wb.save(path)
    return path


def _build_deck(out_dir: Path, slides: list[dict], file_name: str) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = DECK_WIDTH_EMU
    prs.slide_height = DECK_HEIGHT_EMU
    blank = prs.slide_layouts[6]
    for spec in slides:
        slide = prs.slides.add_slide(blank)
        boxes = [(spec["title"], 0.5, 0.4, 24)]
        for i, field in enumerate(spec["fields"]):
            boxes.append((field, 0.5, 1.6 + i * 1.8, 16))
        boxes.append((FOOTER, 0.5, 6.8, 10))
        for text, left, top, size in boxes:
            box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(12.3), Inches(1.2))
            box.text_frame.text = text
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    props = prs.core_properties
    props.author = CREATOR
    props.last_modified_by = CREATOR
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / file_name
    prs.save(path)
    return path


def build_operational_deck(out_dir: Path) -> Path:
    return _build_deck(out_dir, OPERATIONAL_SLIDES, OPERATIONAL_FILE)


def build_executive_dashboard(out_dir: Path) -> Path:
    return _build_deck(out_dir, EXECUTIVE_SLIDES, EXECUTIVE_FILE)


BUILDERS: dict[str, Callable[[Path], Path]] = {
    WORKBOOK_SLOT: build_workbook,
    OPERATIONAL_SLOT: build_operational_deck,
    EXECUTIVE_SLOT: build_executive_dashboard,
}


def build_all(out_dir: Path) -> dict[str, Path]:
    return {slot: builder(Path(out_dir)) for slot, builder in BUILDERS.items()}


# -------------------------------------------------------------- validators

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_MEMBERS = 500
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_MEMBER_RATIO = 100
_BLOCKED_PART_MARKERS = ("vbaproject", "/embeddings/", "oleobject", "/activex/", "/media/")
_BLOCK = re.compile(r"\{\{.*?\}\}", re.S)
_TAGS = re.compile(r"<[^>]+>")
_NAME_ATTR = re.compile(r'name\s*=\s*"([^"]*)"')
_EXTERNAL = re.compile(r"""TargetMode\s*=\s*["']External["']""", re.I)
_CORE_WB_PROPS = (
    "creator", "lastModifiedBy", "title", "subject", "description", "keywords",
    "category", "identifier", "language", "version", "contentStatus", "revision",
)
_CORE_DECK_PROPS = (
    "author", "last_modified_by", "title", "subject", "comments", "keywords",
    "category", "content_status", "identifier", "language", "version", "revision",
)
_HEADER_FOOTERS = (
    "oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter",
)


def _workbook_placeholders() -> set[str]:
    names = {f"{{{{{n}}}}}" for n, _ in SCALARS}
    for _, key in WORKBOOK_TABS:
        names.update(_tab_placeholders(key))
    return names


def _deck_placeholders(slides: list[dict]) -> set[str]:
    return {item for row in _slide_placeholders(slides) for item in row}


def _privacy_problems(where: str, text: str) -> list[str]:
    # Scan the RAW text: anything hidden inside {{...}} is still private content.
    return [f"{where}: contains {name}" for name, rx in _PRIVACY if rx.search(text)]


def _placeholder_problems(where: str, text: str, allowed: set[str]) -> list[str]:
    # Only exact schema placeholder names are allowed; never echo the offending text.
    if any(m.group(0) not in allowed for m in _BLOCK.finditer(text)):
        return [f"{where}: unknown placeholder name"]
    return []


def _dedupe(items: list[str]) -> list[str]:
    return list(dict.fromkeys(items))


def _package_problems(path: Path) -> tuple[list[str], list[str]]:
    """Bounds and safety checks on the zip before any parser loads it.

    Returns (blocking problems, non-blocking findings from package metadata).
    Messages never include file system paths or member names.
    """
    import html
    import zipfile

    try:
        size = Path(path).stat().st_size
    except OSError:
        return ["cannot read file"], []
    if size > MAX_FILE_BYTES:
        return [f"file larger than {MAX_FILE_BYTES} bytes"], []
    try:
        zf = zipfile.ZipFile(path)
    except Exception:
        return ["not a valid Office package (zip)"], []
    problems: list[str] = []
    notes: list[str] = []
    with zf:
        infos = zf.infolist()
        if len(infos) > MAX_MEMBERS:
            return [f"package has more than {MAX_MEMBERS} members"], []
        if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
            return ["package uncompressed size is too large"], []
        for info in infos:
            name = info.filename.replace("\\", "/")
            lower = "/" + name.lower()
            if name.startswith("/") or re.match(r"[A-Za-z]:", name) or ".." in name.split("/"):
                problems.append("package member has an unsafe path")
            if info.flag_bits & 0x1:
                problems.append("package contains encrypted members")
            if any(marker in lower for marker in _BLOCKED_PART_MARKERS):
                problems.append("package contains macro, embedded-object, picture or media parts")
            if info.file_size > 1024 and (
                info.compress_size == 0 or info.file_size / info.compress_size > MAX_MEMBER_RATIO
            ):
                problems.append("package compression ratio is too high")
        if problems:
            return _dedupe(problems), []
        for info in infos:
            lower = "/" + info.filename.replace("\\", "/").lower()
            is_rels = lower.endswith(".rels")
            is_types = lower == "/[content_types].xml"
            is_props = lower.startswith("/docprops/") and lower.endswith(".xml")
            if not (is_rels or is_types or is_props):
                continue
            try:
                with zf.open(info) as fh:
                    data = fh.read(MAX_UNCOMPRESSED_BYTES + 1).decode("utf-8", "replace")
            except Exception:
                problems.append("package part is unreadable")
                continue
            if is_rels and _EXTERNAL.search(data):
                problems.append("package has external relationships")
            if is_types and re.search(r"macroEnabled|vbaProject", data, re.I):
                problems.append("package declares macro content")
            if is_props:
                names = " ".join(_NAME_ATTR.findall(data))
                text = html.unescape(_TAGS.sub(" ", data)) + " " + names
                notes += _privacy_problems("document properties", text)
    return _dedupe(problems), notes


def _check_workbook(path: Path) -> list[str]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path)
    except Exception as exc:  # unreadable or not an xlsx
        return [f"cannot open workbook: {type(exc).__name__}"]
    problems: list[str] = []
    allowed = _workbook_placeholders()
    expected = [s for s, _ in WORKBOOK_TABS]
    if wb.sheetnames != expected:
        problems.append(f"sheet order/names differ from {expected}")
    if wb.properties.creator != CREATOR or wb.properties.lastModifiedBy != CREATOR:
        problems.append("workbook creator metadata is not the generic template creator")
    for prop in _CORE_WB_PROPS:
        value = str(getattr(wb.properties, prop, "") or "")
        problems += _privacy_problems(f"property {prop}", value)
    if getattr(wb, "_external_links", None):
        problems.append("workbook contains external links")
    for name, defined in wb.defined_names.items():
        problems += _privacy_problems("defined name", f"{name} {defined.attr_text or ''}")
    for ws in wb:
        problems += _privacy_problems("sheet name", ws.title)
        for name, defined in ws.defined_names.items():
            text = f"{name} {defined.attr_text or ''}"
            problems += _privacy_problems(f"{ws.title}: defined name", text)
        for dv in ws.data_validations.dataValidation:
            for part in (dv.formula1, dv.formula2, dv.prompt, dv.promptTitle, dv.error,
                         dv.errorTitle, str(dv.sqref)):
                problems += _privacy_problems(f"{ws.title}: data validation", str(part or ""))
        if getattr(ws, "_images", None) or getattr(ws, "_charts", None):
            problems.append(f"{ws.title}: contains pictures or charts")
        if ws._hyperlinks:
            problems.append(f"{ws.title}: contains a hyperlink")
        for attr in _HEADER_FOOTERS:
            hf = getattr(ws, attr)
            for side in ("left", "center", "right"):
                text = getattr(hf, side).text or ""
                if text:
                    problems += _privacy_problems(f"{ws.title}: header/footer", text)
                    problems += _placeholder_problems(f"{ws.title}: header/footer", text, allowed)
                    if PLACEHOLDER.sub("", text).strip():
                        problems.append(f"{ws.title}: header/footer non-placeholder text")
        for c in ws._cells.values():
            if c.comment is not None:
                problems.append(f"{ws.title}!{c.coordinate}: contains a comment")
                text = f"{c.comment.text} {c.comment.author or ''}"
                problems += _privacy_problems(f"{ws.title}!{c.coordinate} comment", text)
            if c.hyperlink is not None:
                problems.append(f"{ws.title}!{c.coordinate}: contains a hyperlink")
                link = c.hyperlink
                text = f"{link.target or ''} {link.location or ''} {link.display or ''}"
                problems += _privacy_problems(f"{ws.title}!{c.coordinate} hyperlink", text)
    for index, (sheet, key) in enumerate(WORKBOOK_TABS):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
        text = "\n".join(str(c.value) for c in cells)
        required = _tab_placeholders(key)
        if index == 0:
            required += [f"{{{{{n}}}}}" for n, _ in SCALARS]
        for item in required:
            if item not in text:
                problems.append(f"{sheet}: missing placeholder {item}")
        marker = REPEAT_MARKER.format(name=key)
        if marker not in text:
            problems.append(f"{sheet}: missing repeat marker {marker}")
        for c in cells:
            value = str(c.value)
            where = f"{sheet}!{c.coordinate}"
            rest = PLACEHOLDER.sub("", value.replace(marker, "")).strip()
            if rest and rest not in ALLOWED_LABELS:
                problems.append(f"{where}: non-placeholder text")
            problems += _placeholder_problems(where, value, allowed)
            problems += _privacy_problems(where, value)
    return problems


def validate_workbook(path: Path) -> list[str]:
    blocking, notes = _package_problems(path)
    if blocking:
        return blocking
    try:
        return _dedupe(_check_workbook(path) + notes)
    except Exception as exc:  # fail closed, no path leakage
        return [f"cannot validate workbook: {type(exc).__name__}"]


def _walk_shapes(shapes, where: str, problems: list[str], texts: list[str]) -> None:
    for shape in shapes:
        kind = shape._element.tag.rsplit("}", 1)[-1]
        if kind == "pic":
            problems.append(f"{where}: contains a picture or media")
        elif kind == "grpSp":
            problems.append(f"{where}: contains a group shape")
            _walk_shapes(shape.shapes, where, problems, texts)
        elif kind == "graphicFrame":
            if getattr(shape, "has_table", False):
                problems.append(f"{where}: contains a table")
                for row in shape.table.rows:
                    texts.extend(cell.text_frame.text for cell in row.cells)
            elif getattr(shape, "has_chart", False):
                problems.append(f"{where}: contains a chart")
            else:
                problems.append(f"{where}: contains an embedded object")
        elif kind == "contentPart":
            problems.append(f"{where}: contains an embedded object")
        if getattr(shape, "has_text_frame", False):
            texts.append(shape.text_frame.text)


def _scan_part(shapes, root, where: str, problems: list[str], texts: list[str]) -> None:
    _walk_shapes(shapes, where, problems, texts)
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        local = el.tag.rsplit("}", 1)[-1]
        if local in ("hlinkClick", "hlinkHover"):
            problems.append(f"{where}: contains a hyperlink")
        elif local in ("videoFile", "audioFile", "wavAudioFile", "quickTimeFile", "oleObj"):
            problems.append(f"{where}: contains media or an embedded object")
        elif local == "t" and el.text:
            problems += _privacy_problems(where, el.text)
        elif local == "cNvPr":
            for attr in ("name", "descr", "title"):
                problems += _privacy_problems(f"{where} shape {attr}", el.get(attr) or "")


def _check_deck(path: Path, slides: list[dict]) -> list[str]:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return [f"cannot open deck: {type(exc).__name__}"]
    problems: list[str] = []
    allowed = _deck_placeholders(slides)
    if len(prs.slides) != len(slides):
        problems.append(f"expected {len(slides)} slides, found {len(prs.slides)}")
    if prs.slide_width != DECK_WIDTH_EMU or prs.slide_height != DECK_HEIGHT_EMU:
        problems.append("slide size is not 13.333 x 7.5 in")
    props = prs.core_properties
    if props.author != CREATOR or props.last_modified_by != CREATOR:
        problems.append("deck author metadata is not the generic template creator")
    for prop in _CORE_DECK_PROPS:
        problems += _privacy_problems(f"property {prop}", str(getattr(props, prop, "") or ""))
    expected = _slide_placeholders(slides)
    for n, slide in enumerate(prs.slides, start=1):
        where = f"slide {n}"
        texts: list[str] = []
        _scan_part(slide.shapes, slide._element, where, problems, texts)
        joined = "\n".join(texts)
        if n <= len(expected):
            for item in expected[n - 1]:
                if item not in joined:
                    problems.append(f"{where}: missing placeholder {item}")
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            _scan_part(notes_slide.shapes, notes_slide._element, f"{where} notes", problems, texts)
        for text in texts:
            problems += _placeholder_problems(where, text, allowed)
            problems += _privacy_problems(where, text)
            if PLACEHOLDER.sub("", text).strip():
                problems.append(f"{where}: non-placeholder text")
    for master in prs.slide_masters:
        _scan_part(master.shapes, master._element, "slide master", problems, [])
        for layout in master.slide_layouts:
            _scan_part(layout.shapes, layout._element, "slide layout", problems, [])
    return problems


def _validate_deck(path: Path, slides: list[dict]) -> list[str]:
    blocking, notes = _package_problems(path)
    if blocking:
        return blocking
    try:
        return _dedupe(_check_deck(path, slides) + notes)
    except Exception as exc:  # fail closed, no path leakage
        return [f"cannot validate deck: {type(exc).__name__}"]


def validate_operational_deck(path: Path) -> list[str]:
    return _validate_deck(path, OPERATIONAL_SLIDES)


def validate_executive_dashboard(path: Path) -> list[str]:
    return _validate_deck(path, EXECUTIVE_SLIDES)


VALIDATORS: dict[str, Callable[[Path], list[str]]] = {
    WORKBOOK_SLOT: validate_workbook,
    OPERATIONAL_SLOT: validate_operational_deck,
    EXECUTIVE_SLOT: validate_executive_dashboard,
}


def validate_all(paths: dict[str, Path]) -> dict[str, list[str]]:
    return {slot: VALIDATORS[slot](Path(path)) for slot, path in paths.items()}
