"""Validated Critical/High workbook data and approved monthly trend history."""

from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from wizexpert.governance.profile import active as active_profile
from wizexpert.governance.profile import ordinal

SEVERITIES = ("CRITICAL", "HIGH")


def baseline_day() -> date:
    return active_profile().baseline


def baseline_label() -> str:
    day = baseline_day()
    return f"{day:%B} {day.day}, {day.year}"


def baseline_open(data: WorkbookData | None = None) -> dict[str, int]:
    """Open CRITICAL/HIGH counts at the fixed baseline.

    Uses the profile's counts, else the workbook's baseline column.
    """
    configured = active_profile().baseline_open
    if configured:
        return dict(configured)
    day = baseline_day()
    require(
        data is not None and day in data.snapshots,
        "Baseline open counts are not in the client profile and the workbook has no "
        f"{baseline_label()} "
        "column. Add them with 'wizexpert governance setup'.",
    )
    return {severity: data.snapshots[day][severity].open for severity in SEVERITIES}


class DashboardError(ValueError):
    """An input cannot safely be used to generate the approved dashboard."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise DashboardError(message)


def count(value: object, context: str) -> int:
    require(
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(value)
        and value >= 0
        and int(value) == value,
        f"{context}: expected a non-negative whole number, got {value!r}.",
    )
    return int(value)


def percent(value: object, context: str) -> float:
    try:
        if isinstance(value, str) and value.endswith("%"):
            result = float(value[:-1])
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            result = float(value) * 100
        else:
            raise ValueError
    except ValueError as exc:
        raise DashboardError(f"{context}: invalid percentage {value!r}.") from exc
    require(math.isfinite(result), f"{context}: percentage must be finite.")
    return result


def parse_date(value: object, context: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    require(isinstance(value, str), f"{context}: missing or invalid date.")
    text = value.strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DashboardError(f"{context}: unsupported date {value!r}.")


@dataclass(frozen=True)
class Snapshot:
    total: int
    open: int
    resolved: int
    ignored: int

    @property
    def closed(self) -> int:
        return self.resolved + self.ignored

    @property
    def addressed_pct(self) -> float:
        return self.closed / self.total * 100 if self.total else 0.0


@dataclass(frozen=True)
class Finding:
    issue_id: str
    title: str
    severity: str
    status: str
    closed_on: date
    reference: str


@dataclass(frozen=True)
class WorkbookData:
    snapshots: dict[date, dict[str, Snapshot]]
    findings: tuple[Finding, ...]
    accepted_risk_supported: bool
    preliminary: bool = False

    @property
    def dates(self) -> list[date]:
        return sorted(self.snapshots)

    @property
    def latest(self) -> date:
        return self.dates[-1]

    @property
    def previous(self) -> date:
        return self.dates[-2]

    @property
    def reporting_date(self) -> date:
        if not self.preliminary:
            return self.latest
        year, month = divmod(self.previous.year * 12 + self.previous.month, 12)
        return date(year, month + 1, baseline_day().day)

    def open_at(self, day: date) -> int:
        return sum(value.open for value in self.snapshots[day].values())

    def closed_at(self, day: date) -> int:
        return sum(value.closed for value in self.snapshots[day].values())


def read_snapshots(ws, preliminary_latest: bool = False) -> dict[date, Snapshot]:
    tab = ws.title
    labels = [str(ws.cell(row, 1).value or "").lower() for row in range(5, 10)]
    for label, prefix in zip(
        labels, ("total", "open", "resolved", "ignored", "resolution %"), strict=False
    ):
        require(label.startswith(prefix), f"{tab}: unrecognized snapshot table layout.")
    snapshots: dict[date, Snapshot] = {}
    for col in range(2, ws.max_column + 1):
        raw = str(ws.cell(4, col).value or "").split("\n")[0].strip()
        if not re.fullmatch(r"[A-Za-z]{3} \d{1,2}, \d{4}", raw):
            require(
                raw.lower().startswith(("vs baseline", "monthly")),
                f"{tab} column {col}: missing or unsupported snapshot header {raw!r}.",
            )
            break
        day = parse_date(raw, f"{tab} row 4")
        following = str(ws.cell(4, col + 1).value or "").split("\n")[0].strip()
        latest_column = not re.fullmatch(r"[A-Za-z]{3} \d{1,2}, \d{4}", following)
        baseline = baseline_day()
        require(
            day >= baseline and (day.day == baseline.day or (preliminary_latest and latest_column)),
            f"{tab}: snapshots must be monthly {ordinal(baseline.day)} dates "
            f"from {baseline:%B %Y}.",
        )
        require(day not in snapshots, f"{tab}: duplicate snapshot {day}.")
        values = [
            count(ws.cell(row, col).value, f"{tab}!{ws.cell(row, col).coordinate}")
            for row in range(5, 9)
        ]
        snapshot = Snapshot(*values)
        require(
            snapshot.total == snapshot.open + snapshot.closed,
            f"{tab} {day}: total does not equal open + resolved + ignored.",
        )
        shown = percent(ws.cell(9, col).value, f"{tab} {day}")
        require(
            abs(shown - snapshot.addressed_pct) <= 0.051,
            f"{tab} {day}: resolution percentage disagrees with counts.",
        )
        snapshots[day] = snapshot
    require(len(snapshots) >= 2, f"{tab}: at least two workbook snapshots are required.")
    require(list(snapshots) == sorted(snapshots), f"{tab}: snapshots must be chronological.")
    return snapshots


def read_workbook(path: Path, allow_preliminary: bool = False) -> WorkbookData:
    require(path.is_file(), f"Workbook not found: {path}")
    by_severity: dict[str, dict[date, Snapshot]] = {}
    findings: list[Finding] = []
    seen_ids: set[str] = set()
    accepted = True
    wb = load_workbook(path, data_only=False)
    try:
        flags = [str(ws["A2"].value or "").startswith("PRELIMINARY |") for ws in wb]
        preliminary = any(flags)
        require(
            not preliminary or allow_preliminary,
            "Preliminary collection workbook cannot be used as a final executive snapshot.",
        )
        require(
            not preliminary or all(flags),
            "Every worksheet in a preliminary workbook must be labelled.",
        )
        for severity in SEVERITIES:
            tab = f"{severity} Issues"
            require(tab in wb.sheetnames, f"Missing required worksheet: {tab}")
            ws = wb[tab]
            accepted = accepted and "accepted risk" in str(ws.cell(8, 1).value or "").lower()
            snapshots = read_snapshots(ws, preliminary)
            by_severity[severity] = snapshots
            header = next(
                (
                    r
                    for r in range(10, min(ws.max_row, 100) + 1)
                    if ws.cell(r, 1).value == "Issue ID"
                ),
                None,
            )
            require(header is not None, f"{tab}: closed-issue detail header missing.")
            names = {
                str(ws.cell(header, c).value or "").strip(): c for c in range(1, ws.max_column + 1)
            }
            # Column positions are read by header name so workbooks from earlier layouts still load.
            c_title, c_sev, c_status, c_created, c_closed = (
                names.get(name, default)
                for name, default in (
                    ("Title", 2),
                    ("Severity", 3),
                    ("Status", 4),
                    ("Created At", 5),
                    ("Resolved Time", 6),
                )
            )
            last_col = get_column_letter(max(names.values()))
            for row in range(header + 1, ws.max_row + 1):
                issue_id = str(ws.cell(row, 1).value or "").strip()
                if not issue_id or issue_id.startswith("Total:"):
                    break
                title = ws.cell(row, c_title).value
                status = str(ws.cell(row, c_status).value or "").strip().upper()
                require(
                    isinstance(title, str) and bool(title.strip()),
                    f"{tab}!{get_column_letter(c_title)}{row}: missing finding title.",
                )
                require(
                    ws.cell(row, c_sev).value == severity, f"{tab} row {row}: severity mismatch."
                )
                require(
                    status in {"RESOLVED", "IGNORED", "REJECTED"},
                    f"{tab} row {row}: invalid closure status.",
                )
                require(
                    issue_id not in seen_ids, f"Duplicate issue ID in detail tables: {issue_id}"
                )
                seen_ids.add(issue_id)
                closed_on = parse_date(
                    ws.cell(row, c_closed).value, f"{tab}!{get_column_letter(c_closed)}{row}"
                )
                created_on = parse_date(
                    ws.cell(row, c_created).value, f"{tab}!{get_column_letter(c_created)}{row}"
                )
                require(created_on <= closed_on, f"{tab} row {row}: closure precedes creation.")
                findings.append(
                    Finding(
                        issue_id,
                        title.strip(),
                        severity,
                        status,
                        closed_on,
                        f"{tab}!A{row}:{last_col}{row}",
                    )
                )
        dates = list(by_severity["CRITICAL"])
        require(dates == list(by_severity["HIGH"]), "Critical and High snapshot dates differ.")
        latest = dates[-1]
        previous = dates[-2]
        next_year, next_month = divmod(previous.year * 12 + previous.month, 12)
        scheduled = date(next_year, next_month + 1, baseline_day().day)
        require(
            previous < latest <= scheduled if preliminary else latest == scheduled,
            "The latest workbook must include the immediately preceding month's snapshot.",
        )
        for severity in SEVERITIES:
            detail = [f for f in findings if f.severity == severity and f.closed_on <= latest]
            expected = by_severity[severity][latest]
            require(
                sum(f.status == "RESOLVED" for f in detail) == expected.resolved
                and sum(f.status != "RESOLVED" for f in detail) == expected.ignored,
                f"{severity}: cutoff-aligned detail counts do not match the latest snapshot.",
            )
        require("Summary" in wb.sheetnames, "Missing Summary worksheet.")
        for severity in SEVERITIES:
            summary = wb["Summary"]
            matches = [
                row
                for row in range(5, summary.max_row + 1)
                if summary.cell(row, 1).value == severity
            ]
            require(len(matches) == 1, f"Summary must contain exactly one {severity} row.")
            row = matches[0]
            values = [
                count(summary.cell(row, col).value, f"Summary row {row}") for col in range(2, 6)
            ]
            expected = by_severity[severity][latest]
            require(
                values == [expected.total, expected.open, expected.resolved, expected.ignored],
                f"Summary and {severity} latest snapshot disagree.",
            )
            require(
                abs(
                    percent(summary.cell(row, 6).value, f"Summary {severity}")
                    - expected.addressed_pct
                )
                <= 0.051,
                f"Summary {severity} percentage mismatch.",
            )
        return WorkbookData(
            {
                day: {severity: by_severity[severity][day] for severity in SEVERITIES}
                for day in dates
            },
            tuple(findings),
            accepted,
            preliminary,
        )
    finally:
        wb.close()


def history_from_json(text: str, allow_preliminary: bool = False) -> dict[date, int]:
    try:
        raw = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DashboardError(f"Invalid history JSON: {exc}") from exc
    require(
        isinstance(raw, dict) and raw.get("schema_version") == 1,
        "History must be a schema_version=1 object.",
    )
    preliminary = raw.get("preliminary") is True
    require(
        not preliminary or allow_preliminary,
        "A preliminary dashboard cannot be used as a template; use the latest final approved deck.",
    )
    require(
        raw.get("scope") == ["CRITICAL", "HIGH"], "History scope must be CRITICAL and HIGH only."
    )
    points = raw.get("points")
    require(isinstance(points, list) and bool(points), "History needs a non-empty points list.")
    result: dict[date, int] = {}
    baseline = baseline_day()
    for index, point in enumerate(points):
        require(isinstance(point, dict), "History points must be objects.")
        day = parse_date(point.get("date"), "History point")
        require(
            (day.day == baseline.day or (preliminary and index == len(points) - 1))
            and day >= baseline
            and day not in result,
            f"Invalid or duplicate monthly history point: {day}",
        )
        result[day] = count(point.get("open"), f"History {day}")
    require(
        baseline in result, f"Approved history must start with the {baseline_label()} baseline."
    )
    total = active_profile().baseline_total
    require(
        total is None or result[baseline] == total,
        f"Approved history must start with {baseline_label()} = {total}.",
    )
    return result


def combine_trend(data: WorkbookData, history: dict[date, int]) -> dict[date, int]:
    require(
        not history or max(history) <= data.latest,
        "Template history is newer than the workbook; supply the newest workbook.",
    )
    baseline = baseline_day()
    points = dict(history)
    for day in data.dates:
        points[day] = data.open_at(day)
    if active_profile().baseline_open or baseline not in history:
        points[baseline] = sum(baseline_open(data).values())
    else:
        points[baseline] = history[baseline]
    last_approved = data.previous if data.preliminary else data.latest
    months = (last_approved.year - baseline.year) * 12 + last_approved.month - baseline.month + 1
    require(data.latest > baseline, "Dashboard requires a post-baseline reporting month.")
    for offset in range(months):
        year, month = divmod(baseline.year * 12 + baseline.month - 1 + offset, 12)
        day = date(year, month + 1, baseline.day)
        require(
            day in points,
            f"Missing approved trend point for {day}; provide history, do not interpolate.",
        )
    return dict(sorted(points.items()))


@dataclass(frozen=True)
class Exposure:
    heading: str
    body: str
    evidence: tuple[Finding, ...]


def exposure_callouts(data: WorkbookData) -> list[Exposure]:
    resolved = [f for f in data.findings if f.status == "RESOLVED" and f.closed_on <= data.latest]
    rules = [
        (
            "Public data exposure closed",
            "Public access to sensitive bucket or blob data addressed.",
            lambda t: "public" in t and "sensitive" in t and ("bucket" in t or "blob" in t),
        ),
        (
            "Broad data access restricted",
            "Broad or external access to sensitive data addressed.",
            lambda t: "sensitive" in t and "access" in t and ("all users" in t or "external" in t),
        ),
        (
            "Threat-related findings resolved",
            "Malware or hacking-tool findings resolved on workloads.",
            lambda t: "infected" in t and ("malware" in t or "hacking tool" in t),
        ),
        (
            "Exposed databases addressed",
            "Public database exposure involving sensitive data addressed.",
            lambda t: "publicly exposed" in t and "database" in t and "sensitive" in t,
        ),
        (
            "External privilege paths reduced",
            "External access paths to high-privilege systems addressed.",
            lambda t: (
                ("third party" in t or "third-party" in t or "external" in t)
                and ("high privileges" in t or "admin" in t)
            ),
        ),
    ]
    options = []
    for priority, (heading, body, matches) in enumerate(rules):
        evidence = tuple(f for f in resolved if matches(f.title.lower()))
        if evidence:
            recent = max(f.closed_on for f in evidence) > data.previous
            options.append((not recent, priority, Exposure(heading, body, evidence)))
    options.sort(key=lambda item: item[:2])
    require(
        len(options) >= 2,
        "Fewer than two supported exposure examples; review detail data before generating.",
    )
    callouts = [item[2] for item in options[:3]]
    if len(callouts) == 2:
        outcome = "resolved or accepted" if data.accepted_risk_supported else "resolved or ignored"
        callouts.append(
            Exposure(
                "Risk outcomes tracked",
                f"{data.closed_at(data.latest)} Critical / High findings {outcome}.",
                (),
            )
        )
    return callouts
