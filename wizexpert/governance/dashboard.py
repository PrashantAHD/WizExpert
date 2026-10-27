"""Generate the client executive dashboard from an approved deck and workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import tempfile
from copy import deepcopy
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from wizexpert.governance.dashboard_charts import chart_fonts, line_chart, severity_chart
from wizexpert.governance.dashboard_data import (
    DashboardError,
    WorkbookData,
    baseline_day,
    baseline_label,
    baseline_open,
    combine_trend,
    exposure_callouts,
    history_from_json,
    read_workbook,
    require,
)
from wizexpert.governance.profile import active as active_profile
from wizexpert.governance.profile import add_profile_argument, ordinal
from wizexpert.governance.profile import ensure as ensure_profile

# Stored inside generated decks and read back next month; keep the original name for compatibility.
HISTORY_PREFIX = "wiz-reporting-history:"
PROGRESS_TEXT_IDS = {3, 4, 6, 7, 9, 10, 12, 13, 15, 16, 18, 19, 22, 23, 24, 25, 26, 27, 28, 30}
ACHIEVEMENT_TEXT_IDS = {
    3,
    4,
    6,
    7,
    9,
    10,
    12,
    13,
    15,
    16,
    20,
    21,
    22,
    23,
    26,
    27,
    28,
    29,
    30,
    33,
    34,
    37,
    38,
    41,
    42,
    43,
}


def shape(slide, shape_id):
    matches = [s for s in slide.shapes if s.shape_id == shape_id]
    require(len(matches) == 1, f"Approved template shape {shape_id} is missing or duplicated.")
    return matches[0]


def validate_template(deck) -> None:
    require(
        len(deck.slides) == 3, "Template must contain the approved cover plus two dashboard slides."
    )
    require(
        (deck.slide_width, deck.slide_height) == (12192000, 6858000),
        "Template must retain the approved 16:9 slide dimensions.",
    )
    for index, expected_ids in ((1, PROGRESS_TEXT_IDS), (2, ACHIEVEMENT_TEXT_IDS)):
        for sid in expected_ids:
            target = shape(deck.slides[index], sid)
            require(target.has_text_frame, f"Slide {index + 1} shape {sid} must be text.")
    for sid in (2, 5):
        require(
            shape(deck.slides[0], sid).has_text_frame, "Approved cover text placeholders missing."
        )
    client = active_profile().client_name
    require(
        shape(deck.slides[0], 5).text.startswith(f"{client} Executive Dashboard"),
        f"Approved {client} cover title not recognized.",
    )
    require(
        shape(deck.slides[1], 3).text == f"{client} Remediation Progress",
        "Progress slide title not recognized.",
    )
    require(
        shape(deck.slides[2], 3).text == "Risk Reduction Achievements",
        "Achievement slide title not recognized.",
    )
    for index, sid in ((1, 31), (2, 44)):
        require(
            shape(deck.slides[index], sid).shape_type == MSO_SHAPE_TYPE.PICTURE,
            f"Slide {index + 1}: approved image chart missing.",
        )


def template_history(
    deck,
    history_path: Path | None,
    template_path: Path | None = None,
    allow_preliminary: bool = False,
) -> dict[date, int]:
    picture = shape(deck.slides[1], 31)
    description = picture._element.xpath(".//p:cNvPr")[0].get("descr", "")
    if history_path is not None:
        require(history_path.is_file(), f"History file not found: {history_path}")
        text = history_path.read_text(encoding="utf-8")
        points = history_from_json(text)
        require(
            template_path is not None, "A template path is required to verify external history."
        )
        expected_hash = json.loads(text).get("template_sha256")
        require(
            expected_hash == hashlib.sha256(template_path.read_bytes()).hexdigest(),
            "History template_sha256 does not match the supplied template. "
            "Verify its chart before using history.",
        )
        return points
    require(
        description.startswith(HISTORY_PREFIX),
        "This template has a raster chart without structured history. Supply --history"
        " with manually verified monthly chart points; "
        "the generator will not guess or OCR numbers.",
    )
    return history_from_json(description.removeprefix(HISTORY_PREFIX), allow_preliminary)


def replace_text(slide, sid: int, paragraphs: list[str]) -> None:
    target = shape(slide, sid)
    frame = target.text_frame
    originals = list(frame.paragraphs)
    require(bool(originals), f"Shape {sid} has no template text paragraph.")
    styles = []
    for paragraph in originals:
        ppr = paragraph._p.find(qn("a:pPr"))
        first_run = next(iter(paragraph.runs), None)
        rpr = first_run._r.find(qn("a:rPr")) if first_run is not None else None
        styles.append((deepcopy(ppr), deepcopy(rpr)))
    frame.clear()
    for index, value in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        ppr, rpr = styles[min(index, len(styles) - 1)]
        existing = paragraph._p.find(qn("a:pPr"))
        if existing is not None:
            paragraph._p.remove(existing)
        if ppr is not None:
            paragraph._p.insert(0, deepcopy(ppr))
        paragraph.text = value.replace("\n", "\v")
        for run in paragraph.runs:
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))


def replace_picture(slide, sid, stream, description):
    picture = shape(slide, sid)
    _, relationship = slide.part.get_or_add_image_part(stream)
    picture._element.xpath(".//a:blip")[0].set(qn("r:embed"), relationship)
    picture._element.xpath(".//p:cNvPr")[0].set("descr", description)


def pct(value: float) -> str:
    return f"{value:.1f}".rstrip("0").rstrip(".") + "%"


def reduction(before: int, after: int) -> str:
    return pct((before - after) / before * 100) if before else "N/A"


def change_description(before: int, after: int) -> str:
    if before == after:
        return "no change"
    if before == 0:
        return "increase from zero"
    return (
        f"{pct(abs(before - after) / before * 100)} {'reduction' if after < before else 'increase'}"
    )


def full_date(day: date) -> str:
    return f"{day.strftime('%B')} {day.day}, {day.year}"


def date_ordinal(day: date) -> str:
    suffix = "th" if 11 <= day.day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day.day % 10, "th")
    return f"{day.strftime('%B')} {day.day}{suffix}, {day.year}"


def trend_language(data: WorkbookData) -> tuple[str, str, str]:
    before, after = data.open_at(data.previous), data.open_at(data.latest)
    if after < before:
        subtitle = "Open findings are declining month-over-month"
        chart = "Open findings declined again this month."
        takeaway = f"Open backlog fell from {before} to {after} this month."
    elif after > before:
        subtitle = "Open findings increased this month"
        chart = "Open findings increased; further action is needed."
        takeaway = f"Open backlog rose from {before} to {after}; further risk reduction is needed."
    else:
        subtitle = "Open findings are unchanged month-over-month"
        chart = "Open findings remained unchanged this month."
        takeaway = f"Open backlog remained at {after}; no further backlog reduction this month."
    critical = data.snapshots[data.latest]["CRITICAL"]
    if critical.open == 0:
        takeaway += " Critical findings are resolved or accepted."
    else:
        takeaway += f" {critical.open} Critical findings remain open."
    return subtitle, chart, takeaway


def build_deck(
    workbook: Path,
    template: Path,
    output: Path | None = None,
    history_path: Path | None = None,
    font_dir: Path | None = None,
    *,
    allow_preliminary: bool = False,
) -> Path:
    workbook, template = workbook.resolve(), template.resolve()
    require(
        template.is_file(),
        f"Approved template not found: {template}. A blank deck will not be generated.",
    )
    require(template.suffix.lower() == ".pptx", "Template must be a .pptx file.")
    data = read_workbook(workbook, allow_preliminary)
    deck = Presentation(template)
    validate_template(deck)
    history = template_history(deck, history_path, template)
    points = combine_trend(data, history)
    profile = active_profile()
    baseline = baseline_day()
    base_counts = baseline_open(data)
    base_total = points[baseline]
    base_month, base_short = baseline.strftime("%B"), baseline.strftime("%b")
    fonts = chart_fonts(font_dir or Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts")
    suffix = "_PRELIMINARY" if data.preliminary else ""
    output = (
        output
        or workbook.parent
        / f"{profile.file_prefix}_Remediation_Executive_Dashboard_"
        f"{data.latest.month}.{data.latest.day}.{data.latest:%y}{suffix}.pptx"
    ).resolve()
    require(output.suffix.lower() == ".pptx", "Output must use the .pptx extension.")
    require(
        output not in {template, workbook}, "Output must not overwrite the workbook or template."
    )
    require(not output.exists(), f"Output already exists: {output}. Choose a new filename.")
    require(output.parent.is_dir(), f"Output directory does not exist: {output.parent}")
    before_layout = [
        [(s.shape_id, s.left, s.top, s.width, s.height) for s in slide.shapes]
        for slide in deck.slides
    ]
    latest, previous = data.latest, data.previous
    latest_open, previous_open = data.open_at(latest), data.open_at(previous)
    latest_closed, previous_closed = data.closed_at(latest), data.closed_at(previous)
    delta_closed = latest_closed - previous_closed
    subtitle, chart_subtitle, takeaway = trend_language(data)
    month, previous_month = latest.strftime("%B"), previous.strftime("%b")
    latest_short = latest.strftime("%b")
    cover_date = date_ordinal(latest)
    if data.preliminary:
        month = f"{data.reporting_date:%B} preview"
        previous_month, latest_short = previous.strftime("%b %d"), latest.strftime("%b %d")
        cover_date = f"{data.reporting_date:%B %Y} (Preliminary)\nAs of {date_ordinal(latest)}"
    addressed = "resolved / accepted" if data.accepted_risk_supported else "resolved / ignored"
    if not data.accepted_risk_supported:
        takeaway = takeaway.replace("resolved or accepted", "resolved or ignored")
    replace_text(deck.slides[0], 2, cover_date.splitlines())
    progress = deck.slides[1]
    fewer = base_total - latest_open
    progress_values = {
        4: [f"{subtitle}  |  {baseline_label()} to {full_date(latest)}"],
        6: [f"{abs(fewer)} {'fewer' if fewer >= 0 else 'more'}"],
        7: [
            f"open findings since {base_month} baseline\n"
            f"{base_total} \u2192 {latest_open} open backlog"
        ],
        9: [reduction(base_total, latest_open)],
        10: [f"overall backlog reduction\n{base_short} baseline \u2192 {latest_short} latest"],
        12: [f"{delta_closed:+d}"],
        13: [
            f"{'closed / accepted' if data.accepted_risk_supported else 'closed / ignored'}"
            f"{' in ' + month if delta_closed >= 0 else ' balance change'}\n"
            f"{previous_month} \u2192 {latest_short}"
        ],
        15: [f"{latest_open - previous_open:+d}"],
        16: [
            f"open backlog in {month}\n{previous_month} \u2192 {latest_short}: "
            f"{change_description(previous_open, latest_open)}"
        ],
        19: [chart_subtitle],
        23: [previous_month],
        24: [f"{previous_open} open findings at the prior snapshot."],
        25: [latest_short],
        26: [
            f"{delta_closed:+d} {addressed}.\nOpen backlog: {previous_open} \u2192 {latest_open}."
        ],
        28: [
            f"Open findings {base_total} \u2192 {latest_open} since {base_month}; "
            f"{change_description(base_total, latest_open)}."
        ],
        30: [takeaway],
    }
    for sid, paragraphs in progress_values.items():
        replace_text(progress, sid, paragraphs)
    achievement = deck.slides[2]
    critical, high = data.snapshots[latest]["CRITICAL"], data.snapshots[latest]["HIGH"]
    achievement_values = {
        4: [f"Severity progress and key exposure risks closed by {full_date(latest)}"],
        6: [str(latest_closed)],
        7: [f"{addressed.capitalize()}\nBy {latest:%b} {latest.day}, across Critical and High."],
        9: [str(latest_open)],
        10: ["Open findings remaining", f"{base_month} baseline: {base_total} open findings."],
        12: [pct(high.addressed_pct)],
        13: [f"High resolution\n{high.closed} of {high.total} High findings addressed."],
        21: [pct(critical.addressed_pct)],
        22: [addressed],
        23: [
            f"{critical.closed} of {critical.total} Critical findings addressed; "
            f"{critical.open} remain open."
        ],
        27: [pct(high.addressed_pct)],
        28: [addressed],
        29: [f"{high.closed} of {high.total} High findings addressed; {high.open} remain open."],
        43: [
            "Accepted risk = team has formally accepted the risk and committed to a fix-by date."
            if data.accepted_risk_supported
            else (
                "Ignored findings are counted separately from resolved findings; "
                "formal risk acceptance is not asserted."
            )
        ],
    }
    callouts = exposure_callouts(data)
    for (heading_id, body_id), callout in zip(
        ((33, 34), (37, 38), (41, 42)), callouts, strict=False
    ):
        achievement_values[heading_id] = [callout.heading]
        achievement_values[body_id] = [callout.body]
    for sid, paragraphs in achievement_values.items():
        replace_text(achievement, sid, paragraphs)
    payload = {
        "schema_version": 1,
        "scope": ["CRITICAL", "HIGH"],
        "points": [{"date": day.isoformat(), "open": value} for day, value in points.items()],
    }
    if data.preliminary:
        payload["preliminary"] = True
    replace_picture(
        progress,
        31,
        line_chart(points, chart_subtitle, fonts),
        HISTORY_PREFIX + json.dumps(payload, separators=(",", ":")),
    )
    replace_picture(
        achievement,
        44,
        severity_chart(
            latest, {"CRITICAL": critical.open, "HIGH": high.open}, fonts, base_counts, baseline
        ),
        f"Open findings: baseline Critical {base_counts['CRITICAL']}, High {base_counts['HIGH']}; "
        f"{latest}: Critical {critical.open}, High {high.open}",
    )
    evidence = {
        "workbook": workbook.name,
        "workbook_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
        "template": template.name,
        "fixed_baseline": {"date": baseline.isoformat(), "open": base_total},
        "latest": latest.isoformat(),
        "previous": previous.isoformat(),
        "latest_open": latest_open,
        "latest_closed_or_accepted": latest_closed,
        "latest_month_closed_change": delta_closed,
        "monthly_backlog_reduction_percent": reduction(previous_open, latest_open),
        "history_sources": {
            day.isoformat(): "fixed baseline"
            if day == baseline
            else "workbook"
            if day in data.snapshots
            else "approved prior deck"
            for day in points
        },
        "callouts": [
            {
                "heading": c.heading,
                "evidence": [
                    {
                        "issue_id": f.issue_id,
                        "cell": f.reference,
                        "status": f.status,
                        "closed_on": str(f.closed_on),
                    }
                    for f in c.evidence
                ],
            }
            for c in callouts
        ],
        "interpretation": (
            "Workbook figures are used as provided; "
            "generation validates consistency, not source completeness."
        ),
    }
    if data.preliminary:
        evidence["preliminary"] = (
            f"Latest snapshot is a preliminary live collection before the scheduled "
            f"{ordinal(baseline.day)} cutoff. Replace with the final dashboard after "
            "the cutoff refresh."
        )
    for slide in deck.slides:
        slide.notes_slide.notes_text_frame.text = (
            "Wiz Governance Reporter dashboard provenance\n" + json.dumps(evidence, indent=2)
        )
    deck.core_properties.title = (
        f"{profile.client_name} Executive Remediation Dashboard - {full_date(latest)}"
        + (" (Preliminary)" if data.preliminary else "")
    )
    # Validate serialization before publishing; exclusive creation avoids overwriting
    # another user's file.
    with tempfile.TemporaryDirectory(prefix="wizexpert governance dashboard-") as directory:
        pending = Path(directory) / "dashboard.pptx"
        deck.save(pending)
        verify = Presentation(pending)
        validate_template(verify)
        after_layout = [
            [(s.shape_id, s.left, s.top, s.width, s.height) for s in slide.shapes]
            for slide in verify.slides
        ]
        require(after_layout == before_layout, "Template geometry or shape order changed.")
        require(
            template_history(verify, None, allow_preliminary=data.preliminary) == points,
            "Saved chart history does not match calculated trend.",
        )
        for slide_index, values in ((1, progress_values), (2, achievement_values)):
            for sid, expected in values.items():
                actual = shape(verify.slides[slide_index], sid).text.replace("\v", "\n")
                require(
                    actual == "\n".join(expected),
                    f"Saved slide {slide_index + 1} shape {sid} text mismatch.",
                )
        with output.open("xb") as stream:
            stream.write(pending.read_bytes())
    return output


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Update the client's approved executive dashboard from an Excel report."
    )
    parser.add_argument(
        "--workbook",
        required=True,
        type=Path,
        help="Generated governance .xlsx report (read-only).",
    )
    parser.add_argument(
        "--template",
        required=True,
        type=Path,
        help="Latest approved 3-slide .pptx deck (required).",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="New .pptx path; defaults beside workbook, named by snapshot date.",
    )
    parser.add_argument(
        "--history",
        type=Path,
        help="Verified monthly history JSON, required for the original raster-only template.",
    )
    parser.add_argument(
        "--font-dir",
        type=Path,
        help=(
            "Directory containing licensed calibri.ttf and calibrib.ttf "
            "(defaults to Windows Fonts)."
        ),
    )
    add_profile_argument(parser)
    args = parser.parse_args()
    ensure_profile(args.profile)
    try:
        output = build_deck(args.workbook, args.template, args.output, args.history, args.font_dir)
    except (DashboardError, OSError) as exc:
        parser.exit(1, f"Error: {exc}\n")
    print(f"Created PowerPoint dashboard: {output}")
    print(
        "Render and visually review all slides before distribution. "
        "Workbook source completeness is not certified."
    )


if __name__ == "__main__":
    main()
