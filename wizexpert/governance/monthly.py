"""Create or refresh a matching monthly governance workbook and executive deck."""

from __future__ import annotations

import argparse
import re
import shutil
import tempfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.exceptions import InvalidFileException

from wizexpert.governance.dashboard import build_deck
from wizexpert.governance.dashboard_data import (
    DashboardError,
    baseline_day,
    read_snapshots,
    read_workbook,
    require,
)
from wizexpert.governance.monthly_files import (
    check_recovery,
    check_unchanged,
    digest,
    monthly_lock,
    publish_pair,
)
from wizexpert.governance.profile import active as active_profile
from wizexpert.governance.profile import add_profile_argument
from wizexpert.governance.profile import ensure as ensure_profile
from wizexpert.governance.wiz_report import (
    SEVERITY_ORDER,
    _parse_dt,
    apply_watermark,
    build_monthly_scope_tab,
    build_severity_tab,
    build_summary_tab,
    compute_metrics,
    default_snapshots,
    load_wiz_csv,
    merge_issue_rows,
    read_issue_rows_from_existing_tab,
)


def reporting_month(value: str) -> datetime:
    require(
        bool(re.fullmatch(r"\d{4}-\d{2}", value)), "Month must use YYYY-MM, for example 2026-10."
    )
    try:
        month = datetime.strptime(value, "%Y-%m")
    except ValueError as exc:
        raise DashboardError(f"Invalid reporting month: {value}") from exc
    baseline = baseline_day()
    require(
        month.year * 12 + month.month > baseline.year * 12 + baseline.month,
        f"Paired dashboards require a month after the {baseline:%B %Y} baseline.",
    )
    return month


@dataclass(frozen=True)
class MonthlyResult:
    workbook: Path
    deck: Path
    archive: Path


def validate_rows(rows: list[dict], source: Path) -> None:
    for index, row in enumerate(rows, 2):
        for field in ("Issue ID", "Title", "Severity", "Status", "Created At"):
            require(
                isinstance(row.get(field), str) and bool(row[field].strip()),
                f"{source} row {index}: missing {field}.",
            )
        row["Severity"] = row["Severity"].strip().upper()
        row["Status"] = row["Status"].strip().upper()
        if row["Status"] == "IGNORED":
            row["Status"] = "REJECTED"
        require(row["Severity"] in SEVERITY_ORDER, f"{source} row {index}: unsupported severity.")
        require(
            row["Status"] in {"OPEN", "IN_PROGRESS", "RESOLVED", "REJECTED"},
            f"{source} row {index}: unsupported status {row['Status']!r}.",
        )
        row["Resolved Time"] = row.get("Resolved Time") or ""
        require(
            isinstance(row["Resolved Time"], str), f"{source} row {index}: invalid Resolved Time."
        )
        created = _parse_dt(row["Created At"])
        closed = _parse_dt(row["Resolved Time"])
        require(created is not None, f"{source} row {index}: invalid Created At.")
        if row["Status"] in {"RESOLVED", "REJECTED"}:
            require(
                closed is not None and closed >= created,
                f"{source} row {index}: missing or invalid closure date.",
            )
        else:
            require(
                not row.get("Resolved Time", "").strip(),
                f"{source} row {index}: an open finding has a closure date; "
                "reconcile its lifecycle first.",
            )


def build_monthly_workbook(
    inputs: list[Path],
    histories: list[Path],
    output: Path,
    month: datetime,
    org: str,
    *,
    collected_at: datetime | None = None,
) -> None:
    snapshots = default_snapshots(month)
    preliminary = None
    if collected_at is not None:
        require(
            collected_at.utcoffset() is not None and collected_at.utcoffset().total_seconds() == 0,
            "Preliminary collection timestamp must be UTC with timezone information.",
        )
        require(collected_at <= datetime.now(UTC), "Collection cannot be in the future.")
        actual = collected_at.replace(tzinfo=None)
        scheduled = snapshots[-1][1]
        require(
            snapshots[-2][1] < actual < scheduled,
            "Preliminary collection must be after the previous scheduled cutoff "
            "and before its scheduled cutoff.",
        )
        snapshots[-1] = (actual.strftime("%b %d, %Y"), actual)
        preliminary = (
            f"PRELIMINARY | Collected through {actual:%Y-%m-%d %H:%M:%S} UTC | "
            f"Scheduled cutoff {scheduled:%Y-%m-%d} EOD | Live collection, not EOD history"
        )
    historical_rows: list[dict] = []
    approved: dict[str, dict] = {}
    for history in histories:
        wb = load_workbook(history)
        try:
            for severity in SEVERITY_ORDER:
                historical_rows.extend(read_issue_rows_from_existing_tab(wb, severity))
                if f"{severity} Issues" in wb.sheetnames:
                    approved.setdefault(severity, {})
                    for day, snapshot in read_snapshots(wb[f"{severity} Issues"]).items():
                        approved[severity].setdefault(day, snapshot)
        finally:
            wb.close()
    historical_rows, _ = merge_issue_rows([], historical_rows)
    validate_rows(historical_rows, Path("retained-history"))
    current_rows: list[dict] = []
    for source in inputs:
        rows = load_wiz_csv(str(source))
        validate_rows(rows, source)
        if collected_at is not None:
            for row in rows:
                require(
                    _parse_dt(row["Created At"]) <= actual
                    and (not row["Resolved Time"] or _parse_dt(row["Resolved Time"]) <= actual),
                    f"{source}: issue {row['Issue ID']} has an event "
                    "after the collection timestamp.",
                )
        if len(rows) == 5000:
            print(
                f"Warning: {source.name} contains exactly 5,000 rows. "
                "Verify export completeness before sharing."
            )
        current_rows.extend(rows)
    current_rows, _ = merge_issue_rows([], current_rows)
    all_rows, supplemented = merge_issue_rows(historical_rows, current_rows)
    wb = Workbook()
    del wb["Sheet"]
    metrics, retained = {}, {}
    try:
        for severity in SEVERITY_ORDER:
            rows = [r for r in all_rows if r["Severity"] == severity]
            if not rows and severity not in {"CRITICAL", "HIGH"} and severity not in approved:
                continue
            retained[severity] = [r for r in supplemented if r["Severity"] == severity]
            metrics[severity] = compute_metrics(rows, snapshots)
            if histories:
                for index, (_, day) in enumerate(snapshots[:-1]):
                    snapshot = approved.get(severity, {}).get(day.date())
                    require(
                        snapshot is not None,
                        f"Missing approved {severity} snapshot for {day.date()}. "
                        "Review the prior workbook rather than reconstructing approved figures.",
                    )
                    for metric in ("total", "open", "resolved", "ignored"):
                        metrics[severity][metric][index] = getattr(snapshot, metric)
                    metrics[severity]["pct"][index] = f"{snapshot.addressed_pct:.1f}%"
            build_severity_tab(
                wb,
                severity,
                rows,
                snapshots,
                retained[severity],
                snapshot_metrics=metrics[severity],
            )
        build_summary_tab(wb, metrics, snapshots)
        build_monthly_scope_tab(wb, metrics, snapshots, supplemented_by_severity=retained)
        apply_watermark(wb, org=org)
        if preliminary:
            for sheet in wb:
                sheet["A2"] = f"{preliminary}\n{sheet['A2'].value}"
                sheet["A2"].alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
                sheet.row_dimensions[2].height = 48
                sheet.oddHeader.center.text = "PRELIMINARY - NOT FINAL MONTH-END DATA"
            wb.properties.description = preliminary
        wb.save(output)
    finally:
        wb.close()


def run_monthly(
    month: str,
    template: Path,
    output_dir: Path,
    *,
    inputs: list[Path] | None = None,
    workbook: Path | None = None,
    previous_report: Path | None = None,
    history: Path | None = None,
    font_dir: Path | None = None,
    org: str | None = None,
) -> MonthlyResult:
    period = reporting_month(month)
    profile = active_profile()
    org = org or profile.watermark
    cutoff = period.date().replace(day=profile.snapshot_day)
    require(
        bool(inputs) != (workbook is not None),
        "Supply CSV --input files OR a corrected --workbook, not both.",
    )
    require(
        workbook is None or previous_report is None,
        "--previous-report is only used with CSV --input.",
    )
    output_dir = output_dir.resolve()
    require(output_dir.is_dir(), f"Output directory does not exist: {output_dir}")
    excel = output_dir / f"{period:%b%y}_{profile.file_prefix}_Governance_Report.xlsx"
    deck = output_dir / (
        f"{profile.file_prefix}_Remediation_Executive_Dashboard_"
        f"{period.month}.{profile.snapshot_day}.{period:%y}.pptx"
    )
    template = template.resolve()
    inputs = [p.resolve() for p in inputs or []]
    workbook = workbook.resolve() if workbook is not None else None
    previous_report = previous_report.resolve() if previous_report is not None else None
    history = history.resolve() if history is not None else None
    sources = [
        template,
        *inputs,
        *([workbook] if workbook else []),
        *([previous_report] if previous_report else []),
        *([history] if history else []),
    ]
    for source in sources:
        require(source.is_file(), f"Source not found: {source}")
    for target in (excel, deck):
        require(not target.is_symlink(), f"Monthly output must not be a symlink: {target}")
    require(previous_report != excel, "The previous report cannot be this month's output.")
    require(
        not set(inputs).intersection((excel, deck)), "CSV inputs cannot be monthly output files."
    )
    with monthly_lock(output_dir, month):
        check_recovery(output_dir, month)
        expected = {p: digest(p) for p in {*sources, excel, deck}}
        if excel.exists() != deck.exists():
            require(
                workbook == excel and excel.exists() and not deck.exists(),
                "Only one monthly output exists. Restore the matching pair, "
                "or explicitly adopt the "
                "existing Excel file with --workbook before continuing.",
            )
        if excel.exists() and workbook != excel:
            require(
                read_workbook(excel).latest == cutoff,
                "Existing monthly workbook has a different reporting date.",
            )
        if previous_report:
            previous = read_workbook(previous_report).latest
            require(
                previous.year * 12 + previous.month == period.year * 12 + period.month - 1,
                "--previous-report must be the immediately preceding month's workbook.",
            )
        with tempfile.TemporaryDirectory(
            prefix=f".wiz-monthly-{month}-", dir=output_dir
        ) as directory:
            stage = Path(directory)
            staged_excel, staged_deck = stage / excel.name, stage / deck.name
            if workbook:
                shutil.copy2(workbook, staged_excel)
            else:
                histories = [p for p in (previous_report, excel if excel.exists() else None) if p]
                require(
                    bool(histories),
                    "First CSV run requires --previous-report to preserve approved snapshots. "
                    "Alternatively, use --workbook to adopt a separately reviewed first report.",
                )
                build_monthly_workbook(inputs, histories, staged_excel, period, org)
            data = read_workbook(staged_excel)
            require(data.latest == cutoff, "Workbook latest snapshot does not match --month.")
            build_deck(staged_excel, template, staged_deck, history, font_dir)
            check_unchanged(expected)
            archive = publish_pair((staged_excel, staged_deck), (excel, deck), expected, month)
        return MonthlyResult(excel, deck, archive)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Create or refresh one month's Excel + PowerPoint pair, "
            "with backups before replacement."
        )
    )
    parser.add_argument(
        "--month",
        required=True,
        help="Explicit reporting month, e.g. 2026-10 (cut off on the profile's snapshot day).",
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument(
        "--input", "-i", nargs="+", type=Path, help="Complete current Wiz CSV export set."
    )
    source.add_argument(
        "--workbook", type=Path, help="Use a corrected workbook as-is, without rebuilding from CSV."
    )
    parser.add_argument(
        "--template", required=True, type=Path, help="Approved PowerPoint template."
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path.cwd(),
        help="Monthly pair directory (default: current directory).",
    )
    parser.add_argument(
        "--previous-report",
        type=Path,
        help="Prior month's workbook for retained closed issues in CSV mode.",
    )
    parser.add_argument(
        "--history", type=Path, help="Verified chart history for an original raster-only template."
    )
    parser.add_argument(
        "--font-dir", type=Path, help="Chart font directory, as for wizexpert governance dashboard."
    )
    parser.add_argument(
        "--org",
        default=None,
        help="Excel watermark text in CSV mode (default: the profile's watermark).",
    )
    add_profile_argument(parser)
    args = parser.parse_args()
    ensure_profile(args.profile)
    try:
        result = run_monthly(
            args.month,
            args.template,
            args.output_dir,
            inputs=args.input,
            workbook=args.workbook,
            previous_report=args.previous_report,
            history=args.history,
            font_dir=args.font_dir,
            org=args.org,
        )
    except (DashboardError, OSError, BadZipFile, InvalidFileException) as exc:
        parser.exit(1, f"Error: {exc}\n")
    print(f"Updated monthly Excel: {result.workbook}")
    print(f"Updated monthly PowerPoint: {result.deck}")
    print(f"Run record / previous-file backups: {result.archive}")
    print("Review both reports before sharing. This command does not watch for later file edits.")


if __name__ == "__main__":
    main()
