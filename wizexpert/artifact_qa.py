"""Read-only checks of actual deliverables against independently reviewed private contracts."""

from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from wizexpert.email_qa import Document, compact, stylesheet, validate_email
from wizexpert.office_qa import package_signature
from wizexpert.workflow_qa import digest, load, require

KINDS = {"email", "html", "pdf", "xlsx", "pptx", "validation_receipt"}
FOOTER_KINDS = {"html", "pdf", "xlsx"}


def footer_line(contract: dict, kind: str) -> str | None:
    """Return the required confidentiality line; an explicit null records 'not required'."""
    if kind not in FOOTER_KINDS and "confidentiality_footer" not in contract:
        return None
    require(
        "confidentiality_footer" in contract,
        "Contract must record the confidentiality-footer decision (text or null).",
    )
    value = contract["confidentiality_footer"]
    require(
        value is None or (isinstance(value, str) and bool(compact(value))),
        "Confidentiality footer must be nonempty text or null.",
    )
    return compact(value) if value else None


def check_text(actual: str, contract: dict) -> None:
    require(
        isinstance(contract.get("expected_text"), str) and bool(contract["expected_text"]),
        "The reviewed full expected text is required; a few keywords are not full coverage.",
    )
    require(
        compact(actual) == compact(contract["expected_text"]),
        "Actual deliverable text differs from the reviewed content.",
    )
    if contract.get("audience", "client") == "client":
        require(
            not re.search(r"\b[A-Za-z]:[\\/]|file://", actual, re.I),
            "Local workstation paths leaked into a client deliverable.",
        )
    for phrase in contract.get("forbidden_text", []):
        require(phrase.casefold() not in actual.casefold(), "Disallowed wording in deliverable.")


def check_workbook(path: Path, contract: dict) -> None:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False, keep_links=True)
    try:
        require(not book._external_links, "External workbook links are not allowed.")
        require(book.sheetnames == contract["sheet_names"], "Workbook sheet names/order differ.")
        if "campaign_workbook" in contract:
            from wizexpert.campaign import check_campaign_workbook

            check_campaign_workbook(book, contract["campaign_workbook"], path)
        for sheet in book:
            require(sheet.sheet_state == "visible", "Hidden worksheets are not allowed.")
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        require(
                            cell.data_type != "f" or contract.get("allow_formulas") is True,
                            "Formula cells are not allowed.",
                        )
                        require(
                            not contract.get("forbid_hyperlinks") or cell.hyperlink is None,
                            "Workbook hyperlinks are prohibited by this contract.",
                        )
            require(
                not any(d.hidden for d in sheet.row_dimensions.values())
                and not any(d.hidden for d in sheet.column_dimensions.values()),
                "Hidden workbook rows/columns are not allowed.",
            )
            line = footer_line(contract, "xlsx")
            if line:
                settings = sheet.HeaderFooter
                footers = [settings.oddFooter]
                if settings.differentOddEven:
                    footers.append(settings.evenFooter)
                if settings.differentFirst:
                    footers.append(settings.firstFooter)
                require(
                    all(
                        any(
                            line in compact(part.text or "") for part in (f.left, f.center, f.right)
                        )
                        for f in footers
                    ),
                    f"Worksheet {sheet.title!r} lacks the confidentiality page footer.",
                )
    finally:
        book.close()


def validate_artifact(item: dict, evidence: dict) -> dict:
    kind = item.get("kind")
    require(kind in KINDS, f"No implemented artifact checker for {kind}; coverage is blocked.")
    path, contract_path = Path(item["path"]), Path(item["contract"])
    for source in (path, contract_path):
        require(
            source.is_absolute() and evidence.get(str(source)) == digest(source),
            f"Artifact or contract was not reviewed: {source}",
        )
    contract = load(contract_path)
    require(contract.get("schema_version") == 1, "Unsupported artifact contract.")
    if kind == "email":
        result = validate_email(path, contract_path)
        require(
            not result["errors"], "Actual email failed template QA: " + "; ".join(result["errors"])
        )
    elif kind == "html":
        document = Document(path.read_text(encoding="utf-8-sig"))
        require(not document.errors, "Malformed HTML deliverable.")
        require(
            not any(
                e.tag in {"script", "iframe", "object", "embed", "form"}
                for e in document.root.elements()
            ),
            "Active HTML content is not allowed.",
        )
        template = Path(contract["approved_template"])
        require(
            template.is_absolute()
            and template.resolve() != path.resolve()
            and evidence.get(str(template)) == digest(template),
            "HTML template must be independently reviewed and hash-bound.",
        )
        approved = Document(template.read_text(encoding="utf-8-sig"))
        require(not approved.errors, "Approved HTML template is malformed.")
        styles = stylesheet(document)
        require(
            compact(styles) == compact(stylesheet(approved)),
            "HTML CSS differs from the approved template.",
        )
        require(
            not re.search(r"@import|url\s*\(", styles, re.I), "CSS imports/assets are not allowed."
        )

        def styling(doc: Document) -> set:
            return {
                (n.tag, n.attrs.get("class", ""), n.attrs.get("style", ""))
                for n in doc.root.elements()
            }

        require(
            styling(document) == styling(approved), "HTML element styling differs from template."
        )
        body = document.by_id(contract["body_id"])
        require(
            any(node is body for node in document.root.elements(visible_only=True)),
            "Report body is hidden.",
        )
        for node in document.root.elements():
            require(
                not any(k.startswith("on") for k in node.attrs),
                "HTML event handlers are not allowed.",
            )
            require(
                not re.search(r"@import|url\s*\(", node.attrs.get("style", "") or "", re.I),
                "Inline CSS imports/assets are not allowed.",
            )
            for key in ("src", "href"):
                value = node.attrs.get(key, "") or ""
                if key == "href" and value in contract.get("allowed_links", []):
                    require(
                        value.startswith("https://"), "Only reviewed HTTPS links are supported."
                    )
                    require(node.tag != "link", "External stylesheets are not allowed.")
                else:
                    require(
                        not re.match(r"\s*(https?:|//|file:|javascript:|vbscript:)", value, re.I)
                        and not (key == "href" and value.strip().casefold().startswith("data:")),
                        "Unreviewed external/active HTML references are not allowed.",
                    )
        check_text(body.text(), contract)
        headings = [
            (e.tag, e.text())
            for e in body.elements(visible_only=True)
            if re.fullmatch(r"h[1-6]", e.tag)
        ]
        require([list(h) for h in headings] == contract["headings"], "HTML section layout differs.")
        line = footer_line(contract, "html")
        if line:
            footers = [
                e
                for e in document.root.elements(visible_only=True)
                if e.tag == "footer" and e.text() == line
            ]
            require(len(footers) == 1, "HTML lacks the exact visible confidentiality footer.")
    elif kind == "pdf":
        try:
            import pymupdf
        except ModuleNotFoundError as error:
            raise ValueError('PDF QA needs the qa extra: pip install "wizexpert[qa]"') from error
        line = footer_line(contract, "pdf")
        with pymupdf.open(path) as document:
            require(not document.is_encrypted, "Encrypted PDF cannot be validated.")
            require(document.page_count == contract["page_count"], "PDF page count differs.")
            pages = []
            for number, page in enumerate(document, 1):
                text = compact(page.get_text())
                if line:
                    require(line in text, f"PDF page {number} lacks the confidentiality footer.")
                    # Expected text describes the content once, without per-page footers.
                    text = text.replace(line, "", 1)
                pages.append(text)
            check_text(" ".join(pages), contract)
            allowed = set(contract.get("allowed_links", []))
            require(
                all(
                    link.get("uri", "").startswith("https://") and link["uri"] in allowed
                    for page in document
                    for link in page.get_links()
                ),
                "PDF contains links outside the reviewed HTTPS allowlist.",
            )
    elif kind in {"xlsx", "pptx"}:
        from defusedxml.ElementTree import fromstring

        expected = Path(contract["reference_artifact"])
        require(
            expected.is_absolute()
            and expected.resolve() != path.resolve()
            and evidence.get(str(expected)) == digest(expected),
            "Independent Office replay/reference must be reviewed and hash-bound.",
        )
        require(bool(contract.get("reference_basis")), "Record the independent reference basis.")
        sources = contract.get("reference_sources")
        require(
            isinstance(sources, list) and bool(sources) and all(s in evidence for s in sources),
            "Office reference requires reviewed input/template evidence.",
        )
        actual = package_signature(path)
        main_part = "xl/workbook.xml" if kind == "xlsx" else "ppt/presentation.xml"
        require(main_part in actual, "Office file has the wrong package type.")
        require(not any("vba" in name.casefold() for name in actual), "Macros are not allowed.")
        require(
            actual == package_signature(expected),
            "Actual Office content, structure or formatting differs from reviewed replay.",
        )
        if kind == "xlsx":
            check_workbook(path, contract)
        else:
            with ZipFile(path) as archive:
                tree = fromstring(archive.read(main_part))
            slides = tree.findall(
                ".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"
            )
            require(len(slides) == contract["slide_count"], "Presentation slide count differs.")
            line = footer_line(contract, "pptx")
            if line:
                with ZipFile(path) as archive:
                    names = sorted(
                        n
                        for n in archive.namelist()
                        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
                    )
                    for name in names:
                        texts = fromstring(archive.read(name)).iter(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
                        )
                        require(
                            line in compact(" ".join(t.text or "" for t in texts)),
                            f"Slide part {name} lacks the confidentiality footer.",
                        )
    else:
        # Existing specialized validators keep responsibility for their own domain logic.
        actual = load(path)
        require(
            contract.get("success_status") in {"passed", "automated_checks_passed", "validated"},
            "A failed/unknown status cannot be configured as a successful validation.",
        )
        require(
            actual.get("status") == contract["success_status"],
            "Specialized validation did not pass.",
        )
        require(
            actual.get("share_ready") is False,
            "Validation receipt must not grant release approval.",
        )
        hashes = actual.get(contract["hash_field"])
        required_sources = contract["required_sources"]
        require(
            isinstance(hashes, dict)
            and isinstance(required_sources, list)
            and bool(required_sources),
            "Validation receipt must bind its actual inputs/outputs.",
        )
        require(
            set(required_sources) <= set(hashes), "Specialized receipt omitted required sources."
        )
        for name in required_sources:
            require(
                evidence.get(name) == hashes[name] == digest(Path(name)),
                "Specialized validator receipt is stale or covers different files.",
            )
        require(
            contract["validator_source"] in evidence and contract["execution_evidence"] in evidence,
            "Specialized validation needs its reviewed implementation and execution evidence.",
        )
        implementation = actual.get("implementation_sha256")
        require(
            isinstance(implementation, dict) and contract["validator_source"] in implementation,
            "Validation receipt must bind the implementation that actually ran.",
        )
        for name, value in implementation.items():
            require(
                evidence.get(name) == value == digest(Path(name)),
                "Specialized validator implementation changed since execution.",
            )
    return {
        "id": item["id"],
        "kind": kind,
        "status": "artifact_checks_passed",
        "path": str(path),
        "sha256": digest(path),
        "share_ready": False,
    }


def validate_native_review(item: dict, review: dict, evidence: dict) -> None:
    from wizexpert.workflow_qa import timestamp

    path = item.get("native_review")
    require(
        isinstance(path, str) and path in evidence,
        "Every rendered deliverable needs its own hash-bound native/visual review.",
    )
    require(evidence[path] == digest(Path(path)), "Native review changed after approval.")
    native = load(Path(path))
    require(
        native.get("schema_version") == 1 and native.get("status") == "passed",
        "Native review is incomplete or unsupported.",
    )
    require(
        native.get("artifact_sha256") == evidence[item["path"]]
        and native.get("kind") == item["kind"],
        "Native review covers a different artifact.",
    )
    require(
        native.get("reviewer") == review["reviewer"] and bool(native.get("application")),
        "Native review must name the reviewing application and separate reviewer.",
    )
    require(
        timestamp(review["started_at"])
        <= timestamp(native["reviewed_at"])
        <= timestamp(review["reviewed_at"]),
        "Native review is outside the reviewed interval.",
    )
    checks = {"all_content_reviewed", "readable", "no_clipping", "no_hidden_or_unexpected_content"}
    if item["kind"] in {"xlsx", "pptx"}:
        checks.add("opened_without_repair")
    require(
        all(native.get("checks", {}).get(k) is True for k in checks),
        "Native/visual checks are missing or failed.",
    )
    refs = native.get("evidence")
    require(
        isinstance(refs, list) and bool(refs) and all(p in evidence for p in refs),
        "Native review needs retained rendering/application evidence.",
    )


def validate_deliverables(review: dict, evidence: dict) -> list:
    requested = review.get("requested_deliverables")
    require(isinstance(requested, list), "Explicit requested-deliverable scope is required.")
    ids = [r["id"] for r in requested]
    require(
        len(ids) == len(set(ids)) and all(isinstance(i, str) and i for i in ids),
        "Requested deliverable IDs must be unique and nonempty.",
    )
    for row in requested:
        require(
            row.get("kind") in KINDS,
            "Requested artifact kind has no checker; do not claim coverage.",
        )
    groups = review.get("parity_groups")
    require(isinstance(groups, list), "Explicit cross-format parity scope is required.")
    kinds = {r["id"]: r["kind"] for r in requested}
    for group in groups:
        require(
            isinstance(group, list)
            and len(group) >= 2
            and len(set(group)) == len(group)
            and all(kinds.get(name) in {"html", "pdf"} for name in group),
            "Text parity groups must name distinct requested HTML/PDF artifacts.",
        )
    if review["phase"] != "handoff":
        return []
    artifacts = review.get("artifacts", [])
    require(
        isinstance(artifacts, list)
        and len(artifacts) == len(requested)
        and {r["id"] for r in artifacts} == set(ids),
        "Requested deliverables are missing, duplicated or substituted.",
    )
    expected = {r["id"]: r["kind"] for r in requested}
    for item in artifacts:
        require(
            item["kind"] == expected[item["id"]], "Deliverable kind differs from approved scope."
        )
    if review["workflow"] in {"governance", "campaign"}:
        require(
            bool(artifacts), "This workflow cannot claim completion without actual deliverables."
        )
    results = []
    for item in artifacts:
        results.append(validate_artifact(item, evidence))
        if item["kind"] != "validation_receipt":
            validate_native_review(item, review, evidence)
    by_id = {item["id"]: item for item in artifacts}

    def content(name: str) -> str:
        # HTML shows the confidentiality footer once; PDF repeats it on each page.
        contract = load(Path(by_id[name]["contract"]))
        text = compact(contract["expected_text"])
        footer = contract.get("confidentiality_footer")
        return compact(text.replace(compact(footer), " ")) if footer else text

    for group in groups:
        texts = [content(name) for name in group]
        footers = [
            load(Path(by_id[name]["contract"])).get("confidentiality_footer") for name in group
        ]
        require(
            len({compact(f) if f else None for f in footers}) == 1,
            "HTML/PDF companions use different confidentiality footer decisions.",
        )
        require(
            len(set(texts)) == 1, "HTML/PDF text parity failed despite individual artifact checks."
        )
    return results
