import json
from pathlib import Path
from shutil import copyfile

import pytest
from openpyxl import Workbook, load_workbook

from wizexpert.artifact_qa import validate_artifact, validate_deliverables
from wizexpert.workflow_qa import digest


def save(path, payload):
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def artifact(tmp_path, kind):
    output = tmp_path / f"report.{kind}"
    contract = {"schema_version": 1}
    source = tmp_path / "source.txt"
    source.write_text("Synthetic reviewed input", encoding="utf-8")
    reference = tmp_path / f"reference.{kind}"
    if kind == "html":
        output.write_text(
            '<html><body><main id="report"><h1>Findings</h1>'
            "<p>Reviewed evidence.</p></main></body></html>",
            encoding="utf-8",
        )
        copyfile(output, reference)
        contract.update(
            body_id="report",
            expected_text="Findings Reviewed evidence.",
            headings=[["h1", "Findings"]],
            approved_template=str(reference),
        )
    elif kind == "pdf":
        pymupdf = pytest.importorskip("pymupdf")
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), "Findings\nReviewed evidence.")
            document.save(output)
        contract.update(page_count=1, expected_text="Findings Reviewed evidence.")
    elif kind == "xlsx":
        book = Workbook()
        book.active.title = "Findings"
        book.active.append(["Issue", "Count"])
        book.active.append(["Synthetic issue", 2])
        book.save(output)
        book.close()
        copyfile(output, reference)
        contract["sheet_names"] = ["Findings"]
    elif kind == "pptx":
        pptx = pytest.importorskip("pptx")
        deck = pptx.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Synthetic findings"
        deck.save(output)
        copyfile(output, reference)
        contract["slide_count"] = 1
    if kind in {"xlsx", "pptx"}:
        contract.update(
            reference_artifact=str(reference),
            reference_basis="Synthetic replay",
            reference_sources=[str(source)],
        )
    contract_path = save(tmp_path / f"{kind}-contract.json", contract)
    evidence = {str(p): digest(p) for p in (source, output, contract_path)}
    if reference.exists():
        evidence[str(reference)] = digest(reference)
    return {"id": kind, "kind": kind, "path": str(output), "contract": str(contract_path)}, evidence


def native_review(tmp_path, item, evidence):
    native = save(
        tmp_path / f"{item['id']}-native.json",
        {
            "schema_version": 1,
            "status": "passed",
            "kind": item["kind"],
            "artifact_sha256": evidence[item["path"]],
            "reviewer": "reviewer",
            "application": "Synthetic application evidence",
            "reviewed_at": "2026-10-27T07:30:00Z",
            "checks": dict.fromkeys(
                [
                    "all_content_reviewed",
                    "readable",
                    "no_clipping",
                    "no_hidden_or_unexpected_content",
                    "opened_without_repair",
                ],
                True,
            ),
            "evidence": [str(tmp_path / "source.txt")],
        },
    )
    item["native_review"] = str(native)
    evidence[str(native)] = digest(native)


def handoff(items):
    return {
        "workflow": "finding",
        "phase": "handoff",
        "reviewer": "reviewer",
        "started_at": "2026-10-27T07:00:00Z",
        "reviewed_at": "2026-10-27T08:00:00Z",
        "requested_deliverables": [{"id": a["id"], "kind": a["kind"]} for a in items],
        "artifacts": items,
        "parity_groups": [],
    }


@pytest.mark.parametrize("kind", ["html", "pdf", "xlsx", "pptx"])
def test_actual_non_email_artifacts_pass_and_cannot_grant_release(tmp_path, kind):
    item, evidence = artifact(tmp_path, kind)
    native_review(tmp_path, item, evidence)
    result = validate_deliverables(handoff([item]), evidence)
    assert result[0]["status"] == "artifact_checks_passed"
    assert result[0]["share_ready"] is False


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_full_office_package_drift_not_just_counts_is_blocked(tmp_path, kind):
    item, evidence = artifact(tmp_path, kind)
    output = Path(item["path"])
    if kind == "xlsx":
        book = load_workbook(output)
        book.active["B2"] = 999
        book.save(output)
        book.close()
    else:
        from pptx import Presentation

        deck = Presentation(output)
        deck.slides[0].shapes.title.text = "Unsupported claim"
        deck.save(output)
    evidence[str(output)] = digest(output)
    with pytest.raises(ValueError, match="differs from reviewed replay"):
        validate_artifact(item, evidence)


@pytest.mark.parametrize(
    "kind,field,value,message",
    [
        ("xlsx", "sheet_names", ["Unexpected"], "sheet names"),
        ("pptx", "slide_count", 3, "slide count"),
        ("pdf", "page_count", 2, "page count"),
        ("pdf", "expected_text", "Wrong content", "text differs"),
        ("html", "headings", [["h3", "Findings"]], "section layout"),
    ],
)
def test_structure_and_content_contracts_are_checked(tmp_path, kind, field, value, message):
    item, evidence = artifact(tmp_path, kind)
    path = Path(item["contract"])
    contract = json.loads(path.read_text())
    contract[field] = value
    save(path, contract)
    evidence[str(path)] = digest(path)
    with pytest.raises(ValueError, match=message):
        validate_artifact(item, evidence)


@pytest.mark.parametrize("mutation", ["style", "hidden", "event"])
def test_html_actual_style_visibility_and_active_content(tmp_path, mutation):
    item, evidence = artifact(tmp_path, "html")
    path = Path(item["path"])
    content = path.read_text()
    replacements = {
        "style": ("<p>", '<p style="color:red">'),
        "hidden": ("<body>", "<body hidden>"),
        "event": ("<p>", '<p onclick="alert(1)">'),
    }
    path.write_text(content.replace(*replacements[mutation]), encoding="utf-8")
    evidence[str(path)] = digest(path)
    with pytest.raises(ValueError, match="styling|hidden|event"):
        validate_artifact(item, evidence)


def test_requested_outputs_cannot_be_omitted_or_replaced_by_unsupported_format(tmp_path):
    item, evidence = artifact(tmp_path, "html")
    review = handoff([item])
    review["artifacts"] = []
    with pytest.raises(ValueError, match="missing"):
        validate_deliverables(review, evidence)
    review["requested_deliverables"][0]["kind"] = "txt"
    with pytest.raises(ValueError, match="no checker"):
        validate_deliverables(review, evidence)


def test_stale_native_review_and_missing_native_evidence_block(tmp_path):
    item, evidence = artifact(tmp_path, "xlsx")
    review = handoff([item])
    with pytest.raises(ValueError, match="native/visual"):
        validate_deliverables(review, evidence)
    native_review(tmp_path, item, evidence)
    path = Path(item["native_review"])
    native = json.loads(path.read_text())
    native["artifact_sha256"] = "0" * 64
    save(path, native)
    evidence[str(path)] = digest(path)
    with pytest.raises(ValueError, match="different artifact"):
        validate_deliverables(review, evidence)


def test_html_pdf_parity_is_not_replaced_by_individual_passes(tmp_path):
    html, first = artifact(tmp_path, "html")
    pdf, second = artifact(tmp_path, "pdf")
    evidence = first | second
    for item in (html, pdf):
        native_review(tmp_path, item, evidence)
    review = handoff([html, pdf])
    review["parity_groups"] = [["html", "pdf"]]
    assert len(validate_deliverables(review, evidence)) == 2
    output, contract_path = Path(html["path"]), Path(html["contract"])
    output.write_text(output.read_text().replace("Reviewed evidence.", "Different evidence."))
    contract = json.loads(contract_path.read_text())
    contract["expected_text"] = "Findings Different evidence."
    save(contract_path, contract)
    evidence.update({str(output): digest(output), str(contract_path): digest(contract_path)})
    native_review(tmp_path, html, evidence)
    with pytest.raises(ValueError, match="parity failed"):
        validate_deliverables(review, evidence)


def test_specialized_receipt_must_bind_current_sources_and_execution(tmp_path):
    source = tmp_path / "input.txt"
    source.write_text("Synthetic input")
    implementation = tmp_path / "validator.py"
    implementation.write_text("# Synthetic validator reference\n")
    execution = tmp_path / "execution.log"
    execution.write_text("Synthetic successful execution record")
    output = save(
        tmp_path / "receipt.json",
        {
            "status": "automated_checks_passed",
            "share_ready": False,
            "source_and_output_sha256": {str(source): digest(source)},
            "implementation_sha256": {str(implementation): digest(implementation)},
        },
    )
    contract = save(
        tmp_path / "contract.json",
        {
            "schema_version": 1,
            "success_status": "automated_checks_passed",
            "hash_field": "source_and_output_sha256",
            "required_sources": [str(source)],
            "validator_source": str(implementation),
            "execution_evidence": str(execution),
        },
    )
    item = {
        "id": "domain",
        "kind": "validation_receipt",
        "path": str(output),
        "contract": str(contract),
    }
    evidence = {str(p): digest(p) for p in (source, implementation, execution, output, contract)}
    assert validate_artifact(item, evidence)["status"] == "artifact_checks_passed"
    implementation.write_text("# Changed implementation\n")
    evidence[str(implementation)] = digest(implementation)
    with pytest.raises(ValueError, match="implementation changed"):
        validate_artifact(item, evidence)
    source.write_text("Changed after validation")
    evidence[str(source)] = digest(source)
    with pytest.raises(ValueError, match="stale"):
        validate_artifact(item, evidence)
