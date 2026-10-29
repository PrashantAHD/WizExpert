import hashlib
import json
import os
import re
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from test_email_qa import email  # noqa: F401  (fixture reuse)

from wizexpert import campaign
from wizexpert.__main__ import main
from wizexpert.artifact_qa import check_workbook
from wizexpert.campaign import (
    CampaignConfig,
    campaign_wording_errors,
    check_campaign_workbook,
    classify_failure,
    finding_count_note,
    lifecycle,
    read_json,
    record_failure,
    write_json,
)
from wizexpert.email_qa import validate_email
from wizexpert.workflow_qa import collection_check, required_checks, rule_hashes

T0 = datetime(2026, 1, 5, 8, 0, tzinfo=UTC)
IDS = ("TEST-VULN-A", "TEST-VULN-B")


def config(**override):
    return CampaignConfig(IDS, **override)


def at(**delta):
    return T0 + timedelta(**delta)


def catalog(external_id):
    return {"vulnerabilityByExternalId": {"externalId": external_id}}


def grouped(count):
    return {
        "vulnerabilityFindingsGroupedByValues": {
            "nodes": [{"cloudPlatform": "Cloud1", "analytics": {"totalFindingCount": count}}]
            if count
            else [],
            "pageInfo": {"endCursor": None, "hasNextPage": False},
            "totalCount": 1 if count else 0,
        }
    }


def finding_page(items, more, cursor):
    return {
        "vulnerabilityFindings": {
            "nodes": [
                {"id": f"f-{a}-{s}", "status": s, "vulnerableAsset": {"id": a}} for a, s in items
            ],
            "pageInfo": {"endCursor": cursor, "hasNextPage": more},
            "totalCount": 2,
        }
    }


def resources(ids):
    return {
        "cloudResourcesV2": {
            "nodes": [{"id": i} for i in ids],
            "pageInfo": {"endCursor": None, "hasNextPage": False},
            "totalCount": len(ids),
        }
    }


def save_catalogs(cfg, directory, now):
    for i in IDS:
        campaign.save(
            cfg, directory, "get_vulnerability_catalog", {"external_id": i}, catalog(i), now=now
        )


def grouped_params(status):
    return {
        "vulnerability_external_ids": list(IDS),
        "status": status,
        "group_by": ["CLOUD_PLATFORM"],
        "first": 20,
    }


def save_grouped(cfg, directory, now, counts=None):
    counts = counts or {"all": 2, "remaining": 1, "resolved": 1, "rejected": 0}
    for kind, status in campaign.GROUPED_KINDS.items():
        campaign.save(
            cfg,
            directory,
            "list_vulnerability_findings_grouped",
            grouped_params(status),
            grouped(counts[kind]),
            now=now,
        )


def finding_params(**extra):
    return {
        "vulnerability_external_ids": list(IDS),
        "status": campaign.ALL_STATUSES,
        "first": 20,
        **extra,
    }


def test_no_default_vulnerability_ids():
    with pytest.raises(ValueError, match="no defaults"):
        CampaignConfig(())
    with pytest.raises(ValueError, match="positive"):
        config(stall_seconds=0)
    assert config().stall_seconds == 900


def test_start_plans_catalogs_then_grouped_then_one_stream_then_resources(tmp_path):
    cfg = config()
    started = campaign.start(cfg, tmp_path, now=T0)
    directory = started["dir"]
    first = started["next"]
    assert first["action"] == "wiz-execute-batch"
    assert [c["parameters"]["external_id"] for c in first["calls"]] == list(IDS)
    save_catalogs(cfg, directory, at(seconds=5))
    nxt = campaign.next_call(cfg, directory, now=at(seconds=6))
    assert [c["parameters"]["status"] for c in nxt["calls"]] == list(
        campaign.GROUPED_KINDS.values()
    )
    assert all(c["parameters"]["first"] == 20 for c in nxt["calls"])
    save_grouped(cfg, directory, at(seconds=10))
    nxt = campaign.next_call(cfg, directory, now=at(seconds=11))
    assert nxt["action"] == "wiz-execute"
    assert nxt["parameters"] == finding_params()
    page = finding_page([("a1", "OPEN")], True, "c1")
    nxt = campaign.save(
        cfg, directory, "list_vulnerability_findings", finding_params(), page, now=at(seconds=12)
    )["next"]
    assert nxt["parameters"]["after"] == "c1"
    page = finding_page([("a2", "RESOLVED")], False, "c2")
    nxt = campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(after="c1"),
        page,
        now=at(seconds=13),
    )["next"]
    assert nxt["action"] == "wiz-execute-batch"
    call = nxt["calls"][0]
    assert call["parameters"] == {
        "resource_id": ["a1", "a2"],
        "include_deleted": True,
        "first": 40,
    }


def test_resource_lookups_batch_by_forty(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    save_grouped(
        cfg, directory, at(seconds=2), {"all": 45, "remaining": 45, "resolved": 0, "rejected": 0}
    )
    items = [(f"a{i}", "OPEN") for i in range(20)]
    campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(),
        finding_page(items, True, "c1"),
        now=at(seconds=3),
    )
    items = [(f"a{i}", "OPEN") for i in range(20, 40)]
    campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(after="c1"),
        finding_page(items, True, "c2"),
        now=at(seconds=4),
    )
    items = [(f"a{i}", "OPEN") for i in range(40, 45)]
    nxt = campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(after="c2"),
        finding_page(items, False, "c3"),
        now=at(seconds=5),
    )["next"]
    assert [len(c["parameters"]["resource_id"]) for c in nxt["calls"]] == [40, 5]


def test_split_status_stream_and_oversized_page_are_rejected(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    page = finding_page([("a1", "OPEN")], False, "c")
    split = finding_params()
    split["status"] = ["OPEN", "IN_PROGRESS"]
    with pytest.raises(ValueError, match="do not split"):
        campaign.save(cfg, directory, "list_vulnerability_findings", split, page, now=at(seconds=1))
    big = finding_params()
    big["first"] = 50
    with pytest.raises(ValueError, match="20"):
        campaign.save(cfg, directory, "list_vulnerability_findings", big, page, now=at(seconds=1))


def complete_run(tmp_path, gap):
    """A full run whose saves are `gap` apart (no total limit)."""
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    clock = T0
    steps = []

    def tick():
        nonlocal clock
        clock += gap
        return clock

    for i in IDS:
        campaign.save(
            cfg, directory, "get_vulnerability_catalog", {"external_id": i}, catalog(i), now=tick()
        )
    counts = {"all": 2, "remaining": 1, "resolved": 1, "rejected": 0}
    for kind, status in campaign.GROUPED_KINDS.items():
        campaign.save(
            cfg,
            directory,
            "list_vulnerability_findings_grouped",
            grouped_params(status),
            grouped(counts[kind]),
            now=tick(),
        )
    page = finding_page([("a1", "OPEN"), ("a2", "RESOLVED")], False, "c")
    campaign.save(cfg, directory, "list_vulnerability_findings", finding_params(), page, now=tick())
    steps.append(
        campaign.save(
            cfg,
            directory,
            "list_cloud_resources",
            {"resource_id": ["a1", "a2"], "include_deleted": True, "first": 40},
            resources(["a1", "a2"]),
            now=tick(),
        )
    )
    return cfg, directory, steps[-1]["next"], clock


def test_no_total_time_limit_while_data_keeps_arriving(tmp_path):
    cfg, directory, nxt, clock = complete_run(tmp_path, timedelta(minutes=14))
    assert (clock - T0) > timedelta(hours=1)  # far beyond any old fixed window
    assert nxt["action"] == "assemble"
    out = tmp_path / "out" / "manifest.json"
    manifest = campaign.assemble(cfg, directory, out)
    assert manifest["observed_from"] == "2026-01-05T08:14:00Z"
    assert manifest["observed_to"] == campaign.iso(clock)
    assert manifest["finding_count"] == 2 and manifest["asset_count"] == 2
    assert read_json(out)["run_id"] == manifest["run_id"]


def test_assemble_blocks_when_group_totals_do_not_reconcile(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    save_grouped(
        cfg, directory, at(seconds=2), {"all": 3, "remaining": 2, "resolved": 1, "rejected": 0}
    )
    page = finding_page([("a1", "OPEN"), ("a2", "RESOLVED")], False, "c")
    campaign.save(
        cfg, directory, "list_vulnerability_findings", finding_params(), page, now=at(seconds=3)
    )
    campaign.save(
        cfg,
        directory,
        "list_cloud_resources",
        {"resource_id": ["a1", "a2"], "include_deleted": True, "first": 40},
        resources(["a1", "a2"]),
        now=at(seconds=4),
    )
    with pytest.raises(ValueError, match="recollect"):
        campaign.assemble(cfg, directory, tmp_path / "m.json")


def test_collection_has_no_elapsed_time_limit(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=10))
    result = campaign.next_call(cfg, directory, now=at(seconds=10 + 30 * 24 * 3600))
    assert result["action"] == "wiz-execute-batch"
    assert not (Path(directory) / "aborted.json").exists()


def test_legacy_stall_config_does_not_limit_collection(tmp_path):
    cfg = config(stall_seconds=60)
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    assert campaign.next_call(cfg, directory, now=at(seconds=60))["action"] == "wiz-execute-batch"
    assert (
        campaign.next_call(cfg, directory, now=at(seconds=61 + 30 * 24 * 3600))["action"]
        == "wiz-execute-batch"
    )


def test_complete_run_is_never_stalled(tmp_path):
    cfg, directory, _, clock = complete_run(tmp_path, timedelta(seconds=1))
    assert campaign.next_call(cfg, directory, now=clock + timedelta(days=1))["action"] == "assemble"


def test_runs_cannot_be_stitched_across_directories(tmp_path):
    cfg = config()
    first = campaign.start(cfg, tmp_path / "a", now=T0)["dir"]
    second = campaign.start(cfg, tmp_path / "b", now=T0)["dir"]
    save_catalogs(cfg, first, at(seconds=1))
    stale = next((tmp_path / "a").glob("staging-*/catalog-TEST-VULN-A.json"))
    (tmp_path / "b" / stale.parent.name / stale.name).write_bytes(stale.read_bytes())
    # Same stamp gives the same run id, so make the second run distinct.
    meta = read_json(f"{second}/meta.json")
    meta["run_id"] = "another-run"
    write_json(f"{second}/meta.json", meta)
    with pytest.raises(ValueError, match="never stitch"):
        campaign.next_call(cfg, second, now=at(seconds=2))


def test_scope_change_is_rejected(tmp_path):
    directory = campaign.start(config(), tmp_path, now=T0)["dir"]
    with pytest.raises(ValueError, match="scope differs"):
        campaign.next_call(CampaignConfig(("OTHER",)), directory, now=T0)


def test_helper_json_is_utf8_without_bom_outside_staging_and_reader_tolerates_bom(tmp_path):
    staging = tmp_path / "staging"
    staging.mkdir()
    outside = tmp_path / "helpers" / "params.json"
    write_json(outside, {"name": "caf\u00e9"}, staging=staging)
    raw = outside.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf") and raw.endswith(b"\n")
    with pytest.raises(ValueError, match="outside the staging"):
        write_json(staging / "params.json", {}, staging=staging)
    bom = tmp_path / "bom.json"
    bom.write_bytes(b"\xef\xbb\xbf" + b'{"first": 20}')
    assert read_json(bom) == {"first": 20}


def test_cli_flow_rejects_helpers_inside_staging_and_accepts_bom(tmp_path, capsys):
    base = ["campaign"]
    ids = ["--vuln-id", IDS[0], "--vuln-id", IDS[1]]
    assert main([*base, "start", *ids, "--base", str(tmp_path / "runs")]) == 0
    directory = json.loads(capsys.readouterr().out)["dir"]
    params = tmp_path / "params.json"
    response = tmp_path / "response.json"
    params.write_bytes(b"\xef\xbb\xbf" + json.dumps({"external_id": IDS[0]}).encode())
    write_json(response, catalog(IDS[0]))
    save_args = [*base, "save", *ids, "--dir", directory, "--tool", "get_vulnerability_catalog"]
    assert main([*save_args, "--params-file", str(params), "--response-file", str(response)]) == 0
    capsys.readouterr()
    inside = tmp_path / "runs" / "x"
    inside_params = f"{directory}/params.json"
    write_json(inside_params, {"external_id": IDS[1]})
    assert main([*save_args, "--params-file", inside_params, "--response-file", str(response)]) == 1
    assert "outside the staging" in capsys.readouterr().err
    assert not inside.exists()
    assert main([*base, "plan"]) == 0
    assert "ONE finding stream" in capsys.readouterr().out


def test_auth_challenge_stops_and_repeated_failures_abort(tmp_path):
    assert classify_failure(401) == "stop_reconnect"
    assert classify_failure(None, "OAuth challenge: invalid_token") == "stop_reconnect"
    assert classify_failure(500, "timeout") == "retry"
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    params = {"external_id": IDS[0]}
    for expected in (1, 2):
        result = record_failure(directory, "get_vulnerability_catalog", params, status=500)
        assert result["action"] == "retry" and result["attempts"] == expected
    result = record_failure(directory, "get_vulnerability_catalog", params, status=500)
    assert result["action"] == "abort_unable_to_fetch"
    assert (
        campaign.next_call(cfg, directory, now=at(seconds=1))["action"] == "abort_unable_to_fetch"
    )
    with pytest.raises(ValueError, match="unable to fetch"):
        save_catalogs(cfg, directory, at(seconds=2))


def test_successful_save_clears_earlier_failures(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    params = {"external_id": IDS[0]}
    record_failure(directory, "get_vulnerability_catalog", params, status=500)
    campaign.save(
        cfg, directory, "get_vulnerability_catalog", params, catalog(IDS[0]), now=at(seconds=1)
    )
    assert read_json(f"{directory}/attempts.json") == {}


def state(**override):
    data = {
        "status": "active",
        "vulnerability_ids": list(IDS),
        "fully_patched_ids": [],
        "closure_gate": {"met": False},
    }
    data.update(override)
    return data


def test_lifecycle_refreshes_until_closure_and_never_sends_or_mutates():
    scheduled = lifecycle(state(fully_patched_ids=[IDS[0]]), "scheduled")
    assert scheduled["action"] == "refresh" and scheduled["prepare_drafts_only"] is True
    assert scheduled["send_email"] is False and scheduled["mutate_wiz_or_cloud"] is False
    assert lifecycle(state(), "manual")["prepare_drafts_only"] is False


def test_lifecycle_stops_on_closure_and_asks_to_disable_automation():
    met = state(fully_patched_ids=list(IDS), closure_gate={"met": True, "confirmed_by": "operator"})
    result = lifecycle(met, "scheduled")
    assert result["action"] == "stop" and result["disable_automation"] is True
    assert result["prepare_final_closure_update"] is True
    closed = lifecycle(state(status="closed"), "scheduled")
    assert closed["action"] == "stop" and closed["prepare_final_closure_update"] is False
    assert "automation" in closed["tell_owner"]


def test_lifecycle_rejects_gate_met_without_full_patching():
    bad = state(fully_patched_ids=[IDS[0]], closure_gate={"met": True, "confirmed_by": "operator"})
    with pytest.raises(ValueError, match="fully patched"):
        lifecycle(bad)
    with pytest.raises(ValueError, match="manual or scheduled"):
        lifecycle(state(), "other")


def test_lifecycle_cli(tmp_path, capsys):
    path = tmp_path / "state.json"
    write_json(path, state())
    assert main(["campaign", "lifecycle", "--state", str(path), "--mode", "scheduled"]) == 0
    assert json.loads(capsys.readouterr().out)["action"] == "refresh"


def sheet_rows(book, title, rows, note):
    sheet = book.create_sheet(title)
    sheet.append(["Finding", "Asset"])
    for row in rows:
        sheet.append(row)
    sheet.append([None, None])
    sheet.append([note, None])


def count_workbook(path, remaining_note=None, resolved_note=None, extra=None):
    book = Workbook()
    book.remove(book.active)
    sheet_rows(
        book,
        "Remaining Findings",
        [["f1", "a1"], ["f2", "a2"]],
        remaining_note or "Total Remaining Un-Resolved Findings = 2",
    )
    sheet_rows(
        book, "Resolved Findings", [["f3", "a3"]], resolved_note or "Total Resolved Findings = 1"
    )
    for title in extra or []:
        sheet_rows(book, title, [["f9", "a9"]], "n/a")
    book.save(path)


SPEC = {
    "count_sheets": {"Remaining Findings": "remaining", "Resolved Findings": "resolved"},
    "group_names": ["Team Alpha"],
}


def test_finding_count_notes_are_exact():
    assert finding_count_note([1, 2, 3], "remaining") == "Total Remaining Un-Resolved Findings = 3"
    assert finding_count_note(4, "resolved") == "Total Resolved Findings = 4"
    with pytest.raises(ValueError, match="remaining or resolved"):
        finding_count_note(1, "rejected")


def test_campaign_workbook_notes_pass_and_fail(tmp_path):
    good = tmp_path / "good.xlsx"
    count_workbook(good)
    book = load_workbook(good)
    check_campaign_workbook(book, SPEC, good)
    for kwargs in (
        {"remaining_note": "Total Remaining Findings = 2"},
        {"remaining_note": "Total Remaining Un-Resolved Findings = 3"},
        {"resolved_note": "Total Resolved Findings = 2"},
    ):
        bad = tmp_path / "bad.xlsx"
        count_workbook(bad, **kwargs)
        with pytest.raises(ValueError, match="bottom note"):
            check_campaign_workbook(load_workbook(bad), SPEC, bad)


def test_campaign_workbook_rejects_group_sheets_and_workbooks(tmp_path):
    extra = tmp_path / "extra.xlsx"
    count_workbook(extra, extra=["Team Alpha"])
    with pytest.raises(ValueError, match="extra or per-group"):
        check_campaign_workbook(load_workbook(extra), SPEC, extra)
    named = tmp_path / "Team Alpha findings.xlsx"
    count_workbook(named)
    with pytest.raises(ValueError, match="Per-group campaign workbooks"):
        check_campaign_workbook(load_workbook(named), SPEC, named)
    permitted = {**SPEC, "other_sheets": ["Team Alpha"], "group_names": []}
    check_campaign_workbook(load_workbook(extra), permitted, extra)


def test_artifact_check_applies_campaign_workbook_spec(tmp_path):
    good = tmp_path / "good.xlsx"
    count_workbook(good)
    contract = {
        "sheet_names": ["Remaining Findings", "Resolved Findings"],
        "confidentiality_footer": None,
        "campaign_workbook": SPEC,
    }
    check_workbook(good, contract)
    bad = tmp_path / "bad.xlsx"
    count_workbook(bad, remaining_note="Total Remaining Un-Resolved Findings = 9")
    with pytest.raises(ValueError, match="bottom note"):
        check_workbook(bad, contract)


@pytest.mark.parametrize(
    "text,fragment",
    [
        ("Remaining 4. Rejected 1.", "Rejected"),
        ("Resolved 5 findings.", "Wiz lifecycle"),
        ("Resolved in Wiz: 5. These are verified patching results.", "verified"),
        ("Reconstructed value for the day.", "UTC cutoff"),
        ("Partial-day observation, resolved in Wiz.", "UTC cutoff"),
    ],
)
def test_campaign_wording_flags(text, fragment):
    assert any(fragment in e for e in campaign_wording_errors(text))


def test_campaign_wording_accepts_qualified_text():
    text = (
        "Remaining 4. Resolved in Wiz 5 (Wiz lifecycle status, not verified patching). "
        "Reconstructed from retained data as of 15:40 UTC; partial-day observation."
    )
    assert campaign_wording_errors(text) == []
    assert campaign_wording_errors("Not verified patching only.") == []
    assert campaign_wording_errors("x", partial_day=True) != []
    assert campaign_wording_errors("x at 09:05:00Z", partial_day=True) == []


def progress_email(email, body, kind="progress", campaign=True):  # noqa: F811
    target, contract = email
    payload = json.loads(contract.read_text())
    payload.update(
        kind=kind,
        approved_subject="Reviewed update",
        approved_body=body,
        issue_mappings=[],
        attachments={},
    )
    if campaign:
        payload["campaign"] = True
    contract.write_text(json.dumps(payload), encoding="utf-8")
    content = target.read_text().replace(
        "Review requested: Synthetic finding | SNOW: INC123", "Reviewed update"
    )
    content = re.sub(
        r'<section id="email-body">.*?</section>',
        f'<section id="email-body"><p>{body}</p></section>',
        content,
        flags=re.S,
    )
    target.write_text(content, encoding="utf-8")
    return target, contract


def test_progress_email_flags_rejected_unqualified_resolved_and_missing_cutoff(email):  # noqa: F811
    wrapper = "Hi Team, {} Best Regards, Example Security Team"
    ok = wrapper.format("Remaining 4. Resolved in Wiz 5 (lifecycle status).")
    assert validate_email(*progress_email(email, ok))["errors"] == []
    rejected = wrapper.format("Remaining 4. Resolved in Wiz 5. Rejected 1.")
    errors = validate_email(*progress_email(email, rejected))["errors"]
    assert any("Rejected" in e for e in errors)
    bare = wrapper.format("Resolved 5.")
    assert any("Wiz lifecycle" in e for e in validate_email(*progress_email(email, bare))["errors"])
    partial = wrapper.format("Resolved in Wiz 5; reconstructed daily value.")
    assert any("UTC cutoff" in e for e in validate_email(*progress_email(email, partial))["errors"])
    cutoff = wrapper.format("Resolved in Wiz 5; reconstructed daily value as of 15:40 UTC.")
    assert validate_email(*progress_email(email, cutoff))["errors"] == []


def test_wording_applies_to_campaign_closure_but_not_non_campaign_progress(email):  # noqa: F811
    body = "Hi Team, Rejected items are resolved. Best Regards, Example Security Team"
    # Non-campaign progress email: not subject to campaign wording.
    plain = progress_email(email, body, kind="progress", campaign=False)
    assert validate_email(*plain)["errors"] == []
    # Campaign closure draft: checked.
    closure = progress_email(email, body, kind="closure", campaign=True)
    errors = validate_email(*closure)["errors"]
    assert any("Rejected" in e for e in errors)
    # Campaign flag on a progress email: checked.
    flagged = progress_email(email, body, kind="progress", campaign=True)
    assert validate_email(*flagged)["errors"]


def write_log(tmp_path):
    log = tmp_path / "activity.jsonl"
    log.write_text(
        json.dumps(
            {
                "timestamp": (T0 - timedelta(minutes=5)).isoformat(),
                "type": "tool.execution_start",
                "data": {"toolName": "wiz-execute", "arguments": {}},
            }
        )
        + "\n",
        encoding="utf-8",
    )
    return log


def qa_manifest(block=None, **extra):
    data = {
        "started_at": (T0 - timedelta(hours=1)).isoformat(),
        "ended_at": (T0 + timedelta(hours=3)).isoformat(),
        "reused_sources": ["prior"],
        "counts": [{"name": "n", "expected": 1, "collected": 1}],
        **extra,
    }
    if block is not None:
        data["campaign_collection"] = block
    return data


def bound_block(tmp_path, cfg, directory):
    out = tmp_path / "assembled" / "manifest.json"
    campaign.assemble(cfg, directory, out)
    return {
        "assemble_manifest": str(out),
        "assemble_manifest_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
    }


def flags_for(tmp_path, block, workflow="campaign"):
    log = write_log(tmp_path)
    return collection_check(log, qa_manifest(block), workflow)


def test_campaign_review_requires_block_bound_to_assemble_manifest(tmp_path):
    cfg, directory, _, _ = complete_run(tmp_path / "run", timedelta(seconds=1))
    block = bound_block(tmp_path, cfg, directory)
    assert flags_for(tmp_path, block)["status"] == "pass"
    missing = flags_for(tmp_path, None)
    assert missing["status"] == "flagged" and any("missing" in f for f in missing["flags"])
    # Self-reported claims are not accepted.
    claims = {"staging_dirs": ["x"], "complete": True, "auth_ok": True}
    assert flags_for(tmp_path, claims)["status"] == "flagged"
    # Non-campaign workflows only check the block when one is supplied.
    assert collection_check(write_log(tmp_path), qa_manifest(), "finding")["status"] == "pass"
    wrong_hash = {**block, "assemble_manifest_sha256": "0" * 64}
    assert flags_for(tmp_path, wrong_hash)["status"] == "flagged"


def test_bound_block_detects_changed_extra_aborted_and_stale_evidence(tmp_path):
    cfg, directory, _, _ = complete_run(tmp_path / "run", timedelta(seconds=1))
    block = bound_block(tmp_path, cfg, directory)
    staging = Path(directory)
    catalog_file = staging / "catalog-TEST-VULN-A.json"
    original = catalog_file.read_bytes()
    catalog_file.write_bytes(original.replace(b"TEST-VULN-A", b"TEST-VULN-A", 1) + b" ")
    assert any("changed" in f for f in flags_for(tmp_path, block)["flags"])
    catalog_file.write_bytes(original)
    (staging / "findings-0009.json").write_bytes(b"{}")
    assert any("differs" in f for f in flags_for(tmp_path, block)["flags"])
    (staging / "findings-0009.json").unlink()
    (staging / "aborted.json").write_text("{}", encoding="utf-8")
    assert any("aborted" in f for f in flags_for(tmp_path, block)["flags"])


def test_bound_block_detects_evidence_that_predates_the_run(tmp_path):
    cfg, directory, _, _ = complete_run(tmp_path / "run", timedelta(seconds=1))
    stale_file = Path(directory) / "catalog-TEST-VULN-A.json"
    record = read_json(stale_file)
    record["observed_at"] = campaign.iso(at(hours=-1))  # older than the run's start
    write_json(stale_file, record)
    block = bound_block(tmp_path, cfg, directory)
    assert any("predates" in f for f in flags_for(tmp_path, block)["flags"])


def test_collection_check_cli_workflow_option(tmp_path, capsys):
    log = write_log(tmp_path)
    path = tmp_path / "m.json"
    write_json(path, qa_manifest())
    base = ["workflow-qa", "collection-check", "--activity-log", str(log), "--manifest", str(path)]
    assert main(base) == 0
    capsys.readouterr()
    assert main([*base, "--workflow", "campaign"]) == 1


def test_campaign_checks_have_no_fixed_window_and_cover_auth_and_partial_packs():
    collection = required_checks("campaign", "collection")
    assert "snapshot-inventory-stall-and-identity" in collection
    assert "snapshot-inventory-window-and-identity" not in collection
    assert "wiz-auth-and-model-availability" in collection
    assert "no-partial-pack-and-qualified-wording" in required_checks("campaign", "handoff")


def test_campaign_module_is_in_every_rule_hash():
    for workflow in ("governance", "campaign", "finding"):
        assert "wizexpert/campaign.py" in rule_hashes(workflow)


def test_legacy_stall_metadata_does_not_limit_collection(tmp_path):
    directory = campaign.start(config(stall_seconds=60), tmp_path, now=T0)["dir"]
    assert read_json(f"{directory}/meta.json")["stall_seconds"] == 60
    lax = config(stall_seconds=10**6)
    result = campaign.next_call(lax, directory, now=at(seconds=61))
    assert result["action"] == "wiz-execute-batch"


def test_abort_is_sticky_for_next_save_and_assemble(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    write_json(
        Path(directory) / "aborted.json",
        {
            "action": "abort_unable_to_fetch",
            "reason": "A fetch failed three times.",
        },
    )
    assert (
        campaign.next_call(cfg, directory, now=at(seconds=901))["action"] == "abort_unable_to_fetch"
    )
    assert (Path(directory) / "aborted.json").is_file()
    # Even with a "fresh" clock, the run never resumes.
    again = campaign.next_call(cfg, directory, now=T0)
    assert again["action"] == "abort_unable_to_fetch" and again["aborted"] is True
    with pytest.raises(ValueError, match="closed"):
        save_catalogs(cfg, directory, T0)
    with pytest.raises(ValueError, match="closed"):
        campaign.assemble(cfg, directory, tmp_path / "m.json")


def test_auth_challenge_marker_blocks_planning_until_new_start(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path / "a", now=T0)["dir"]
    assert (
        record_failure(directory, "get_vulnerability_catalog", {}, status=401)["action"]
        == "stop_reconnect"
    )
    nxt = campaign.next_call(cfg, directory, now=at(seconds=1))
    assert nxt["action"] == "stop_reconnect" and "calls" not in nxt
    with pytest.raises(ValueError, match="reconnect"):
        save_catalogs(cfg, directory, at(seconds=2))
    fresh = campaign.start(cfg, tmp_path / "b", now=at(minutes=5))
    assert fresh["next"]["action"] == "wiz-execute-batch"


def test_more_than_one_hundred_finding_pages_complete_in_numeric_order(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    pages = 105
    save_grouped(
        cfg,
        directory,
        at(seconds=2),
        {"all": pages, "remaining": pages, "resolved": 0, "rejected": 0},
    )
    nxt = None
    for number in range(pages):
        params = finding_params(**({"after": f"c{number - 1}"} if number else {}))
        page = finding_page([(f"a{number:03d}", "OPEN")], number < pages - 1, f"c{number}")
        nxt = campaign.save(
            cfg, directory, "list_vulnerability_findings", params, page, now=at(seconds=3 + number)
        )["next"]
    names = sorted(p.name for p in Path(directory).glob("findings-*.json"))
    assert len(names) == pages and names[100] == "findings-0100.json"
    assert nxt["action"] == "wiz-execute-batch"
    assert [len(c["parameters"]["resource_id"]) for c in nxt["calls"]] == [40, 40, 25]
    first = nxt["calls"][0]["parameters"]["resource_id"]
    assert first == [f"a{i:03d}" for i in range(40)]


def test_repeated_or_empty_cursors_and_empty_pages_are_refused(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    for cursor in (None, ""):
        page = finding_page([("a1", "OPEN")], True, cursor)
        with pytest.raises(ValueError, match="endCursor"):
            campaign.save(
                cfg,
                directory,
                "list_vulnerability_findings",
                finding_params(),
                page,
                now=at(seconds=1),
            )
    empty = finding_page([], True, "c1")
    with pytest.raises(ValueError, match="Empty finding page"):
        campaign.save(
            cfg,
            directory,
            "list_vulnerability_findings",
            finding_params(),
            empty,
            now=at(seconds=1),
        )
    first = finding_page([("a1", "OPEN")], True, "c1")
    campaign.save(
        cfg, directory, "list_vulnerability_findings", finding_params(), first, now=at(seconds=2)
    )
    repeated = finding_page([("a2", "OPEN")], True, "c1")
    with pytest.raises(ValueError, match="repeated a cursor"):
        campaign.save(
            cfg,
            directory,
            "list_vulnerability_findings",
            finding_params(after="c1"),
            repeated,
            now=at(seconds=3),
        )


def test_assemble_refuses_duplicate_finding_ids(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    save_grouped(
        cfg, directory, at(seconds=2), {"all": 2, "remaining": 2, "resolved": 0, "rejected": 0}
    )
    same = [("a1", "OPEN")]
    campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(),
        finding_page(same, True, "c1"),
        now=at(seconds=3),
    )
    campaign.save(
        cfg,
        directory,
        "list_vulnerability_findings",
        finding_params(after="c1"),
        finding_page(same, False, "c2"),
        now=at(seconds=4),
    )
    campaign.save(
        cfg,
        directory,
        "list_cloud_resources",
        {"resource_id": ["a1"], "include_deleted": True, "first": 40},
        resources(["a1"]),
        now=at(seconds=5),
    )
    with pytest.raises(ValueError, match="Duplicate finding IDs"):
        campaign.assemble(cfg, directory, tmp_path / "m.json")


def test_cli_refuses_replayed_helper_responses(tmp_path, capsys):
    ids = ["--vuln-id", IDS[0], "--vuln-id", IDS[1]]
    assert main(["campaign", "start", *ids, "--base", str(tmp_path / "runs")]) == 0
    directory = json.loads(capsys.readouterr().out)["dir"]
    params = tmp_path / "params.json"
    response = tmp_path / "response.json"
    write_json(params, {"external_id": IDS[0]})
    write_json(response, catalog(IDS[0]))
    old = time.time() - 3600
    os.utime(response, (old, old))
    args = ["campaign", "save", *ids, "--dir", directory, "--tool", "get_vulnerability_catalog"]
    assert main([*args, "--params-file", str(params), "--response-file", str(response)]) == 1
    assert "replay" in capsys.readouterr().err
    os.utime(response, None)
    assert main([*args, "--params-file", str(params), "--response-file", str(response)]) == 0


@pytest.mark.parametrize("bad", ["", "a b", "../x", "x/y", "x;y", "a\\b", "caf\u00e9"])
def test_vulnerability_ids_are_validated(bad):
    with pytest.raises(ValueError, match="A-Za-z0-9"):
        CampaignConfig(("OK-1", bad))
    assert CampaignConfig(("A.b_c-1",)).vulnerability_ids == ("A.b_c-1",)


def test_assemble_output_is_never_overwritten(tmp_path):
    cfg, directory, _, _ = complete_run(tmp_path / "run", timedelta(seconds=1))
    out = tmp_path / "out.json"
    campaign.assemble(cfg, directory, out)
    first = out.read_bytes()
    with pytest.raises(ValueError, match="overwrite"):
        campaign.assemble(cfg, directory, out)
    assert out.read_bytes() == first


def test_save_uses_and_releases_an_exclusive_lock(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    lock = Path(directory) / "save.lock"
    lock.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="Another save is in progress"):
        save_catalogs(cfg, directory, at(seconds=1))
    lock.unlink()
    save_catalogs(cfg, directory, at(seconds=1))
    assert not lock.exists()
    # The lock is also released when a save is refused.
    with pytest.raises(ValueError):
        campaign.save(cfg, directory, "unknown_tool", {}, {}, now=at(seconds=2))
    assert not lock.exists()


def test_assemble_reports_unreturned_resources_without_closing_them(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    save_grouped(cfg, directory, at(seconds=2))
    page = finding_page([("a1", "OPEN"), ("a2", "RESOLVED")], False, "c")
    campaign.save(
        cfg, directory, "list_vulnerability_findings", finding_params(), page, now=at(seconds=3)
    )
    campaign.save(
        cfg,
        directory,
        "list_cloud_resources",
        {"resource_id": ["a1", "a2"], "include_deleted": True, "first": 40},
        resources(["a1"]),
        now=at(seconds=4),
    )
    manifest = campaign.assemble(cfg, directory, tmp_path / "m.json")
    assert manifest["finding_count"] == 2 and manifest["asset_count"] == 2
    assert manifest["returned_resource_count"] == 1
    assert manifest["collection_qa"]["status"] == "requires_reconciliation"
    flags = campaign.campaign_collection_flags(
        {
            "assemble_manifest": str((tmp_path / "m.json").resolve()),
            "assemble_manifest_sha256": hashlib.sha256(
                (tmp_path / "m.json").read_bytes()
            ).hexdigest(),
        }
    )
    assert any("requires reconciliation" in f for f in flags)
    assert manifest["unreturned_resource_ids"] == ["a2"]
    assert manifest["unreturned_resource_count"] == 1
    assert "requires reconciliation" in manifest["unreturned_resource_label"]


def _to_resource_stage(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    save_catalogs(cfg, directory, at(seconds=1))
    save_grouped(cfg, directory, at(seconds=2))
    page = finding_page([("a1", "OPEN"), ("a2", "RESOLVED")], False, "c")
    campaign.save(
        cfg, directory, "list_vulnerability_findings", finding_params(), page, now=at(seconds=3)
    )
    return cfg, directory


def test_empty_resource_batch_is_rejected_not_accepted_as_complete(tmp_path):
    cfg, directory = _to_resource_stage(tmp_path)
    params = {"resource_id": ["a1", "a2"], "include_deleted": True, "first": 40}
    with pytest.raises(ValueError, match="Empty resource batch"):
        campaign.save(
            cfg, directory, "list_cloud_resources", params, resources([]), now=at(seconds=4)
        )
    assert not list(Path(directory).glob("resources-*.json"))
    assert campaign.next_call(cfg, directory, now=at(seconds=5))["action"] == "wiz-execute-batch"


def test_exhausted_attempts_block_save_before_marker_exists(tmp_path):
    cfg = config()
    directory = campaign.start(cfg, tmp_path, now=T0)["dir"]
    params = {"external_id": IDS[0]}
    write_json(
        Path(directory) / "attempts.json",
        {campaign._call_key("get_vulnerability_catalog", params): 3},
    )
    with pytest.raises(ValueError, match="unable to fetch"):
        campaign.save(
            cfg, directory, "get_vulnerability_catalog", params, catalog(IDS[0]), now=at(seconds=1)
        )


def test_output_inside_public_repo_is_refused(tmp_path, monkeypatch):
    monkeypatch.delenv("WIZEXPERT_ALLOW_REPO_OUTPUT", raising=False)
    repo = Path(campaign.__file__).resolve().parents[1]
    if not (repo / "pyproject.toml").is_file():
        pytest.skip("not a source checkout")
    with pytest.raises(ValueError, match="outside the public WizExpert repository"):
        campaign.start(config(), repo / "tests" / "never-created", now=T0)
    campaign.start(config(), tmp_path, now=T0)
