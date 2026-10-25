"""The WizExpert rename preserves old entry points and workbook updates."""

import importlib
import subprocess
import sys
import tomllib
from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from wizexpert.csv_report import (
    CREATOR,
    LEGACY_CREATOR,
    REPOSITORY_ROOT,
    SHEETS,
    _update_digest,
)


@pytest.mark.parametrize(
    "module",
    ["__main__", "csv_report", "csv_layout", "execution_update", "workflow", "issue_metadata"],
)
def test_legacy_imports_share_canonical_module(module):
    assert importlib.import_module(f"remediation.{module}") is importlib.import_module(
        f"wizexpert.{module}"
    )


@pytest.mark.parametrize("package", ["wizexpert", "remediation"])
@pytest.mark.parametrize("command", [[], ["workflow"], ["update-remediation"]])
def test_module_commands_remain_available(package, command):
    result = subprocess.run(
        [sys.executable, "-m", package, *command, "--help"],
        cwd=REPOSITORY_ROOT, capture_output=True, text=True, check=False,
    )
    assert result.returncode == 0, result.stderr
    assert "usage:" in result.stdout
    assert result.stderr == ""


@pytest.mark.parametrize("creator", [CREATOR, LEGACY_CREATOR])
def test_update_validation_accepts_current_and_legacy_workbooks(tmp_path, creator):
    book = Workbook()
    book.remove(book.active)
    for name in SHEETS:
        book.create_sheet(name)
    book.properties.creator = creator
    book[SHEETS[0]]["A1"] = "Synthetic retained evidence"
    stream = BytesIO()
    book.save(stream)
    book.close()
    path = tmp_path / "report.xlsx"
    original = stream.getvalue()
    path.write_bytes(original)
    assert _update_digest(path) == sha256(original).digest()
    assert path.read_bytes() == original
    with BytesIO(path.read_bytes()) as content:
        restored = load_workbook(content)
        assert restored[SHEETS[0]]["A1"].value == "Synthetic retained evidence"
        restored.close()


def test_distribution_and_both_commands_target_wizexpert():
    metadata = tomllib.loads((REPOSITORY_ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    assert metadata["project"]["name"] == "wizexpert"
    assert metadata["project"]["scripts"] == {
        "wizexpert": "wizexpert.__main__:main",
        "cloud-remediations": "wizexpert.__main__:main",
    }
