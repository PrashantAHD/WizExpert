import hashlib
import json
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt
from synthetic import mail

from wizexpert.governance.dashboard import shape
from wizexpert.governance.dashboard_data import DashboardError
from wizexpert.governance.governance import (
    TABLE_BODY_FONT_SIZE,
    TABLE_CAPACITY,
    TEXT_IDS,
    build_governance,
    main,
    table_row,
)
from wizexpert.governance.governance_data import (
    AGGREGATE_KEYS,
    ISSUE_TYPES,
    SEVERITIES,
    capture_groups,
    cutoff_counts,
    load_snapshot,
)


def severity_counts(**values):
    return {severity: values.get(severity, 0) for severity in SEVERITIES}


def group(identifier, count=2, resolved=True, severity="MEDIUM"):
    return {
        "id": identifier,
        "name": f"Synthetic rule {identifier}",
        "counts": severity_counts(**{severity: count}),
        "recommendation": "### Protect resources\n* Apply least privilege to the service account.",
        "samples": [
            {
                "id": f"issue-{identifier}",
                "created_at": "2025-01-01T00:00:00Z",
                "resolved_at": "2025-10-01T00:00:00Z" if resolved else None,
            }
        ],
    }


def make_template(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = 12192000, 6858000
    for index, ids in TEXT_IDS.items():
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for sid in sorted(set(ids) | {2}):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
            box._element.xpath(".//p:cNvPr")[0].set("id", str(sid))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = "Stale September narrative"
            run.font.name, run.font.size = "Arial", Pt(18)
        if index in TABLE_CAPACITY:
            table = slide.shapes.add_table(
                TABLE_CAPACITY[index] + 1, 4, Inches(0.25), Inches(0.9), Inches(12.8), Inches(5.5)
            )
            table._element.xpath(".//p:cNvPr")[0].set("id", "7")
            for row in table.table.rows:
                for cell in row.cells:
                    cell.text = "Stale September narrative"
    shape(deck.slides[0], 5).text = "Monthly Cloud Security Governance"
    shape(deck.slides[0], 4).text = "Example Health"
    # Unused background text and an extra table must not survive generation.
    table = deck.slides[8].shapes.add_table(1, 4, 0, 0, Inches(4), Inches(0.5))
    table.table.cell(0, 0).text = "Stale extra table"
    deck.save(path)


def snapshot(template):
    score = {"value": 95, "source": "Synthetic metric query"}
    return {
        "schema_version": 1,
        "scope": {
            "kind": "TENANT",
            "name": "Synthetic dedicated tenant",
            "tenant_id": "test-tenant",
            "issue_types": sorted(ISSUE_TYPES),
        },
        "period": {
            "start_exclusive": "2025-09-15T23:59:59Z",
            "end_inclusive": "2025-10-15T23:59:59Z",
            "resolved_through": "2025-10-13T09:00:00Z",
            "collected_from": "2025-10-13T09:00:00Z",
            "collected_through": "2025-10-13T10:00:00Z",
        },
        "prior": {
            "date": "2025-09-15",
            "source": "Approved synthetic governance deck",
            "source_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
            "open": dict(zip(SEVERITIES[:4], (0, 0, 10, 1), strict=False)),
            **{key: deepcopy(score) for key in ("security_score", "hipaa", "hitrust")},
        },
        "open": severity_counts(MEDIUM=8, LOW=2, INFORMATIONAL=3),
        "in_progress": severity_counts(MEDIUM=2),
        "resolved": severity_counts(MEDIUM=2, INFORMATIONAL=1),
        **{key: deepcopy(score) for key in ("security_score", "hipaa", "hitrust")},
        "resolved_groups": {
            "total_groups": 2,
            "groups": [group("resolved"), group("info", 1, severity="INFORMATIONAL")],
        },
        "in_progress_groups": {"total_groups": 1, "groups": [group("progress", resolved=False)]},
        "limitations": [
            "Resolution actions and current team blockers were not independently verified."
        ],
    }


class GovernanceTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.template, self.source = (
            self.root / "approved.pptx",
            self.root / "source.governance.json",
        )
        make_template(self.template)
        self.payload = snapshot(self.template)
        self.save()

    def save(self):
        self.source.write_text(json.dumps(self.payload), encoding="utf-8")

    def build(self, **kwargs):
        self.save()
        return build_governance(self.source, self.template, **kwargs)

    def test_slide_edits_import_into_snapshot_and_survive_rebuild(self):
        team_group = group("resolved", 1)
        team_group["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Resolved assignee query",
            "entries": [{"id": "u1", "name": "Alex", "count": 1}],
        }
        team_group["resolution"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "reasons": {"ISSUE_FIXED": 1},
            "changes": [{"text": "1 VPC/VNet deleted", "count": 1}],
            "risks": ["UNPROTECTED_PRINCIPAL"],
        }
        self.payload["team_resolved"] = {
            "basis": "Resolved in the period and currently assigned to a team member in Wiz.",
            "team": [{"id": "u1", "name": "Alex"}],
            "total_groups": 1,
            "groups": [team_group],
        }
        deck = Presentation(self.build())
        self.assertEqual(shape(deck.slides[4], 7).table.cell(1, 2).text, "VPC/VNet deleted.")
        shape(deck.slides[4], 7).table.cell(1, 2).text = "Default VPCs removed."
        shape(deck.slides[7], 7).table.cell(1, 3).text = "Awaiting  CloudOps\nreply."
        edited = self.root / "edited.pptx"
        deck.save(edited)
        command = [
            "--snapshot",
            str(self.source),
            "--import-edits",
            str(edited),
            "--editor",
            "Alex",
        ]
        self.assertEqual(main(command), 0)
        payload = json.loads(self.source.read_text(encoding="utf-8"))
        self.assertEqual(
            [(e["id"], e["changed"], e["source"].split()[0]) for e in payload["closure_wording"]],
            [("resolved", "Default VPCs removed.", "Alex")],
        )
        self.assertEqual(
            [(e["id"], e["status"], "next_step" in e) for e in payload["rule_status"]],
            [("progress", "Awaiting CloudOps reply.", False)],
        )
        self.assertEqual(len(list(self.root.glob("source.governance.before-import-*.json"))), 1)
        before = self.source.read_bytes()
        self.assertEqual(main(command), 0)
        self.assertEqual(self.source.read_bytes(), before)
        rebuilt = Presentation(
            build_governance(self.source, self.template, output=self.root / "rebuilt.pptx")
        )
        self.assertEqual(shape(rebuilt.slides[4], 7).table.cell(1, 2).text, "Default VPCs removed.")
        self.assertEqual(
            shape(rebuilt.slides[4], 7).table.cell(1, 0).text,
            "Synthetic rule resolved\n(Alex) (1\u00a0resolved)",
        )
        self.assertEqual(
            shape(rebuilt.slides[7], 7).table.cell(1, 3).text, "Awaiting CloudOps reply."
        )

    def test_counts_dates_and_evidence_semantics(self):
        data = load_snapshot(self.source)
        self.assertEqual(data.total("open"), 10)
        self.assertEqual(data.total("resolved"), 2)
        self.assertTrue(data.preliminary)
        self.assertEqual(len(data.ranked("resolved_groups")), 1)
        row = table_row(data.ranked("resolved_groups")[0], True)
        self.assertIn("not verified", row[2])
        self.assertIn("needs evidence", row[3])
        self.assertTrue(
            table_row(group("work", resolved=False), False)[2].startswith("Apply least privilege")
        )
        malformed = group("ai", resolved=False)
        malformed["recommendation"] = (
            "###Identity management\n*Apply least privilege.\n*Restrict data access."
        )
        self.assertEqual(table_row(malformed, False)[2], "Apply least privilege.")
        malformed["recommendation"] = (
            "* Apply least-privilege principles to the model's service account"
        )
        self.assertEqual(
            table_row(malformed, False)[2],
            "Apply least-privilege principles to the model's service account.",
        )

    def test_uniform_body_font_and_assignee_labels_on_all_tables(self):
        assigned = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Full assignment census from Wiz",
            "entries": [
                {"id": "user-a", "name": "Alex", "count": 1},
                {"id": "user-b", "name": "Blair", "count": 1},
            ],
        }
        self.payload["resolved_groups"]["groups"][0]["assignments"] = deepcopy(assigned)
        assigned["entries"][1] = {"id": None, "name": "Unassigned", "count": 1}
        self.payload["in_progress_groups"]["groups"][0]["assignments"] = assigned
        output = self.build()
        deck = Presentation(output)
        self.assertEqual(
            "Synthetic rule resolved\n(2\u00a0findings)",
            shape(deck.slides[4], 7).table.cell(1, 0).text,
        )
        self.assertIn(
            "Synthetic rule progress\n(Alex)", shape(deck.slides[7], 7).table.cell(1, 0).text
        )
        for index in TABLE_CAPACITY:
            table = shape(deck.slides[index], 7).table
            for row in list(table.rows)[1:]:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            self.assertEqual(run.font.size.pt, TABLE_BODY_FONT_SIZE)
        self.assertEqual(TABLE_BODY_FONT_SIZE, 12)
        self.assertIn("ETA and blockers", shape(deck.slides[7], 7).table.cell(1, 3).text)
        self.assertNotIn("Assignee,", shape(deck.slides[7], 7).table.cell(1, 3).text)
        notes = json.loads(deck.slides[4].notes_slide.notes_text_frame.text)
        self.assertEqual(notes["groups"][0]["assignments"]["entries"][0]["name"], "Alex")

    def test_invalid_assignment_evidence_fails(self):
        good = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Verified full-group query",
            "entries": [{"id": "user-a", "name": "Alex", "count": 2}],
        }
        mutations = [
            lambda a: a["entries"][0].update(count=3),
            lambda a: a["entries"][0].update(count=True),
            lambda a: a["entries"][0].update(id=None),
            lambda a: a["entries"][0].pop("id"),
            lambda a: a["entries"][0].update(name="Alex/Blair"),
            lambda a: a["entries"].append(deepcopy(a["entries"][0])),
            lambda a: a.update(source=""),
            lambda a: a.update(verified_at="2999-01-01T00:00:00Z"),
            lambda a: a.update(entries=[]),
        ]
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                evidence = deepcopy(good)
                mutation(evidence)
                self.payload["resolved_groups"]["groups"][0]["assignments"] = evidence
                self.save()
                with self.assertRaises(DashboardError):
                    load_snapshot(self.source)
        row = table_row(group("no-evidence", resolved=False), False)
        self.assertIn("(Assignee not verified)", row[0])

    def test_partial_assignees_do_not_filter_status_counts(self):
        for section, resolved in (("resolved_groups", True), ("in_progress_groups", False)):
            with self.subTest(section=section):
                item = self.payload[section]["groups"][0]
                item["counts"]["MEDIUM"] = 100
                self.payload["resolved" if resolved else "in_progress"]["MEDIUM"] = 100
                self.payload["open"]["MEDIUM"] = 100
                item["assignments"] = {
                    "verified_at": "2025-10-13T11:00:00Z",
                    "source": "Same-status rule assignee query",
                    "entries": [{"id": "user-a", "name": "Alex", "count": 1}],
                }
                self.save()
                data = load_snapshot(self.source)
                row = table_row(item, resolved)
                if resolved:
                    self.assertNotIn("(Alex)", row[0])
                else:
                    self.assertIn("(Alex)", row[0])
                self.assertIn("(100\u00a0findings)", row[0])
                self.assertEqual(data.total("resolved" if resolved else "in_progress"), 100)
                item["assignments"]["entries"] = [{"id": None, "name": "Unassigned", "count": 1}]
                if not resolved:
                    self.assertIn("(Assignee not verified)", table_row(item, resolved)[0])
                item["assignments"]["entries"][0]["count"] = 100
                if not resolved:
                    self.assertIn("(Unassigned)", table_row(item, resolved)[0])
        output = self.build()
        deck = Presentation(output)
        for index in TABLE_CAPACITY:
            self.assertFalse(
                any(s.name == "Governance coverage" for s in deck.slides[index].shapes)
            )

    def test_team_resolved_shows_only_assigned_closures(self):
        self.payload["resolved_groups"]["groups"][0]["counts"]["MEDIUM"] = 2
        team_group = group("resolved", 1)
        team_group["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Resolved assignee query",
            "entries": [{"id": "u1", "name": "Alex", "count": 1}],
        }
        self.payload["team_resolved"] = {
            "basis": "Resolved in the period and currently assigned to a team member in Wiz.",
            "team": [{"id": "u1", "name": "Alex"}],
            "total_groups": 1,
            "groups": [team_group],
        }
        deck = Presentation(self.build())
        for index in (4, 5, 6):
            self.assertIn("Resolved by Team", shape(deck.slides[index], 4).text)
        row = shape(deck.slides[4], 7).table.cell(1, 0).text
        self.assertEqual(row, "Synthetic rule resolved\n(Alex) (1\u00a0resolved)")
        self.assertEqual(
            shape(deck.slides[4], 7).table.cell(1, 2).text, "Change details: team to confirm."
        )
        self.assertNotIn("Alex", shape(deck.slides[2], 20).text)
        self.payload["accepted"] = severity_counts(MEDIUM=3)
        accepted_only = group("accepted-only", 1)
        accepted_only["counts"] = severity_counts()
        accepted_only["accepted"] = severity_counts(MEDIUM=2)
        accepted_only["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Rejected assignee query",
            "entries": [{"id": "u1", "name": "Alex", "count": 2}],
        }
        team_group["accepted"] = severity_counts(MEDIUM=1)
        team_group["assignments"]["entries"][0]["count"] = 2
        self.payload["team_resolved"].update(total_groups=2, groups=[team_group, accepted_only])
        deck = Presentation(self.build(refresh=True))
        self.assertIn("Resolved / Risk Accepted by Team", shape(deck.slides[4], 4).text)
        table = shape(deck.slides[4], 7).table
        cells = {
            table.cell(r, 0).text.splitlines()[0]: (table.cell(r, 0).text, table.cell(r, 2).text)
            for r in range(1, 3)
        }
        self.assertEqual(
            cells["Synthetic rule resolved"],
            (
                "Synthetic rule resolved\n(Alex) (1\u00a0resolved, 1\u00a0accepted\u00a0risk)",
                "Change details: team to confirm.",
            ),
        )
        self.assertEqual(
            cells["Synthetic rule accepted-only"],
            (
                "Synthetic rule accepted-only\n(Alex) (2\u00a0accepted\u00a0risk)",
                "Change details: team to confirm.",
            ),
        )
        accepted_only["resolution"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "reasons": {"EXCEPTION": 2},
            "changes": [],
            "risks": ["UNPROTECTED_PRINCIPAL"],
        }
        deck = Presentation(self.build(refresh=True))
        table = shape(deck.slides[4], 7).table
        row = next(
            r
            for r in range(1, 3)
            if table.cell(r, 0).text.startswith("Synthetic rule accepted-only")
        )
        self.assertEqual(table.cell(row, 2).text, "Risk reviewed and accepted by the team.")
        self.assertEqual(table.cell(row, 3).text, "Accepted risk documented in Wiz.")
        team_group["resolution"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "reasons": {"ISSUE_FIXED": 2},
            "changes": [{"text": "1 VPC/VNet deleted", "count": 1}],
            "risks": ["UNPROTECTED_PRINCIPAL"],
        }
        self.payload["closure_wording"] = [
            {
                "id": "resolved",
                "changed": "Query logging enabled. Removed default VPCs.",
                "risk": "Improved DNS query visibility.",
                "source": "Syed",
                "as_of": "2025-10-14T06:30:00Z",
            }
        ]
        deck = Presentation(self.build(refresh=True))
        table = shape(deck.slides[4], 7).table
        row = next(
            r for r in range(1, 3) if table.cell(r, 0).text.startswith("Synthetic rule resolved")
        )
        self.assertEqual(table.cell(row, 2).text, "Query logging enabled. Removed default VPCs.")
        self.assertEqual(table.cell(row, 3).text, "Improved DNS query visibility.")
        self.payload["closure_wording"][0]["id"] = "progress"
        self.save()
        with self.assertRaisesRegex(DashboardError, "must name a team-resolved rule group"):
            load_snapshot(self.source)
        del self.payload["closure_wording"], team_group["resolution"]
        accepted_only["resolution"]["reasons"] = {"EXCEPTION": 1}
        self.save()
        with self.assertRaisesRegex(DashboardError, "count every team closure"):
            load_snapshot(self.source)
        del accepted_only["resolution"]
        self.assertIn(
            "5 findings were closed this period (2 remediated, 3 risk-accepted)",
            shape(deck.slides[2], 20).text,
        )
        self.payload["accepted"] = severity_counts(MEDIUM=1)
        self.save()
        with self.assertRaisesRegex(DashboardError, "accepted risk exceeds"):
            load_snapshot(self.source)
        self.payload["accepted"] = severity_counts(MEDIUM=3)
        for mutate, message in (
            (
                lambda t: t["groups"][0]["assignments"]["entries"].__setitem__(
                    0, {"id": None, "name": "Unassigned", "count": 1}
                ),
                "only list team members",
            ),
            (lambda t: t["groups"][0]["counts"].__setitem__("MEDIUM", 3), "exceed tenant closures"),
            (lambda t: t["groups"][0].__setitem__("id", "unknown"), "present in resolved_groups"),
        ):
            payload = deepcopy(self.payload)
            mutate(payload["team_resolved"])
            self.source.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaisesRegex(DashboardError, message):
                load_snapshot(self.source)

    def test_final_run_after_cutoff_drops_preliminary_labels(self):
        self.assertIn(
            "Scheduled period end; live snapshot",
            shape(Presentation(self.build()).slides[0], 7).text,
        )
        self.payload["period"].update(
            resolved_through="2025-10-15T23:59:59Z",
            collected_from="2025-10-16T18:00:00Z",
            collected_through="2025-10-16T18:30:00Z",
        )
        self.save()
        with self.assertRaisesRegex(DashboardError, "belong to the next report"):
            load_snapshot(self.source)
        self.payload["period"]["backlog_as_of"] = "2025-10-15T23:59:59Z"
        deck = Presentation(self.build(refresh=True))
        cover = shape(deck.slides[0], 7).text
        self.assertEqual(
            cover,
            "Data as of Oct 15, 2025 EOD\n"
            "Reporting period closed October 15, 2025 EOD; "
            "later changes roll into the next report.",
        )
        self.assertIn("Oct 15, 2025", shape(deck.slides[3], 12).text)
        self.assertNotIn("Coverage", shape(deck.slides[9], 8).text)
        for slide in deck.slides:
            texts = [s.text for s in slide.shapes if s.has_text_frame]
            self.assertFalse(
                any(
                    "PRELIMINARY" in t or "Scheduled period end" in t or "Oct 16" in t
                    for t in texts
                )
            )
        early = {
            "start_exclusive": "2025-09-15T23:59:59Z",
            "end_inclusive": "2025-10-15T23:59:59Z",
            "resolved_through": "2025-10-15T23:59:59Z",
            "backlog_as_of": "2025-10-15T23:59:59Z",
            "collected_from": "2025-10-14T16:00:00Z",
            "collected_through": "2025-10-14T16:30:00Z",
        }
        self.payload["period"] = dict(early)
        self.save()
        with self.assertRaisesRegex(DashboardError, "cannot extend beyond collection"):
            load_snapshot(self.source)
        self.payload["period"]["final_approval"] = {
            "approved_by": "Account owner",
            "approved_at": "2025-10-14T16:35:00Z",
            "reason": "Publish as final",
        }
        deck = Presentation(self.build(refresh=True))
        self.assertTrue(shape(deck.slides[0], 7).text.startswith("Data as of Oct 15, 2025 EOD"))
        self.assertEqual(
            json.loads(deck.slides[0].notes_slide.notes_text_frame.text)["final_approval"][
                "approved_by"
            ],
            "Account owner",
        )
        self.payload["period"]["collected_through"] = "2025-10-10T16:30:00Z"
        self.payload["period"]["collected_from"] = "2025-10-10T16:00:00Z"
        self.save()
        with self.assertRaisesRegex(DashboardError, "within 3 days"):
            load_snapshot(self.source)

    def test_cutoff_counts_ignore_everything_after_the_period_end(self):
        start, cutoff = datetime(2025, 9, 15, 23, 59, 59), datetime(2025, 10, 15, 23, 59, 59)
        rows = [
            {
                "Severity": "MEDIUM",
                "Status": "RESOLVED",
                "Created At": "2025-09-01",
                "Resolved Time": "2025-10-10",
            },
            {
                "Severity": "MEDIUM",
                "Status": "RESOLVED",
                "Created At": "2025-09-01",
                "Resolved Time": "2025-10-16",
            },
            {
                "Severity": "MEDIUM",
                "Status": "REJECTED",
                "Created At": "2025-09-01",
                "Resolved Time": "2025-10-01",
            },
            {
                "Severity": "MEDIUM",
                "Status": "RESOLVED",
                "Created At": "2025-08-01",
                "Resolved Time": "2025-09-10",
            },
            {"Severity": "LOW", "Status": "OPEN", "Created At": "2025-10-16", "Resolved Time": ""},
            {
                "Severity": "LOW",
                "Status": "IN_PROGRESS",
                "Created At": "2025-09-01",
                "Resolved Time": "",
                "Status Changed At": "2025-10-01",
            },
            {
                "Severity": "LOW",
                "Status": "IN_PROGRESS",
                "Created At": "2025-09-01",
                "Resolved Time": "",
                "Status Changed At": "2025-10-16",
            },
            {
                "Severity": "INFORMATIONAL",
                "Status": "OPEN",
                "Created At": "2025-09-01",
                "Resolved Time": "",
            },
        ]
        counts = cutoff_counts(rows, start, cutoff)
        self.assertEqual(counts["resolved"]["MEDIUM"], 1)
        self.assertEqual(counts["accepted"]["MEDIUM"], 1)
        self.assertEqual(counts["open"], {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 1, "LOW": 2})
        self.assertEqual(counts["in_progress"], {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 1})

    def test_key_messages_and_takeaways_are_client_facing(self):
        self.payload["hipaa"]["value"] = 96
        deck = Presentation(self.build())
        messages = shape(deck.slides[2], 20).text
        takeaways = shape(deck.slides[9], 8).text
        self.assertEqual(
            messages.splitlines(),
            [
                "Key messages",
                "Critical and High findings remain at zero; the Wiz Security Score held at 95% "
                "and HIPAA compliance rose to 96% (up 1 point).",
                "2 findings were closed this period (2 remediated), reducing the open backlog "
                "to 10, 1 fewer (9.1%) than September.",
                "Remediation is under way on 2 findings, led by Synthetic rule progress (2).",
                "Next-month focus: sustain zero Critical and High findings and keep reducing "
                "the Medium backlog (8 open).",
            ],
        )
        self.assertTrue(
            takeaways.startswith(
                "The Wiz Security Score held at 95%.\nCritical and High issues remain at zero.\n"
            )
        )
        self.assertIn(
            "2 findings closed this period (2 remediated); open backlog down 9.1% to 10.", takeaways
        )
        for text in (messages, takeaways):
            for internal in (
                "team",
                "Confirm",
                "tenant",
                "covered interval",
                "JSON",
                "Coverage",
                "C/H/M/L",
            ):
                self.assertNotIn(internal, text)

    def test_deferred_in_progress_rule_is_hidden_but_still_counted(self):
        extra = group("later", 3, resolved=False)
        self.payload["in_progress_groups"] = {
            "total_groups": 2,
            "groups": [self.payload["in_progress_groups"]["groups"][0], extra],
        }
        self.payload["in_progress"] = severity_counts(MEDIUM=5)
        self.payload["open"]["MEDIUM"] = 8
        self.payload["deferred_in_progress"] = [{"id": "later", "source": "Owner: next period"}]
        deck = Presentation(self.build())
        names = [shape(deck.slides[7], 7).table.cell(r, 0).text for r in range(1, 7)]
        self.assertFalse(any("Synthetic rule later" in n for n in names))
        self.assertTrue(any("Synthetic rule progress" in n for n in names))
        self.assertIn("Remediation is under way on 5 findings", shape(deck.slides[2], 20).text)
        self.payload["deferred_in_progress"] = [{"id": "unknown", "source": "x"}]
        self.save()
        with self.assertRaisesRegex(DashboardError, "must name an in-progress rule group"):
            load_snapshot(self.source)

    def test_next_step_is_a_complete_sentence(self):
        from wizexpert.governance.governance import first_sentence

        text = (
            "Use the \u201cleast privilege\u201d principle when assigning permissions, "
            "meaning that each account is assigned the exact permissions that it needs "
            "to function properly. When assigning permissions, avoid wild-card permissions."
        )
        self.assertEqual(
            first_sentence(text),
            "Use the \u201cleast privilege\u201d principle when assigning permissions, "
            "meaning that each account is assigned the exact permissions that it needs "
            "to function properly.",
        )
        self.assertEqual(first_sentence("Remove stale keys"), "Remove stale keys")
        self.assertTrue(first_sentence("word " * 60).endswith("..."))
        from wizexpert.governance.governance import next_step

        guardrail = (
            "### Require a guardrail in the IAM policy\n"
            "* This finding is on an IAM identity policy that allows `bedrock:InvokeModel` "
            "without a guardrail.\n"
            "* Add a `bedrock:GuardrailIdentifier` condition to the `Allow` statement "
            "so inference requests must carry a guardrail.\n"
        )
        self.assertEqual(
            next_step(guardrail),
            "Add a bedrock:GuardrailIdentifier condition to the Allow statement so "
            "inference requests must carry a guardrail.",
        )
        stale = (
            "### Protect highly privileged principals\n* " + text + "\n* Enable MFA for users.\n"
            "### Remove unnecessary access\n* Remove stale principals/keys.\n"
        )
        self.assertEqual(next_step(stale), "Remove stale principals/keys.")
        self.assertEqual(
            next_step("### Protect highly privileged principals\n* " + text), first_sentence(text)
        )

    def test_owner_status_replaces_pending_confirmation_for_that_owner(self):
        progress = self.payload["in_progress_groups"]["groups"][0]
        progress["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "In-progress checks",
            "entries": [{"id": "a", "name": "Alex", "count": 2}],
        }
        self.payload["owner_status"] = [
            {
                "name": "Alex",
                "status": "Pending review",
                "source": "Team update",
                "as_of": "2025-10-14T06:30:00Z",
            }
        ]
        self.payload["rule_status"] = [
            {
                "id": "progress",
                "status": "Awaiting permission-review outcome.",
                "next_step": "Add a private endpoint.",
                "source": "Team update",
            }
        ]
        row = shape(Presentation(self.build()).slides[7], 7).table.rows[1]
        self.assertEqual(
            (row.cells[2].text, row.cells[3].text),
            ("Add a private endpoint.", "Awaiting permission-review outcome."),
        )
        del self.payload["rule_status"]
        deck = Presentation(self.build(refresh=True))
        self.assertEqual(shape(deck.slides[7], 7).table.cell(1, 3).text, "Status: Pending review.")
        self.assertEqual(
            json.loads(deck.slides[7].notes_slide.notes_text_frame.text)["owner_status"][0]["name"],
            "Alex",
        )
        mixed = table_row(
            {
                **progress,
                "assignments": {
                    **progress["assignments"],
                    "entries": [
                        {"id": "a", "name": "Alex", "count": 1},
                        {"id": "b", "name": "Blair", "count": 1},
                    ],
                },
            },
            False,
            owner_status={"Alex": "Pending review"},
        )
        self.assertEqual(mixed[3], "Alex: Pending review.")
        self.payload["owner_status"].append(dict(self.payload["owner_status"][0]))
        self.save()
        with self.assertRaisesRegex(DashboardError, "Duplicate owner_status"):
            load_snapshot(self.source)

    def test_snapshot_without_informational_loads_and_slides_never_mention_it(self):
        def strip(value):
            if isinstance(value, dict):
                if set(value) == set(SEVERITIES):
                    value.pop("INFORMATIONAL")
                for item in value.values():
                    strip(item)
            elif isinstance(value, list):
                for item in value:
                    strip(item)

        self.payload["resolved_groups"]["groups"] = self.payload["resolved_groups"]["groups"][:1]
        self.payload["resolved_groups"]["total_groups"] = 1
        strip(self.payload)
        self.assertNotIn("INFORMATIONAL", json.dumps(self.payload))
        deck = Presentation(self.build())
        for slide in deck.slides:
            texts = [s.text for s in slide.shapes if s.has_text_frame]
            texts += [
                c.text
                for s in slide.shapes
                if s.has_table
                for row in s.table.rows
                for c in row.cells
            ]
            self.assertFalse(any("informational" in t.lower() for t in texts))

    def test_rule_assignees_never_credit_resolved_work_without_history(self):
        self.payload["assignment_scope"] = "RULE"
        resolved = self.payload["resolved_groups"]["groups"][0]
        progress = self.payload["in_progress_groups"]["groups"][0]
        progress["id"] = resolved["id"]
        resolved["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "Resolved checks",
            "entries": [{"id": None, "name": "Unassigned", "count": 2}],
        }
        progress["assignments"] = {
            "verified_at": "2025-10-13T11:00:00Z",
            "source": "In-progress checks",
            "entries": [{"id": "a", "name": "Alex", "count": 2}],
        }
        self.payload["rule_assignment_evidence"] = [
            {
                "id": resolved["id"],
                "status": "OPEN",
                "total_count": 10,
                "assignments": {
                    "verified_at": "2025-10-13T11:00:00Z",
                    "source": "Open finding assignee query",
                    "entries": [
                        {"id": "b", "name": "Blair", "count": 1},
                        {"id": "a", "name": "Alex", "count": 1},
                    ],
                },
            }
        ]
        output = self.build()
        deck = Presentation(output)
        for index in (4, 7):
            title = shape(deck.slides[index], 7).table.cell(1, 0).text
            if index == 7:
                self.assertIn("(Alex/Blair)", title)
            else:
                self.assertNotIn("(Alex/Blair)", title)
            self.assertIn("(2\u00a0findings)", title)
        self.assertEqual(
            json.loads(deck.slides[4].notes_slide.notes_text_frame.text)["rule_assignees"][
                resolved["id"]
            ],
            ["Alex", "Blair"],
        )
        self.assertNotIn("Assignee", table_row(resolved, True, [])[0])
        for index in (4, 5, 6):
            self.assertIn("Attribution Pending", shape(deck.slides[index], 4).text)
        self.assertNotIn("Alex", shape(deck.slides[2], 20).text)
        self.assertIn(
            "UNVERIFIED",
            json.loads(deck.slides[4].notes_slide.notes_text_frame.text)["resolved_attribution"],
        )
        self.payload["rule_assignment_evidence"][0]["status"] = "REJECTED"
        self.save()
        with self.assertRaises(DashboardError):
            load_snapshot(self.source)

    def test_generated_deck_no_stale_text_and_sources_unchanged(self):
        original = self.template.read_bytes(), self.source.read_bytes()
        output = self.build()
        deck = Presentation(output)
        self.assertEqual(len(deck.slides), 10)
        for index, slide in enumerate(deck.slides):
            self.assertEqual(shape(slide, 2).text, "")
            texts = [s.text for s in slide.shapes if s.has_text_frame]
            texts += [
                c.text
                for s in slide.shapes
                if s.has_table
                for row in s.table.rows
                for c in row.cells
            ]
            # Header/customer labels are untouched, all narrative placeholders are overwritten.
            dynamic_ids = set(TEXT_IDS[index]) - {4, 5}
            self.assertFalse(any("Stale" in shape(slide, sid).text for sid in dynamic_ids))
            self.assertNotIn("Stale extra table", texts)
            self.assertIn("snapshot_sha256", slide.notes_slide.notes_text_frame.text)
        self.assertIn("PRELIMINARY", shape(deck.slides[0], 7).text)
        self.assertIn("1 fewer", shape(deck.slides[2], 20).text)
        self.assertIn("1 more", shape(deck.slides[3], 26).text)
        trend = shape(deck.slides[3], 25).text
        self.assertTrue(trend.startswith("Trend statement"))
        self.assertIn("Critical and High remain at zero", trend)
        self.assertIn(
            "Focus: Medium remains the largest volume (8; 80% of backlog), 2 in progress.", trend
        )
        self.assertIn(
            "Backlog 1 fewer (9.1%) since September; Medium and Low still call for steady", trend
        )
        self.assertEqual(len(shape(deck.slides[3], 25).text_frame.paragraphs), 4)
        self.assertEqual(
            shape(deck.slides[2], 9).text_frame.paragraphs[0].runs[0].font.name, "Arial"
        )
        self.assertEqual(self.template.read_bytes(), original[0])
        self.assertEqual(self.source.read_bytes(), original[1])
        self.assertEqual(output.with_suffix(".governance.json").read_bytes(), original[1])

    def test_missing_or_bad_counts_and_incomplete_groups_fail(self):
        original = deepcopy(self.payload)
        mutations = [
            lambda p: p["open"].pop("LOW"),
            lambda p: p["open"].update(MEDIUM=True),
            lambda p: p["open"].update(MEDIUM=-1),
            lambda p: p["in_progress"].update(MEDIUM=50),
            lambda p: p["resolved_groups"].update(total_groups=50),
            lambda p: p["resolved"].update(MEDIUM=3),
            lambda p: p["scope"].update(issue_types=["TOXIC_COMBINATION"]),
            lambda p: p["resolved_groups"]["groups"].append(
                deepcopy(p["resolved_groups"]["groups"][0])
            ),
        ]
        for mutate in mutations:
            self.payload = deepcopy(original)
            mutate(self.payload)
            self.save()
            with self.assertRaises(DashboardError):
                load_snapshot(self.source)

    def test_period_boundaries_and_timestamp_validation(self):
        original = deepcopy(self.payload)
        for value in (
            "2025-09-15T23:59:59Z",
            "2025-10-13T09:00:01Z",
            "2025-10-01T00:00:00",
            "2025-10-01T00:00:00+05:30",
        ):
            self.payload = deepcopy(original)
            self.payload["resolved_groups"]["groups"][0]["samples"][0]["resolved_at"] = value
            self.save()
            with self.assertRaises(DashboardError):
                load_snapshot(self.source)
        self.payload = original
        self.payload["resolved_groups"]["groups"][0]["samples"][0]["resolved_at"] = (
            "2025-10-13T09:00:00Z"
        )
        self.save()
        load_snapshot(self.source)

    def test_unavailable_compliance_explicit_not_zero(self):
        self.payload["hitrust"] = {
            "value": None,
            "source": "Wiz query",
            "unavailable_reason": "HTTP 504",
        }
        output = self.build()
        self.assertIn("HITRUST Unavailable", shape(Presentation(output).slides[2], 18).text)
        self.payload["hitrust"].pop("unavailable_reason")
        self.save()
        with self.assertRaises(DashboardError):
            load_snapshot(self.source)
        self.payload["hitrust"] = {"source": "Wiz query", "unavailable_reason": "HTTP 504"}
        self.save()
        with self.assertRaises(DashboardError):
            load_snapshot(self.source)

    def test_annotations_require_captured_evidence_and_do_not_extrapolate(self):
        item = self.payload["resolved_groups"]["groups"][0]
        item["annotation"] = {
            "change": "Previously accepted; lifecycle closure only.",
            "risk": "Not new remediation.",
            "source": "Approved prior review",
            "issue_ids": ["issue-resolved"],
        }
        self.save()
        data = load_snapshot(self.source)
        self.assertIn("covers 1 of 2", table_row(data.ranked("resolved_groups")[0], True)[2])
        item["annotation"]["issue_ids"] = ["unrelated"]
        self.save()
        with self.assertRaises(DashboardError):
            load_snapshot(self.source)

    def test_highlight_capacity_and_full_evidence_retention(self):
        groups = [group(f"r-{i}", 1) for i in range(20)]
        self.payload["resolved_groups"] = {"total_groups": 20, "groups": groups}
        self.payload["resolved"] = severity_counts(MEDIUM=20)
        output = self.build()
        slide = Presentation(output).slides[4]
        self.assertFalse(any(s.name == "Governance coverage" for s in slide.shapes))
        self.assertEqual(sum(1 for r in list(shape(slide, 7).table.rows)[1:] if r.cells[0].text), 4)
        notes = json.loads(slide.notes_slide.notes_text_frame.text)
        self.assertEqual(len(notes["groups"]), 20)

    def test_refresh_backups_and_template_hash_binding(self):
        output = self.build()
        old = output.read_bytes()
        with self.assertRaises(DashboardError):
            self.build()
        self.payload["open"]["LOW"] = 3
        self.build(refresh=True)
        backups = list((self.root / "report-backups").glob(f"governance-*\\{output.name}"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), old)
        self.payload["prior"]["source_sha256"] = "0" * 64
        with self.assertRaises(DashboardError):
            self.build(refresh=True)

    def test_bad_template_and_source_output_collision(self):
        with self.assertRaises(DashboardError):
            self.build(output=self.template)
        deck = Presentation(self.template)
        shape(deck.slides[0], 5).text = "Wrong template"
        deck.save(self.template)
        self.payload["prior"]["source_sha256"] = hashlib.sha256(
            self.template.read_bytes()
        ).hexdigest()
        with self.assertRaises(DashboardError):
            self.build()

    def test_next_month_reuses_generated_template_without_coverage_captions(self):
        october = self.build()
        # A deck generated before captions were dropped still carries one; it must not survive.
        legacy = Presentation(october)
        legacy.slides[4].shapes.add_textbox(
            Inches(0.45), Inches(6.55), Inches(12), Inches(0.38)
        ).name = "Governance coverage"
        legacy.save(october)
        self.payload["prior"].update(
            date="2025-10-15",
            source=october.name,
            source_sha256=hashlib.sha256(october.read_bytes()).hexdigest(),
            open={s: self.payload["open"][s] for s in SEVERITIES[:4]},
        )
        self.payload["period"].update(
            start_exclusive="2025-10-15T23:59:59Z",
            end_inclusive="2025-11-15T23:59:59Z",
            resolved_through="2025-11-13T09:00:00Z",
            collected_from="2025-11-13T09:00:00Z",
            collected_through="2025-11-13T10:00:00Z",
        )
        for item in self.payload["resolved_groups"]["groups"]:
            item["samples"][0]["resolved_at"] = "2025-11-01T00:00:00Z"
        self.save()
        november = build_governance(self.source, october)
        self.assertNotEqual(november, october)
        deck = Presentation(november)
        self.assertIn("November 15, 2025", shape(deck.slides[0], 6).text)
        for index in TABLE_CAPACITY:
            self.assertFalse(
                any(s.name == "Governance coverage" for s in deck.slides[index].shapes)
            )

    def test_publication_failure_preserves_previous_output(self):
        output = self.build()
        old = output.read_bytes()
        with patch(
            "wizexpert.governance.governance.publish_pair", side_effect=OSError("File locked")
        ), self.assertRaises(OSError):
            self.build(refresh=True)
        self.assertEqual(output.read_bytes(), old)
        self.assertFalse(list(self.root.glob(".wiz-monthly-*.lock")))

    def test_cli_errors_are_visible(self):
        with self.assertRaises(SystemExit) as result:
            main(["--snapshot", str(self.root / "missing.json"), "--template", str(self.template)])
        self.assertEqual(result.exception.code, 1)

    def test_cli_reports_assignee_actions_and_validates_gaps(self):
        self.payload["assignment_gaps"] = [
            {
                "issue_id": "issue-1",
                "rule_id": "r1",
                "rule_name": "Synthetic rule",
                "status": "IN_PROGRESS",
                "suggested_owner": "Alex",
                "suggested_owner_id": "u1",
                "basis": "ISSUE_IN_PROGRESS by " + mail("alex") + "",
                "event_at": "2025-06-09T00:00:00Z",
            }
        ]
        self.save()
        actions = self.root / "actions.csv"
        with patch("builtins.print") as printed:
            self.assertEqual(
                main(
                    [
                        "--snapshot",
                        str(self.source),
                        "--template",
                        str(self.template),
                        "--actions-csv",
                        str(actions),
                    ]
                ),
                0,
            )
        output = "\n".join(str(call.args[0]) for call in printed.call_args_list)
        self.assertIn("Wiz assignee actions: 1", output)
        self.assertIn("Assign Alex to issue-1 [IN_PROGRESS] Synthetic rule", output)
        self.assertIn(
            "issue-1,r1,Synthetic rule,IN_PROGRESS,Alex,u1", actions.read_text(encoding="utf-8")
        )
        self.payload["assignment_gaps"][0]["suggested_owner_id"] = None
        self.save()
        with self.assertRaisesRegex(DashboardError, "given together"):
            load_snapshot(self.source)

    def test_capture_minimizes_resources_and_retains_aggregates(self):
        response = {
            "issuesGroupedByValue": {
                "totalCount": 1,
                "nodes": [
                    {
                        "id": "rule",
                        "issues": {
                            **{
                                key: (3 if severity == "MEDIUM" else 0)
                                for severity, key in AGGREGATE_KEYS.items()
                            },
                            "totalCount": 3,
                            "nodes": [
                                {
                                    "id": "issue",
                                    "createdAt": "2025-01-01T00:00:00Z",
                                    "resolvedAt": None,
                                    "entitySnapshot": {"name": "sensitive-resource-name"},
                                    "sourceRules": [
                                        {
                                            "name": "Synthetic rule",
                                            "resolutionRecommendation": "* Fix permissions.",
                                        }
                                    ],
                                }
                            ],
                        },
                    }
                ],
            }
        }
        captured = capture_groups(response)
        self.assertEqual(captured["groups"][0]["counts"]["MEDIUM"], 3)
        self.assertEqual(len(captured["groups"][0]["samples"]), 1)
        self.assertNotIn("sensitive-resource-name", json.dumps(captured))


if __name__ == "__main__":
    unittest.main()
