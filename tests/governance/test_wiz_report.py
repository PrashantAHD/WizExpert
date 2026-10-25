import unittest
from datetime import datetime

from openpyxl import Workbook

from wizexpert.governance.wiz_report import (
    build_monthly_scope_tab,
    build_severity_tab,
    build_summary_tab,
    compute_metrics,
    default_snapshots,
    merge_issue_rows,
    read_issue_rows_from_existing_tab,
)

SNAPSHOTS = [
    ("Apr 15, 2026", datetime(2026, 4, 15, 23, 59, 59)),
    ("Jul 15, 2026", datetime(2026, 7, 15, 23, 59, 59)),
    ("Aug 15, 2026", datetime(2026, 8, 15, 23, 59, 59)),
]


def issue(issue_id, status, created, resolved=""):
    return {
        "Issue ID": issue_id,
        "Title": f"Issue {issue_id}",
        "Severity": "HIGH",
        "Status": status,
        "Created At": created,
        "Resolved Time": resolved,
    }


class GovernanceReportTests(unittest.TestCase):
    def test_september_defaults_use_fixed_march_baseline(self):
        self.assertEqual(
            [label for label, _ in default_snapshots(datetime(2026, 9, 21))],
            ["Mar 15, 2026", "Aug 15, 2026", "Sep 15, 2026"],
        )

    def test_baseline_does_not_roll_forward_next_year(self):
        self.assertEqual(
            [label for label, _ in default_snapshots(datetime(2027, 1, 21))],
            ["Mar 15, 2026", "Dec 15, 2026", "Jan 15, 2027"],
        )
        self.assertEqual(
            [label for label, _ in default_snapshots(datetime(2027, 4, 21))],
            ["Mar 15, 2026", "Mar 15, 2027", "Apr 15, 2027"],
        )

    def test_baseline_month_does_not_duplicate_or_include_pre_go_live_dates(self):
        self.assertEqual(
            [label for label, _ in default_snapshots(datetime(2026, 3, 21))],
            ["Mar 15, 2026"],
        )
        with self.assertRaisesRegex(ValueError, "precedes"):
            default_snapshots(datetime(2026, 2, 21))

    def setUp(self):
        self.rows = [
            issue("baseline-closed", "RESOLVED", "2026-04-01", "2026-06-01 10:00"),
            issue("baseline-open", "OPEN", "2026-04-01"),
            issue("new-closed", "REJECTED", "2026-08-01", "2026-08-10 09:30"),
            issue("new-open", "OPEN", "2026-08-01"),
        ]

    def test_monthly_scope_metrics_match_august_report_logic(self):
        metrics = compute_metrics(self.rows, SNAPSHOTS)
        self.assertEqual(metrics["total"], [2, 2, 4])
        self.assertEqual(metrics["open"], [2, 1, 2])
        self.assertEqual(metrics["resolved"], [0, 1, 1])
        self.assertEqual(metrics["ignored"], [0, 0, 1])

        workbook = Workbook()
        del workbook["Sheet"]
        build_severity_tab(workbook, "HIGH", self.rows, SNAPSHOTS)
        sheet = workbook["HIGH Issues"]

        self.assertEqual(sheet["G5"].value, "New in month: +2")
        self.assertEqual(sheet["G6"].value, "Scope: 3 | Closed: 1 | Remaining: 2")
        self.assertEqual(sheet["G9"].value, "Closure rate: 33.3%")
        self.assertEqual(
            sheet["A12"].value,
            "Monthly flow: Start Open (Jul 15, 2026) 1 + New +2"
            " - Closed 1 = End Open (Aug 15, 2026) 2",
        )

    def test_summary_uses_open_backlog_progress_and_dynamic_trend(self):
        metrics = compute_metrics(self.rows, SNAPSHOTS)
        workbook = Workbook()
        build_summary_tab(workbook, {"HIGH": metrics}, SNAPSHOTS)
        sheet = workbook["Summary"]

        self.assertEqual(
            sheet["A1"].value,
            "Example Health Cloud Security Governance Report - August 2026",
        )
        self.assertEqual(sheet["G5"].value, "+0.0%")
        self.assertEqual(sheet["H5"].value, "→ Open issues unchanged")

    def test_monthly_scope_view_matches_management_calculations(self):
        metrics = compute_metrics(self.rows, SNAPSHOTS)
        workbook = Workbook()
        build_monthly_scope_tab(workbook, {"HIGH": metrics}, SNAPSHOTS)
        sheet = workbook["Monthly Scope View"]

        self.assertEqual(
            [sheet.cell(4, column).value for column in range(2, 11)],
            [2, 4, 1, 2, "100.0%", 3, 1, 2, "33.3%"],
        )

    def test_baseline_progress_excludes_prior_resolutions_and_ignored_issues(self):
        rows = [
            issue("prior-resolved", "RESOLVED", "2026-04-01", "2026-04-10"),
            issue("prior-ignored", "REJECTED", "2026-04-01", "2026-04-11"),
            issue("closed-later", "RESOLVED", "2026-04-01", "2026-08-01"),
            issue("still-open", "OPEN", "2026-04-01"),
        ]
        metrics = compute_metrics(rows, SNAPSHOTS)
        workbook = Workbook()
        severities = ("CRITICAL", "HIGH", "MEDIUM", "LOW")
        build_monthly_scope_tab(workbook, {sev: metrics for sev in severities}, SNAPSHOTS)
        sheet = workbook["Monthly Scope View"]
        for row, severity in enumerate(severities, 4):
            with self.subTest(severity=severity):
                self.assertEqual(sheet.cell(row, 2).value, 2)
                self.assertEqual(sheet.cell(row, 3).value, 4)
                self.assertEqual(sheet.cell(row, 5).value, 3)
                self.assertEqual(sheet.cell(row, 6).value, "50.0%")
        self.assertEqual(metrics["total"][0], 4)
        self.assertEqual(metrics["resolved"][0], 1)
        self.assertEqual(metrics["ignored"][0], 1)

    def test_high_baseline_is_open_backlog_without_deleting_prebaseline_resolution(self):
        snapshots = default_snapshots(datetime(2026, 9, 21))
        metrics = {
            "labels": [label for label, _ in snapshots],
            "total": [97, 106, 107],
            "open": [96, 4, 2],
            "resolved": [1, 78, 81],
            "ignored": [0, 24, 24],
            "pct": ["1.0%", "96.2%", "98.1%"],
        }
        workbook = Workbook()
        build_monthly_scope_tab(workbook, {"HIGH": metrics}, snapshots)
        sheet = workbook["Monthly Scope View"]
        self.assertEqual(sheet["B4"].value, 96)
        self.assertEqual(sheet["E4"].value, 105)
        self.assertEqual(sheet["F4"].value, "108.3%")
        self.assertIn("Open Backlog", sheet["B3"].value)
        self.assertEqual(metrics["total"][0], 97)
        self.assertEqual(metrics["resolved"][0], 1)

    def test_append_carries_forward_closed_issues_missing_from_latest_export(self):
        historical = [issue("old", "RESOLVED", "2026-04-01", "2026-06-01 10:00")]
        current = [issue("current", "OPEN", "2026-08-01")]
        merged, supplemented = merge_issue_rows(historical, current)

        self.assertEqual({row["Issue ID"] for row in merged}, {"old", "current"})
        self.assertEqual([row["Issue ID"] for row in supplemented], ["old"])

        workbook = Workbook()
        del workbook["Sheet"]
        build_severity_tab(workbook, "HIGH", merged, SNAPSHOTS, supplemented)
        restored = read_issue_rows_from_existing_tab(workbook, "HIGH")
        self.assertEqual(restored[0]["Issue ID"], "old")
        self.assertEqual(restored[0]["Resolved Time"], "2026-06-01 10:00")


if __name__ == "__main__":
    unittest.main()
