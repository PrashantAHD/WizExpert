import csv
import json
import os
import tempfile
import unittest
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import ImageFont
from pptx import Presentation
from test_dashboard import history_payload, make_template

from wizexpert.governance.dashboard import shape, template_history
from wizexpert.governance.dashboard_data import DashboardError, read_workbook
from wizexpert.governance.monthly import build_monthly_workbook, reporting_month, run_monthly
from wizexpert.governance.monthly_files import digest, monthly_lock, publish_pair
from wizexpert.governance.wiz_report import REQUIRED_COLUMNS


def synthetic_rows():
    result = []
    titles = (
        "Publicly readable bucket contains sensitive data",
        "VM infected with a hacking tool",
        "Third party service account with high privileges",
    )
    for severity, size in (("CRITICAL", 6), ("HIGH", 96), ("MEDIUM", 1), ("LOW", 1)):
        for index in range(size):
            opened = (severity == "HIGH" and index == 95) or severity in {"MEDIUM", "LOW"}
            result.append(
                {
                    "Issue ID": f"{severity}-{index}",
                    "Title": titles[index % len(titles)],
                    "Severity": severity,
                    "Status": "OPEN" if opened else "RESOLVED",
                    "Created At": "2026-01-01",
                    "Resolved Time": ""
                    if opened
                    else "2026-08-20 10:00"
                    if severity == "CRITICAL" and index == 5
                    else "2026-08-10 10:00",
                }
            )
    return result


def save_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=sorted(REQUIRED_COLUMNS))
        writer.writeheader()
        writer.writerows(rows)


class MonthlyWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.template = self.root / "approved.pptx"
        make_template(self.template)
        self.history = self.root / "approved.dashboard-history.json"
        self.history.write_text(json.dumps(history_payload(self.template)), encoding="utf-8")
        self.csv = self.root / "exports.csv"
        self.rows = synthetic_rows()
        save_csv(self.csv, self.rows)
        self.previous = self.root / "approved-august.xlsx"
        build_monthly_workbook(
            [self.csv], [], self.previous, reporting_month("2026-08"), "Example Provider"
        )
        font = ImageFont.load_default(size=24)
        for mock in (
            patch(
                "wizexpert.governance.dashboard.chart_fonts",
                return_value=(Path("regular"), Path("bold")),
            ),
            patch("wizexpert.governance.dashboard_charts.ImageFont.truetype", return_value=font),
        ):
            mock.start()
            self.addCleanup(mock.stop)

    def run_september(self, **kwargs):
        return run_monthly(
            "2026-09",
            self.template,
            self.root,
            history=self.history,
            inputs=[self.csv],
            previous_report=self.previous,
            **kwargs,
        )

    def close_remaining_high(self, when="2026-09-10"):
        row = next(row for row in self.rows if row["Issue ID"] == "HIGH-95")
        row.update(Status="RESOLVED", **{"Resolved Time": when})
        save_csv(self.csv, self.rows)

    def test_first_run_then_same_month_refresh_and_verified_backups(self):
        first = self.run_september()
        self.assertEqual(first.workbook.name, "Sep26_ExampleHealth_Governance_Report.xlsx")
        self.assertEqual(
            first.deck.name, "ExampleHealth_Remediation_Executive_Dashboard_9.15.26.pptx"
        )
        self.assertEqual(
            read_workbook(first.workbook).open_at(read_workbook(first.workbook).latest), 1
        )
        self.assertEqual(shape(Presentation(first.deck).slides[2], 9).text, "1")
        before_excel, before_deck = first.workbook.read_bytes(), first.deck.read_bytes()
        original_template = self.template.read_bytes()
        self.close_remaining_high()
        refreshed = self.run_september()
        self.assertEqual((first.workbook, first.deck), (refreshed.workbook, refreshed.deck))
        self.assertNotEqual(first.archive, refreshed.archive)
        self.assertEqual((refreshed.archive / first.workbook.name).read_bytes(), before_excel)
        self.assertEqual((refreshed.archive / first.deck.name).read_bytes(), before_deck)
        self.assertEqual(shape(Presentation(refreshed.deck).slides[2], 9).text, "0")
        self.assertEqual(
            read_workbook(refreshed.workbook).open_at(read_workbook(refreshed.workbook).latest), 0
        )
        self.assertEqual(original_template, self.template.read_bytes())
        self.assertEqual(
            json.loads((refreshed.archive / "pair.monthly-run.json").read_text())["status"],
            "complete",
        )
        self.assertFalse((self.root / ".wiz-monthly-2026-09.lock").exists())
        self.assertFalse(list(self.root.glob(".wiz-monthly-2026-09-*")))

    def test_preliminary_workbook_uses_actual_collection_and_preserves_prior(self):
        original = self.previous.read_bytes()
        target = self.root / "preliminary-september.xlsx"
        collected = datetime(2026, 9, 13, 12, 30, tzinfo=UTC)
        build_monthly_workbook(
            [self.csv],
            [self.previous],
            target,
            reporting_month("2026-09"),
            "Example Provider",
            collected_at=collected,
        )
        from wizexpert.governance.dashboard_data import read_snapshots

        wb = load_workbook(target)
        prior = load_workbook(self.previous)
        try:
            for sheet in wb:
                self.assertIn("PRELIMINARY", sheet["A2"].value)
                self.assertIn("2026-09-13 12:30:00 UTC", sheet["A2"].value)
                self.assertIn("Scheduled cutoff 2026-09-15 EOD", sheet["A2"].value)
                self.assertIn("PRELIMINARY", sheet.oddHeader.center.text)
                if sheet.title.endswith(" Issues"):
                    approved = read_snapshots(prior[sheet.title])
                    self.assertEqual(sheet.cell(4, 4).value.splitlines()[0], "Sep 13, 2026")
                    for col in (2, 3):
                        day = datetime.strptime(
                            sheet.cell(4, col).value.splitlines()[0], "%b %d, %Y"
                        ).date()
                        self.assertEqual(
                            [sheet.cell(row, col).value for row in range(5, 9)],
                            [
                                getattr(approved[day], key)
                                for key in ("total", "open", "resolved", "ignored")
                            ],
                        )
        finally:
            wb.close()
            prior.close()
        self.assertEqual(self.previous.read_bytes(), original)
        with self.assertRaisesRegex(DashboardError, "Preliminary collection"):
            read_workbook(target)
        build_monthly_workbook(
            [self.csv],
            [self.previous],
            target,
            reporting_month("2026-09"),
            "Example Provider",
            collected_at=datetime(2026, 9, 15, 12, 30, tzinfo=UTC),
        )
        with self.assertRaisesRegex(DashboardError, "Preliminary collection"):
            read_workbook(target)

    def test_preliminary_dashboard_is_explicit_labelled_and_not_reusable_as_template(self):
        from wizexpert.governance.dashboard import build_deck

        workbook = self.root / "preliminary.xlsx"
        build_monthly_workbook(
            [self.csv],
            [self.previous],
            workbook,
            reporting_month("2026-09"),
            "Example Provider",
            collected_at=datetime(2026, 9, 13, 12, 30, tzinfo=UTC),
        )
        with self.assertRaisesRegex(DashboardError, "Preliminary collection"):
            build_deck(workbook, self.template, history_path=self.history)
        deck_path = build_deck(
            workbook, self.template, history_path=self.history, allow_preliminary=True
        )
        self.assertEqual(
            deck_path.name, "ExampleHealth_Remediation_Executive_Dashboard_9.13.26_PRELIMINARY.pptx"
        )
        deck = Presentation(deck_path)
        self.assertEqual(shape(deck.slides[0], 2).text, "September 13th, 2026 (Preliminary)")
        self.assertIn("Preliminary", deck.core_properties.title)
        history = template_history(deck, None, allow_preliminary=True)
        self.assertEqual(max(history).isoformat(), "2026-09-13")
        self.assertTrue(all(day.day == 15 for day in history if day != max(history)))
        with self.assertRaisesRegex(
            DashboardError, "preliminary dashboard cannot be used as a template"
        ):
            template_history(deck, None)

    def test_kpi_splits_remaining_into_in_progress_and_open(self):
        next(r for r in self.rows if r["Issue ID"] == "LOW-0")["Status"] = "IN_PROGRESS"
        save_csv(self.csv, self.rows)
        target = self.root / "kpi.xlsx"
        build_monthly_workbook(
            [self.csv], [self.previous], target, reporting_month("2026-09"), "Example Provider"
        )
        wb = load_workbook(target)
        try:
            kpi = {
                sev: next(
                    c.value
                    for c in wb[f"{sev} Issues"]["A"]
                    if str(c.value or "").startswith("KPI:")
                )
                for sev in ("HIGH", "MEDIUM", "LOW")
            }
        finally:
            wb.close()
        self.assertTrue(kpi["LOW"].endswith("| Remaining = 1 (In progress = 1 | Open = 0)"))
        self.assertTrue(kpi["MEDIUM"].endswith("| Remaining = 1 (In progress = 0 | Open = 1)"))
        self.assertTrue(kpi["HIGH"].endswith("(In progress = 0 | Open = 1)"))

    def test_changes_after_cutoff_are_left_for_next_report(self):
        fields = sorted(REQUIRED_COLUMNS) + ["Status Changed At"]
        rows = [dict(r, **{"Status Changed At": ""}) for r in self.rows]
        by_id = {r["Issue ID"]: r for r in rows}
        by_id["LOW-0"].update(Status="RESOLVED", **{"Resolved Time": "2026-09-16 10:00"})
        by_id["MEDIUM-0"].update(Status="IN_PROGRESS", **{"Status Changed At": "2026-09-16 09:00"})
        by_id["HIGH-95"].update(Status="IN_PROGRESS", **{"Status Changed At": "2026-09-10 09:00"})
        rows.append(
            {
                **by_id["LOW-0"],
                "Issue ID": "LOW-NEW",
                "Status": "OPEN",
                "Created At": "2026-09-16",
                "Resolved Time": "",
                "Status Changed At": "",
            }
        )
        with self.csv.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        target = self.root / "after-cutoff.xlsx"
        build_monthly_workbook(
            [self.csv], [self.previous], target, reporting_month("2026-09"), "Example Provider"
        )
        wb = load_workbook(target)
        try:
            low, medium, high = (wb[f"{s} Issues"] for s in ("LOW", "MEDIUM", "HIGH"))
            self.assertEqual([low.cell(r, 4).value for r in range(5, 9)], [1, 1, 0, 0])
            ids = [c.value for c in low["A"]]
            self.assertNotIn("LOW-0", ids)
            self.assertNotIn("LOW-NEW", ids)
            kpi = {
                name: next(c.value for c in ws["A"] if str(c.value or "").startswith("KPI:"))
                for name, ws in (("LOW", low), ("MEDIUM", medium), ("HIGH", high))
            }
            footer = next(c.value for c in low["A"] if str(c.value or "").startswith("Total:"))
        finally:
            wb.close()
        self.assertTrue(kpi["LOW"].endswith("Remaining = 1 (In progress = 0 | Open = 1)"))
        self.assertTrue(kpi["MEDIUM"].endswith("(In progress = 0 | Open = 1)"))
        self.assertTrue(kpi["HIGH"].endswith("(In progress = 1 | Open = 0)"))
        self.assertTrue(
            footer.startswith("Total: 0  |  Resolved: 0  |  Ignored: 0  |  1 issue(s) still Open")
        )

    def test_detail_columns_have_no_owner_and_report_snapshot_month(self):
        september = self.root / "detail-september.xlsx"
        build_monthly_workbook(
            [self.csv], [self.previous], september, reporting_month("2026-09"), "Example Provider"
        )
        october = self.root / "detail-october.xlsx"
        build_monthly_workbook(
            [self.csv], [september], october, reporting_month("2026-10"), "Example Provider"
        )
        for path in (september, october):
            wb = load_workbook(path)
            try:
                ws = wb["HIGH Issues"]
                header = next(
                    r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Issue ID"
                )
                self.assertEqual(
                    [ws.cell(header, c).value for c in range(1, 9)],
                    [
                        "Issue ID",
                        "Title",
                        "Severity",
                        "Status",
                        "Created At",
                        "Resolved Time",
                        "Snapshot Month",
                        None,
                    ],
                )
            finally:
                wb.close()
        wb = load_workbook(september)
        try:
            ws = wb["CRITICAL Issues"]
            header = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Issue ID")
            months = {
                ws.cell(r, 1).value: ws.cell(r, 7).value
                for r in range(header + 1, ws.max_row + 1)
                if str(ws.cell(r, 1).value or "").startswith("CRITICAL-")
            }
        finally:
            wb.close()
        # Resolved Aug 10 -> August report; Aug 20 (after the Aug 15 cutoff) -> September report.
        self.assertEqual(months["CRITICAL-0"], "August 2026")
        self.assertEqual(months["CRITICAL-5"], "September 2026")
        self.assertEqual(read_workbook(september).latest.isoformat(), "2026-09-15")

    def test_invalid_preliminary_collection_cannot_publish(self):
        target = self.root / "invalid-preliminary.xlsx"
        for when in (
            datetime(2026, 9, 13),
            datetime(2026, 8, 13, tzinfo=UTC),
            datetime(2026, 9, 16, tzinfo=UTC),
            datetime(2999, 9, 13, tzinfo=UTC),
        ):
            with self.subTest(when=when), self.assertRaises(DashboardError):
                build_monthly_workbook(
                    [self.csv],
                    [self.previous],
                    target,
                    reporting_month("2026-09"),
                    "Example Provider",
                    collected_at=when,
                )
            self.assertFalse(target.exists())
        self.rows[0]["Resolved Time"] = "2026-09-13 12:31"
        save_csv(self.csv, self.rows)
        with self.assertRaisesRegex(DashboardError, "after the collection timestamp"):
            build_monthly_workbook(
                [self.csv],
                [self.previous],
                target,
                reporting_month("2026-09"),
                "Example Provider",
                collected_at=datetime(2026, 9, 13, 12, 30, tzinfo=UTC),
            )
        self.assertFalse(target.exists())

    def test_october_rollover_reuses_september_history_and_retained_findings(self):
        september = self.run_september()
        before = (september.workbook.read_bytes(), september.deck.read_bytes())
        self.close_remaining_high("2026-10-05")
        # Only current open/changed rows remain in this synthetic export; closed
        # records missing through retention must come from September's workbook.
        kept = [row for row in self.rows if row["Issue ID"] in {"HIGH-95", "MEDIUM-0", "LOW-0"}]
        kept.append(
            {
                "Issue ID": "HIGH-new",
                "Title": "Third party service account with high privileges",
                "Severity": "HIGH",
                "Status": "RESOLVED",
                "Created At": "2026-10-01",
                "Resolved Time": "2026-10-06",
            }
        )
        save_csv(self.csv, kept)
        october = run_monthly(
            "2026-10",
            september.deck,
            self.root,
            inputs=[self.csv],
            previous_report=september.workbook,
        )
        self.assertEqual(october.workbook.name, "Oct26_ExampleHealth_Governance_Report.xlsx")
        self.assertEqual(
            october.deck.name, "ExampleHealth_Remediation_Executive_Dashboard_10.15.26.pptx"
        )
        data = read_workbook(october.workbook)
        self.assertEqual(data.closed_at(data.latest), 103)
        self.assertEqual(data.open_at(data.previous), 1)
        self.assertEqual(data.open_at(data.latest), 0)
        history = template_history(Presentation(october.deck), None)
        self.assertEqual(list(history.values()), [102, 95, 40, 8, 5, 2, 1, 0])
        self.assertEqual(before, (september.workbook.read_bytes(), september.deck.read_bytes()))

    def test_manual_workbook_correction_is_copied_without_rebuilding(self):
        existing = self.run_september()
        self.close_remaining_high()
        corrected = self.root / "corrected.xlsx"
        build_monthly_workbook(
            [self.csv], [], corrected, reporting_month("2026-09"), "Example Provider"
        )
        save_csv(self.csv, synthetic_rows())  # Old exports must not undo a workbook correction.
        expected = corrected.read_bytes()
        result = run_monthly("2026-09", existing.deck, self.root, workbook=corrected)
        self.assertEqual(result.workbook.read_bytes(), expected)
        self.assertEqual(shape(Presentation(result.deck).slides[2], 9).text, "0")
        self.assertEqual(corrected.read_bytes(), expected)
        self.assertEqual(
            json.loads(
                Presentation(result.deck)
                .slides[0]
                .notes_slide.notes_text_frame.text.split("\n", 1)[1]
            )["workbook_sha256"],
            digest(result.workbook),
        )
        # Editing/saving the managed workbook itself is also an explicit supported input.
        repeat = run_monthly("2026-09", existing.deck, self.root, workbook=existing.workbook)
        self.assertEqual(repeat.workbook.read_bytes(), expected)

    def test_approved_baseline_and_previous_month_survive_latest_lifecycle_changes(self):
        september = self.run_september()
        original = read_workbook(september.workbook)
        self.close_remaining_high("2026-10-05")
        changed = self.rows[0]
        changed["Created At"] = "2026-04-01"
        changed["Resolved Time"] = "2026-10-05"
        save_csv(self.csv, self.rows)
        october = run_monthly(
            "2026-10",
            september.deck,
            self.root,
            inputs=[self.csv],
            previous_report=september.workbook,
        )
        data = read_workbook(october.workbook)
        self.assertEqual(data.snapshots[data.dates[0]], original.snapshots[original.dates[0]])
        self.assertEqual(data.snapshots[data.previous], original.snapshots[original.latest])
        self.assertEqual(data.open_at(data.previous), 1)
        self.assertEqual(data.open_at(data.latest), 0)
        wb = load_workbook(october.workbook)
        try:
            self.assertEqual(wb["Monthly Scope View"]["B4"].value, 6)
            self.assertEqual(wb["Monthly Scope View"]["D4"].value, 0)
            self.assertEqual(wb["Monthly Scope View"]["D5"].value, 1)
            self.assertEqual(
                wb["CRITICAL Issues"]["G6"].value, "Scope: 0 | Closed: 0 | Remaining: 0"
            )
        finally:
            wb.close()
        # A same-month refresh can use its already-preserved columns without
        # requiring the previous report to be supplied again.
        again = run_monthly("2026-10", october.deck, self.root, inputs=[self.csv])
        self.assertEqual(read_workbook(again.workbook).snapshots, data.snapshots)

    def test_deck_generation_failure_leaves_both_previous_outputs_unchanged(self):
        existing = self.run_september()
        before = (existing.workbook.read_bytes(), existing.deck.read_bytes())
        self.close_remaining_high()
        with patch(
            "wizexpert.governance.monthly.build_deck",
            side_effect=DashboardError("unsupported template"),
        ), self.assertRaisesRegex(DashboardError, "unsupported template"):
            self.run_september()
        self.assertEqual(before, (existing.workbook.read_bytes(), existing.deck.read_bytes()))
        self.assertFalse((self.root / ".wiz-monthly-2026-09.lock").exists())

    def test_second_publication_failure_restores_previous_workbook_and_keeps_backups(self):
        existing = self.run_september()
        before = (existing.workbook.read_bytes(), existing.deck.read_bytes())
        self.close_remaining_high()
        original_replace = os.replace

        def fail_deck(source, target):
            if Path(target) == existing.deck and Path(source).parent.name.startswith(
                ".wiz-monthly-"
            ):
                raise PermissionError("PowerPoint is open")
            return original_replace(source, target)

        with (
            patch("wizexpert.governance.monthly_files.os.replace", side_effect=fail_deck),
            self.assertRaisesRegex(DashboardError, "rolled back"),
        ):
            self.run_september()
        self.assertEqual(before, (existing.workbook.read_bytes(), existing.deck.read_bytes()))
        records = [
            json.loads(path.read_text())["status"]
            for path in (self.root / "report-backups").glob("*/pair.monthly-run.json")
        ]
        self.assertEqual(sorted(records), ["complete", "rolled-back"])

    def test_invalid_inputs_month_or_snapshot_never_publish(self):
        for month in ("2026-13", "../2026", "2026-03"):
            with self.subTest(month=month), self.assertRaises(DashboardError):
                run_monthly(
                    month, self.template, self.root, inputs=[self.csv], history=self.history
                )
        self.rows[0]["Created At"] = "invalid"
        save_csv(self.csv, self.rows)
        with self.assertRaisesRegex(DashboardError, "Created At"):
            self.run_september()
        self.assertFalse(list(self.root.glob("*_Governance_Report.xlsx")))
        save_csv(self.csv, synthetic_rows())
        existing = self.run_september()
        with self.assertRaisesRegex(DashboardError, "does not match"):
            run_monthly("2026-10", existing.deck, self.root, workbook=existing.workbook)

    def test_existing_excel_can_be_adopted_explicitly_but_orphan_pair_is_not_overwritten(self):
        excel = self.root / "Sep26_ExampleHealth_Governance_Report.xlsx"
        build_monthly_workbook(
            [self.csv], [], excel, reporting_month("2026-09"), "Example Provider"
        )
        before = excel.read_bytes()
        with self.assertRaisesRegex(DashboardError, "Only one monthly output"):
            self.run_september()
        result = run_monthly(
            "2026-09", self.template, self.root, workbook=excel, history=self.history
        )
        self.assertEqual(excel.read_bytes(), before)
        self.assertEqual((result.archive / excel.name).read_bytes(), before)
        self.assertTrue(result.deck.is_file())

    def test_concurrent_run_lock_and_interrupted_publication_block_refresh(self):
        with (
            monthly_lock(self.root, "2026-09"),
            self.assertRaisesRegex(DashboardError, "lock already exists"),
        ):
            self.run_september()
        archive = self.root / "report-backups" / "2026-09-interrupted"
        archive.mkdir(parents=True)
        journal = archive / "pair.monthly-run.json"
        journal.write_text(json.dumps({"status": "publishing"}))
        with self.assertRaisesRegex(DashboardError, "needs recovery"):
            self.run_september()


class PairPublicationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.stage = self.root / "stage"
        self.stage.mkdir()
        self.targets = (self.root / "report.xlsx", self.root / "deck.pptx")
        self.staged = (self.stage / "report.xlsx", self.stage / "deck.pptx")
        for path in self.staged:
            path.write_bytes(b"new")

    def test_first_run_partial_publish_is_removed_if_second_file_fails(self):
        expected = {p: None for p in self.targets}
        original_link = os.link

        def fail_second(source, target):
            if target == self.targets[1]:
                raise PermissionError("cannot create deck")
            return original_link(source, target)

        with (
            patch("wizexpert.governance.monthly_files.os.link", side_effect=fail_second),
            self.assertRaisesRegex(DashboardError, "rolled back"),
        ):
            publish_pair(self.staged, self.targets, expected, "2026-09")
        self.assertTrue(all(not p.exists() for p in self.targets))

    def test_external_edits_are_not_overwritten(self):
        for path in self.targets:
            path.write_bytes(b"old")
        expected = {p: digest(p) for p in self.targets}
        self.targets[0].write_bytes(b"external edit")
        with self.assertRaisesRegex(DashboardError, "changed during generation"):
            publish_pair(self.staged, self.targets, expected, "2026-09")
        self.assertEqual(self.targets[0].read_bytes(), b"external edit")
        self.assertEqual(self.targets[1].read_bytes(), b"old")

    def test_backup_failure_never_replaces_either_output(self):
        for path in self.targets:
            path.write_bytes(b"old")
        expected = {p: digest(p) for p in self.targets}
        with patch(
            "wizexpert.governance.monthly_files.shutil.copy2",
            side_effect=OSError("backup disk full"),
        ), self.assertRaisesRegex(OSError, "backup disk full"):
            publish_pair(self.staged, self.targets, expected, "2026-09")
        self.assertTrue(all(p.read_bytes() == b"old" for p in self.targets))

    def test_failed_rollback_retains_recovery_record_and_original_backups(self):
        for path in self.targets:
            path.write_bytes(b"old")
        expected = {p: digest(p) for p in self.targets}
        original_replace = os.replace

        def fail_publication_and_restore(source, target):
            if Path(target) == self.targets[1] or Path(source).name.startswith("restore-"):
                raise PermissionError("locked")
            return original_replace(source, target)

        with patch(
            "wizexpert.governance.monthly_files.os.replace",
            side_effect=fail_publication_and_restore,
        ), self.assertRaisesRegex(DashboardError, "manual recovery"):
            publish_pair(self.staged, self.targets, expected, "2026-09")
        archive = next((self.root / "report-backups").iterdir())
        self.assertEqual((archive / "report.xlsx").read_bytes(), b"old")
        self.assertEqual((archive / "deck.pptx").read_bytes(), b"old")
        self.assertEqual(
            json.loads((archive / "pair.monthly-run.json").read_text())["status"], "publishing"
        )


if __name__ == "__main__":
    unittest.main()
