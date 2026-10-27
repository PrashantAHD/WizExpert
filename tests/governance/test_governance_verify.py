import json
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from PIL import ImageFont
from pptx import Presentation
from test_dashboard import history_payload, make_template
from test_governance import group, snapshot
from test_governance import make_template as make_operational_template
from test_monthly import save_csv, synthetic_rows

from wizexpert.__main__ import main
from wizexpert.governance.dashboard_data import DashboardError
from wizexpert.governance.governance import build_governance
from wizexpert.governance.governance_data import cutoff_counts
from wizexpert.governance.monthly import build_monthly_workbook, run_monthly
from wizexpert.governance.monthly_files import digest
from wizexpert.governance.profile import active, save
from wizexpert.governance.verify import package_signature, verify_preparation


@pytest.fixture
def prepared(tmp_path):
    source = tmp_path / "current.csv"
    rows = synthetic_rows()
    medium = next(row for row in rows if row["Severity"] == "MEDIUM")
    medium.update(Status="RESOLVED", **{"Resolved Time": "2026-10-01T00:00:00Z"})
    save_csv(source, rows)
    previous = tmp_path / "approved-september.xlsx"
    build_monthly_workbook(
        [source], [], previous, datetime(2026, 9, 1), "PROPRIETARY TO EXAMPLE PROVIDER"
    )
    executive_template = tmp_path / "executive-template.pptx"
    make_template(executive_template)
    history = tmp_path / "history.json"
    history.write_text(json.dumps(history_payload(executive_template)), encoding="utf-8")
    operational_template = tmp_path / "operational-template.pptx"
    make_operational_template(operational_template)
    payload = snapshot(operational_template)
    payload["period"] = {
        "start_exclusive": "2026-09-15T23:59:59Z",
        "end_inclusive": "2026-10-15T23:59:59Z",
        "resolved_through": "2026-10-15T23:59:59Z",
        "backlog_as_of": "2026-10-15T23:59:59Z",
        "collected_from": "2026-10-16T00:00:00Z",
        "collected_through": "2026-10-16T00:01:00Z",
    }
    payload["prior"]["date"] = "2026-09-15"
    payload.update(
        cutoff_counts(rows, datetime(2026, 9, 15, 23, 59, 59), datetime(2026, 10, 15, 23, 59, 59))
    )
    resolved = group("resolved", count=1)
    resolved["samples"][0].update(
        id=medium["Issue ID"],
        created_at="2026-01-01T00:00:00Z",
        resolved_at=medium["Resolved Time"],
    )
    payload["resolved_groups"] = {"total_groups": 1, "groups": [resolved]}
    payload["in_progress_groups"] = {"total_groups": 0, "groups": []}
    snapshot_path = tmp_path / "snapshot.json"
    snapshot_path.write_text(json.dumps(payload), encoding="utf-8")
    font = ImageFont.load_default(size=24)
    with (
        patch(
            "wizexpert.governance.dashboard.chart_fonts",
            return_value=(Path("regular"), Path("bold")),
        ),
        patch("wizexpert.governance.dashboard_charts.ImageFont.truetype", return_value=font),
        patch("wizexpert.governance.governance_data.datetime", wraps=datetime) as clock,
    ):
        clock.now.return_value = datetime(2030, 1, 1, tzinfo=UTC)
        pair = run_monthly(
            "2026-10",
            executive_template,
            tmp_path,
            inputs=[source],
            previous_report=previous,
            history=history,
        )
        operational = build_governance(
            snapshot_path, operational_template, tmp_path / "operational.pptx"
        )
        yield {
            "inputs": [source],
            "previous_report": previous,
            "workbook": pair.workbook,
            "executive_template": executive_template,
            "executive": pair.deck,
            "snapshot": snapshot_path,
            "operational_template": operational_template,
            "operational": operational,
            "history": history,
        }


def test_complete_replay_is_read_only_and_not_release_approval(prepared):
    files = [value for value in prepared.values() if isinstance(value, Path)]
    before = {path: digest(path) for path in files}
    result = verify_preparation(**prepared)
    assert result["status"] == "automated_checks_passed"
    assert result["share_ready"] is False
    assert result["current_unique_issues"] == 104
    assert result["counts"]["resolved"]["MEDIUM"] == 1
    assert len(result["manual_gates_required"]) == 6
    assert before == {path: digest(path) for path in files}


def test_history_reuse_requires_reconciliation_record(prepared):
    payload = json.loads(prepared["snapshot"].read_text())
    payload["collection_strategy"] = "reuse_history"
    prepared["snapshot"].write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(DashboardError, match="requires its reconciliation audit"):
        verify_preparation(**prepared)


@pytest.mark.parametrize("change", ["duplicate", "informational", "counts"])
def test_source_scope_and_count_drift_fail_closed(prepared, change):
    from wizexpert.governance.wiz_report import load_wiz_csv

    path = prepared["inputs"][0]
    rows = load_wiz_csv(str(path))
    if change == "duplicate":
        rows.append(dict(rows[0]))
        message = "Duplicate current Issue ID"
    elif change == "informational":
        rows.append({**rows[0], "Issue ID": "info", "Severity": "INFO"})
        message = "informational findings"
    else:
        rows.append({**rows[-1], "Issue ID": "new-open"})
        message = "disagree on open"
    save_csv(path, rows)
    with pytest.raises(DashboardError, match=message):
        verify_preparation(**prepared)


@pytest.mark.parametrize("target", ["workbook", "executive", "operational"])
def test_report_content_or_template_drift_is_blocked(prepared, target):
    path = prepared[target]
    if target == "workbook":
        workbook = load_workbook(path)
        workbook["MEDIUM Issues"]["D7"] = 999
        workbook.save(path)
        workbook.close()
    else:
        deck = Presentation(path)
        box = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame)
        box.text = "Unreviewed replacement title"
        deck.save(path)
    before = digest(path)
    with pytest.raises(DashboardError, match="differs from its reviewed-input replay"):
        verify_preparation(**prepared)
    assert digest(path) == before


def test_missing_assignee_actions_are_not_silently_skipped(prepared):
    path = prepared["snapshot"]
    payload = json.loads(path.read_text())
    payload["assignment_gaps"] = [
        {
            "issue_id": "MEDIUM-0",
            "rule_id": "resolved",
            "rule_name": "Synthetic rule resolved",
            "status": "RESOLVED",
            "suggested_owner": None,
            "suggested_owner_id": None,
            "basis": "No team activity in retained history",
            "event_at": None,
        }
    ]
    content = json.dumps(payload)
    path.write_text(content, encoding="utf-8")
    prepared["operational"].with_suffix(".governance.json").write_text(content, encoding="utf-8")
    with pytest.raises(DashboardError, match="actions CSV is required"):
        verify_preparation(**prepared)


def test_companion_and_cutoff_mismatch_block_verification(prepared):
    companion = prepared["operational"].with_suffix(".governance.json")
    companion.write_text("{}", encoding="utf-8")
    with pytest.raises(DashboardError, match="companion differs"):
        verify_preparation(**prepared)


def test_package_comparison_ignores_xml_prefixes_not_values(tmp_path):
    paths = [tmp_path / f"{name}.xlsx" for name in ("first", "second", "changed")]
    for path, xml in zip(
        paths,
        (
            '<root xmlns="urn:test"><value b="2" a="1">one</value></root>',
            '<n:root xmlns:n="urn:test"><n:value a="1" b="2">one</n:value></n:root>',
            '<root xmlns="urn:test"><value a="1" b="2">two</value></root>',
        ),
        strict=True,
    ):
        with ZipFile(path, "w") as archive:
            archive.writestr("sheet.xml", xml)
    assert package_signature(paths[0]) == package_signature(paths[1])
    assert package_signature(paths[0]) != package_signature(paths[2])


def test_package_comparison_rejects_xml_entities_before_canonicalization(tmp_path):
    path = tmp_path / "unsafe.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr(
            "sheet.xml", "<!DOCTYPE root [<!ENTITY value 'expanded'>]><root>&value;</root>"
        )
    with pytest.raises(DefusedXmlException):
        package_signature(path)


def test_verify_cli_is_discoverable_and_missing_arguments_fail(capsys):
    assert main(["governance"]) == 0
    assert "verify" in capsys.readouterr().out
    with pytest.raises(SystemExit) as error:
        main(["governance", "verify"])
    assert error.value.code == 2


def test_verify_cli_returns_machine_readable_gate_result(prepared, tmp_path, capsys):
    profile_path = save(active(), tmp_path / "profile.json")
    review = tmp_path / "workflow-review.json"
    review.write_text('{"synthetic": true}', encoding="utf-8")
    arguments = [
        "governance",
        "verify",
        "--profile",
        str(profile_path),
        "--workflow-review",
        str(review),
    ]
    for name, value in prepared.items():
        if name == "inputs":
            arguments.extend(["--input", *(str(path) for path in value)])
        else:
            arguments.extend(["--" + name.replace("_", "-"), str(value)])
    capsys.readouterr()

    def rebuilding_with_warning(*args, **kwargs):
        print("Warning: verify export completeness.")
        return build_monthly_workbook(*args, **kwargs)

    with (
        patch(
            "wizexpert.governance.verify.build_monthly_workbook",
            side_effect=rebuilding_with_warning,
        ),
        patch(
            "wizexpert.workflow_qa.require_generation_review",
            return_value={"status": "workflow_checks_passed"},
        ) as workflow_qa,
    ):
        assert main(arguments) == 0
    assert workflow_qa.call_args.args[0] == review
    captured = capsys.readouterr()
    assert "verify export completeness" in captured.err
    result = json.loads(captured.out)
    assert result["status"] == "automated_checks_passed"
    assert result["share_ready"] is False
    assert result["source_and_output_sha256"][str(profile_path)] == digest(profile_path)
    assert result["workflow_qa"]["status"] == "workflow_checks_passed"
