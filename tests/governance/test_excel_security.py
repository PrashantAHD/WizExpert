import unittest
from datetime import datetime
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from wizexpert.governance.excel_cells import write_literal_cell
from wizexpert.governance.wiz_report import (
    build_severity_tab,
    build_summary_tab,
    compute_metrics,
    merge_issue_rows,
    parse_snapshots,
    read_issue_rows_from_existing_tab,
)

HEADERS = ["Issue ID", "Title", "Severity", "Status", "Created At", "Resolved Time"]
SNAPSHOTS = parse_snapshots("Mar 26,Apr 26,May 26")


def csv_issue(severity="HIGH", status="RESOLVED"):
    return {
        "Issue ID": "  =1+1",
        "Title": "\t=2+2",
        "Severity": severity,
        "Status": status,
        "Created At": "2026-03-01T00:00:00Z",
        "Resolved Time": "2026-04-01T00:00:00Z",
    }


class ExcelSecurityTests(unittest.TestCase):
    def assert_no_formula_xml(self, content):
        with ZipFile(BytesIO(content)) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    xml = ElementTree.fromstring(archive.read(name))
                    self.assertFalse(
                        xml.findall(
                            ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f"
                        ),
                        f"Executable formula found in {name}",
                    )

    def round_trip(self, workbook):
        stream = BytesIO()
        workbook.save(stream)
        self.assert_no_formula_xml(stream.getvalue())
        stream.seek(0)
        return load_workbook(stream, data_only=False)

    def test_literal_writer_preserves_strings_numbers_and_dates(self):
        workbook = Workbook()
        values = [
            "=1+1",
            " =2+2",
            "\t=3+3",
            "+1+1",
            "-1+1",
            "@SUM(A1:A2)",
            "#N/A",
            "ordinary text",
            96,
            0.5,
            True,
            datetime(2026, 3, 15),
        ]
        for row, value in enumerate(values, 1):
            write_literal_cell(workbook.active, row, 1, value)
        restored = self.round_trip(workbook)
        for row, value in enumerate(values, 1):
            with self.subTest(value=value):
                cell = restored.active.cell(row, 1)
                self.assertEqual(cell.value, value)
                if isinstance(value, str):
                    self.assertEqual(cell.data_type, "s")
        self.assertEqual(restored.active["A9"].data_type, "n")
        self.assertEqual(restored.active["A12"].data_type, "d")
        restored.close()

    def test_main_generator_writes_issue_fields_and_custom_title_as_text(self):
        for status in ("RESOLVED", "REJECTED"):
            with self.subTest(status=status):
                workbook = Workbook()
                row = csv_issue(status=status)
                build_severity_tab(workbook, "HIGH", [row], SNAPSHOTS)
                build_summary_tab(
                    workbook,
                    {"HIGH": compute_metrics([row], SNAPSHOTS)},
                    SNAPSHOTS,
                    title="=3+3",
                )
                restored = self.round_trip(workbook)
                for coordinate, expected in (("A16", "=1+1"), ("B16", "=2+2")):
                    cell = restored["HIGH Issues"][coordinate]
                    self.assertEqual(cell.value, expected)
                    self.assertEqual(cell.data_type, "s")
                self.assertEqual(restored["Summary"]["A1"].value, "=3+3")
                self.assertEqual(restored["Summary"]["A1"].data_type, "s")
                self.assertEqual(restored["HIGH Issues"]["B5"].data_type, "n")
                restored.close()

    def test_historical_formula_fields_become_literal_when_rebuilt(self):
        workbook = Workbook()
        build_severity_tab(workbook, "HIGH", [csv_issue()], SNAPSHOTS)
        # Simulate a report created by the old vulnerable writer.
        workbook["HIGH Issues"]["A16"] = "=1+1"
        workbook["HIGH Issues"]["B16"] = "=2+2"
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        prior = load_workbook(stream, data_only=False)
        historical = read_issue_rows_from_existing_tab(prior, "HIGH")
        rows, retained = merge_issue_rows(historical, [])
        build_severity_tab(prior, "HIGH", rows, SNAPSHOTS, retained)
        restored = self.round_trip(prior)
        self.assertEqual(restored["HIGH Issues"]["A16"].value, "=1+1")
        self.assertEqual(restored["HIGH Issues"]["B16"].value, "=2+2")
        self.assertEqual(restored["HIGH Issues"]["A16"].data_type, "s")
        self.assertEqual(restored["HIGH Issues"]["B16"].data_type, "s")
        prior.close()
        restored.close()


if __name__ == "__main__":
    unittest.main()
