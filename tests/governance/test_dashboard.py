import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from wizexpert.governance.dashboard import (
    ACHIEVEMENT_TEXT_IDS,
    PROGRESS_TEXT_IDS,
    build_deck,
    change_description,
    shape,
    template_history,
    trend_language,
)
from wizexpert.governance.dashboard_charts import line_chart, severity_chart
from wizexpert.governance.dashboard_data import (
    DashboardError,
    Snapshot,
    combine_trend,
    exposure_callouts,
    history_from_json,
    read_workbook,
)

BASELINE = date(2026, 3, 15)
PREVIOUS = date(2026, 8, 15)
LATEST = date(2026, 9, 15)


def make_workbook(path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    # Synthetic, deliberately different from client figures.
    sets = {
        "CRITICAL": [Snapshot(6, 6, 0, 0), Snapshot(3, 1, 1, 1), Snapshot(3, 0, 2, 1)],
        "HIGH": [Snapshot(97, 96, 1, 0), Snapshot(5, 2, 2, 1), Snapshot(5, 0, 4, 1)],
    }
    titles = [
        "Publicly readable bucket contains sensitive data",
        "VM infected with a hacking tool",
        "Third party service account with high privileges",
        "Sensitive data can be accessed by all users including external users",
    ]
    for summary_row, (severity, values) in enumerate(sets.items(), 5):
        ws = wb.create_sheet(f"{severity} Issues")
        for row, label in enumerate(
            (
                "Total Issues Created",
                "Open / In Progress",
                "Resolved (Cumulative)",
                "Ignored / Accepted Risk",
                "Resolution %",
            ),
            5,
        ):
            ws.cell(row, 1, label)
        for col, (day, snapshot) in enumerate(
            zip((BASELINE, PREVIOUS, LATEST), values, strict=False), 2
        ):
            ws.cell(4, col, day.strftime("%b %d, %Y"))
            for row, value in enumerate(
                (snapshot.total, snapshot.open, snapshot.resolved, snapshot.ignored), 5
            ):
                ws.cell(row, col, value)
            ws.cell(9, col, f"{snapshot.addressed_pct:.1f}%")
        ws.cell(4, 5, "vs Baseline")
        for col, label in enumerate(
            (
                "Issue ID",
                "Title",
                "Severity",
                "Status",
                "Created At",
                "Resolved Time",
                "Snapshot Month",
            ),
            1,
        ):
            ws.cell(15, col, label)
        current = values[-1]
        for i in range(current.closed):
            row = [
                f"{severity}-{i}",
                titles[i % len(titles)],
                severity,
                "RESOLVED" if i < current.resolved else "IGNORED",
                "2026-01-01",
                "2026-08-24 21:54",
                "September 2026",
            ]
            for col, value in enumerate(row, 1):
                ws.cell(16 + i, col, value)
        for col, value in enumerate(
            (
                severity,
                current.total,
                current.open,
                current.resolved,
                current.ignored,
                f"{current.addressed_pct:.1f}%",
            ),
            1,
        ):
            summary.cell(summary_row, col, value)
    wb.save(path)
    wb.close()


def make_template(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = 12192000, 6858000
    for ids in ({2, 5}, PROGRESS_TEXT_IDS, ACHIEVEMENT_TEXT_IDS):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for sid in sorted(ids):
            text = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.8))
            text._element.xpath(".//p:cNvPr")[0].set("id", str(sid))
            run = text.text_frame.paragraphs[0].add_run()
            run.text = "Template label"
            run.font.name, run.font.size, run.font.bold = "Arial", Pt(12), True
            run.font.color.rgb = RGBColor.from_string("16324F")
    shape(deck.slides[0], 5).text_frame.paragraphs[0].runs[
        0
    ].text = "Example Health Executive Dashboard"
    shape(deck.slides[1], 3).text_frame.paragraphs[0].runs[
        0
    ].text = "Example Health Remediation Progress"
    shape(deck.slides[2], 3).text_frame.paragraphs[0].runs[0].text = "Risk Reduction Achievements"
    stream = BytesIO()
    Image.new("RGB", (300, 100), "white").save(stream, format="PNG")
    for index, sid in ((1, 31), (2, 44)):
        stream.seek(0)
        picture = deck.slides[index].shapes.add_picture(
            stream, Inches(0.9), Inches(3.9), Inches(5), Inches(2)
        )
        picture._element.xpath(".//p:cNvPr")[0].set("id", str(sid))
    deck.save(path)


def history_payload(template):
    return {
        "schema_version": 1,
        "scope": ["CRITICAL", "HIGH"],
        "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
        "points": [
            {"date": date(2026, month, 15).isoformat(), "open": value}
            for month, value in zip(range(3, 9), (102, 95, 40, 8, 5, 99), strict=False)
        ],
    }


class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.workbook, self.template = self.root / "report.xlsx", self.root / "approved.pptx"
        self.history = self.root / "approved-history.json"
        make_workbook(self.workbook)
        make_template(self.template)
        self.history.write_text(json.dumps(history_payload(self.template)), encoding="utf-8")
        self.data = read_workbook(self.workbook)

    def edit_workbook(self, tab, cell, value):
        wb = load_workbook(self.workbook)
        wb[tab][cell] = value
        wb.save(self.workbook)
        wb.close()

    def render(self, **kwargs):
        font = ImageFont.load_default(size=24)
        with (
            patch(
                "wizexpert.governance.dashboard.chart_fonts",
                return_value=(Path("regular"), Path("bold")),
            ),
            patch("wizexpert.governance.dashboard_charts.ImageFont.truetype", return_value=font),
        ):
            return build_deck(self.workbook, self.template, history_path=self.history, **kwargs)

    def test_snapshot_math_and_fixed_baseline_history_precedence(self):
        self.assertEqual(self.data.closed_at(LATEST), 8)
        self.assertEqual(self.data.closed_at(LATEST) - self.data.closed_at(PREVIOUS), 3)
        history = history_from_json(self.history.read_text())
        points = combine_trend(self.data, history)
        self.assertEqual(list(points.values()), [102, 95, 40, 8, 5, 3, 0])
        restated = dict(self.data.snapshots)
        restated[BASELINE] = {"CRITICAL": Snapshot(7, 7, 0, 0), "HIGH": Snapshot(98, 98, 0, 0)}
        self.assertEqual(
            combine_trend(replace(self.data, snapshots=restated), history)[BASELINE], 102
        )
        rolled = {day: values for day, values in restated.items() if day != BASELINE}
        self.assertEqual(
            combine_trend(replace(self.data, snapshots=rolled), history)[BASELINE], 102
        )

    def test_invalid_history_is_explicit(self):
        payload = history_payload(self.template)
        for value in (-1, True, "5", 1.2):
            with self.subTest(value=value):
                payload["points"][-1]["open"] = value
                with self.assertRaises(DashboardError):
                    history_from_json(json.dumps(payload))
        payload = history_payload(self.template)
        payload["points"][0]["open"] = 116
        with self.assertRaisesRegex(DashboardError, "102"):
            history_from_json(json.dumps(payload))
        payload = history_payload(self.template)
        payload["points"].append(payload["points"][1])
        with self.assertRaisesRegex(DashboardError, "duplicate"):
            history_from_json(json.dumps(payload))
        with self.assertRaisesRegex(DashboardError, "Invalid history JSON"):
            history_from_json("{broken")

    def test_missing_or_future_months_are_not_guessed(self):
        history = history_from_json(self.history.read_text())
        del history[date(2026, 5, 15)]
        with self.assertRaisesRegex(DashboardError, "Missing approved trend"):
            combine_trend(self.data, history)
        history[date(2026, 10, 15)] = 0
        with self.assertRaisesRegex(DashboardError, "newer than"):
            combine_trend(self.data, history)

    def test_bad_workbook_counts_formulas_dates_and_summary_rejected(self):
        cases = [
            ("HIGH Issues", "D5", -1, "whole number"),
            ("HIGH Issues", "D6", "=1+1", "whole number"),
            ("HIGH Issues", "D5", 99, "does not equal"),
            ("HIGH Issues", "D9", "95.6%", "percentage"),
            ("HIGH Issues", "D4", "Sep 16, 2026", "15th"),
            ("HIGH Issues", "D4", "bad date", "header"),
            ("Summary", "B6", 99, "disagree"),
            ("HIGH Issues", "F16", "2026-10-01", "detail counts"),
            ("HIGH Issues", "F16", "2025-01-01", "precedes creation"),
            ("HIGH Issues", "A17", "HIGH-0", "Duplicate"),
        ]
        for tab, cell, value, error in cases:
            with self.subTest(cell=cell, value=value):
                make_workbook(self.workbook)
                self.edit_workbook(tab, cell, value)
                with self.assertRaisesRegex(DashboardError, error):
                    read_workbook(self.workbook)

    def test_missing_sheet_and_nonconsecutive_dates_rejected(self):
        wb = load_workbook(self.workbook)
        del wb["HIGH Issues"]
        wb.save(self.workbook)
        wb.close()
        with self.assertRaisesRegex(DashboardError, "Missing required"):
            read_workbook(self.workbook)
        make_workbook(self.workbook)
        for severity in ("CRITICAL", "HIGH"):
            self.edit_workbook(f"{severity} Issues", "C4", "Jul 15, 2026")
        with self.assertRaisesRegex(DashboardError, "immediately preceding"):
            read_workbook(self.workbook)

    def test_exposure_claims_use_only_resolved_cutoff_evidence(self):
        callouts = exposure_callouts(self.data)
        self.assertEqual(len(callouts), 3)
        self.assertTrue(all(f.status == "RESOLVED" for c in callouts for f in c.evidence))
        ignored = tuple(replace(f, status="IGNORED") for f in self.data.findings)
        with self.assertRaisesRegex(DashboardError, "Fewer than two"):
            exposure_callouts(replace(self.data, findings=ignored))
        future = tuple(replace(f, closed_on=date(2026, 10, 1)) for f in self.data.findings)
        with self.assertRaises(DashboardError):
            exposure_callouts(replace(self.data, findings=future))

    def test_flat_increase_and_zero_denominator_wording(self):
        for before, after, expected in (
            (0, 0, "unchanged"),
            (0, 2, "increased"),
            (4, 5, "increased"),
            (3, 0, "declining"),
        ):
            snapshots = dict(self.data.snapshots)
            snapshots[PREVIOUS] = {
                "HIGH": Snapshot(before, before, 0, 0),
                "CRITICAL": Snapshot(0, 0, 0, 0),
            }
            snapshots[LATEST] = {
                "HIGH": Snapshot(after, after, 0, 0),
                "CRITICAL": Snapshot(0, 0, 0, 0),
            }
            self.assertIn(expected, trend_language(replace(self.data, snapshots=snapshots))[0])
        self.assertEqual(change_description(0, 2), "increase from zero")
        self.assertEqual(change_description(0, 0), "no change")
        self.assertEqual(change_description(4, 5), "25% increase")
        self.assertEqual(change_description(3, 0), "100% reduction")

    def test_template_required_and_history_bound_to_correct_template(self):
        with self.assertRaisesRegex(DashboardError, "template not found"):
            build_deck(self.workbook, self.root / "missing.pptx")
        with self.assertRaisesRegex(DashboardError, "raster chart"):
            build_deck(self.workbook, self.template)
        payload = history_payload(self.template)
        payload["template_sha256"] = "wrong"
        self.history.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaisesRegex(DashboardError, "does not match"):
            self.render()

    def test_serialized_deck_values_geometry_style_sources_and_next_month_reuse(self):
        originals = [p.read_bytes() for p in (self.workbook, self.template, self.history)]
        output = self.render()
        self.assertEqual(output.name, "ExampleHealth_Remediation_Executive_Dashboard_9.15.26.pptx")
        self.assertEqual(
            originals, [p.read_bytes() for p in (self.workbook, self.template, self.history)]
        )
        deck = Presentation(output)
        before = Presentation(self.template)
        self.assertEqual(len(deck.slides), 3)
        for slide, original in zip(deck.slides, before.slides, strict=False):
            self.assertEqual(
                [(s.shape_id, s.left, s.top, s.width, s.height) for s in slide.shapes],
                [(s.shape_id, s.left, s.top, s.width, s.height) for s in original.shapes],
            )
        expected = {6: "102 fewer", 9: "100%", 12: "+3", 15: "-3"}
        for sid, value in expected.items():
            target = shape(deck.slides[1], sid)
            self.assertEqual(target.text, value)
            font = target.text_frame.paragraphs[0].runs[0].font
            self.assertEqual(
                (font.name, font.size, font.bold, str(font.color.rgb)),
                ("Arial", Pt(12), True, "16324F"),
            )
        self.assertEqual(shape(deck.slides[2], 6).text, "8")
        self.assertEqual(shape(deck.slides[2], 9).text, "0")
        self.assertEqual(shape(deck.slides[2], 12).text, "100%")
        self.assertEqual(
            shape(deck.slides[2], 43).text,
            "Accepted risk = team has formally accepted the risk and committed to a fix-by date.",
        )
        self.assertIn("5 of 5", shape(deck.slides[2], 13).text)
        self.assertEqual(list(template_history(deck, None).values()), [102, 95, 40, 8, 5, 3, 0])
        notes = deck.slides[2].notes_slide.notes_text_frame.text
        self.assertIn("workbook_sha256", notes)
        self.assertIn("CRITICAL Issues!A", notes)
        self.assertNotEqual(
            shape(deck.slides[1], 31).image.blob, shape(before.slides[1], 31).image.blob
        )
        font = ImageFont.load_default(size=24)
        with (
            patch(
                "wizexpert.governance.dashboard.chart_fonts",
                return_value=(Path("regular"), Path("bold")),
            ),
            patch("wizexpert.governance.dashboard_charts.ImageFont.truetype", return_value=font),
        ):
            regenerated = build_deck(self.workbook, output, self.root / "reused.pptx")
        self.assertEqual(
            template_history(Presentation(regenerated), None), template_history(deck, None)
        )

    def test_no_overwrite_and_wrong_template_rejected(self):
        existing = self.root / "already.pptx"
        existing.write_bytes(b"original")
        with self.assertRaisesRegex(DashboardError, "already exists"):
            self.render(output=existing)
        self.assertEqual(existing.read_bytes(), b"original")
        with self.assertRaisesRegex(DashboardError, "must not overwrite"):
            self.render(output=self.template)
        deck = Presentation(self.template)
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(self.template)
        with self.assertRaisesRegex(DashboardError, "cover plus two"):
            self.render()

    def test_chart_labels_and_pixels_have_expected_data(self):
        font = ImageFont.load_default(size=24)
        points = combine_trend(self.data, history_from_json(self.history.read_text()))
        from wizexpert.governance import dashboard_charts

        with (
            patch("wizexpert.governance.dashboard_charts.ImageFont.truetype", return_value=font),
            patch(
                "wizexpert.governance.dashboard_charts.text", wraps=dashboard_charts.text
            ) as drawn,
        ):
            line = Image.open(
                line_chart(points, "Open findings declined this month.", (Path("a"), Path("b")))
            )
            self.assertEqual(line.size, (1343, 513))
            labels = [call.args[2] for call in drawn.call_args_list]
            self.assertIn("Mar 15", labels)
            self.assertIn("Sep 15", labels)
            self.assertIn(102, labels)
            self.assertIn(95, labels)
            drawn.reset_mock()
            bars = Image.open(
                severity_chart(
                    LATEST,
                    {"CRITICAL": 0, "HIGH": 0},
                    (Path("a"), Path("b")),
                    {"CRITICAL": 6, "HIGH": 96},
                    BASELINE,
                )
            )
            self.assertEqual(bars.size, (1283, 430))
            labels = [call.args[2] for call in drawn.call_args_list]
            self.assertIn(6, labels)
            self.assertIn(96, labels)
            self.assertIn("Sep 15 latest", labels)
            self.assertEqual(bars.getpixel((155, 345)), (18, 58, 97))
            self.assertEqual(bars.getpixel((985, 340)), (255, 255, 255))


if __name__ == "__main__":
    unittest.main()
